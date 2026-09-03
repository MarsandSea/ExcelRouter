"""
ExcelRouter - 图形界面

三步式单窗口布局（面向非技术办公用户，按任务顺序组织界面）：
  ① 选择要拆的表格：「选文件 / 选文件夹」两个直选按钮（类型由路径自动推断），
     选完自动统计文件数、后台扫描字段、自动推荐输出目录；
  ② 按哪个字段拆分：拆分字段下拉（扫描后按关键词智能预选）＋批量选项
     （仅文件夹显示：每组 ZIP / 同时拆到人）；
  ③ 开始拆分：输出目录＋主按钮＋进度＋状态行，固定在窗口底部随时可达。
  「▸ 高级设置」折叠在步骤区末尾；「▸ 处理详情」日志默认收起，出错自动展开。

拆分在子线程里跑，日志/进度经 queue 由主线程 UI 泵（_pump_ui，每 100ms 批量刷新）
更新界面——子线程绝不直接碰 Tk 控件（v2.4 机制，勿回退）。

Copyright (c) 2026 AbeLin
MIT License
"""

import os
import re
import sys
import json
import queue
import threading
import subprocess
import webbrowser
from typing import Any

import customtkinter as ctk
from tkinter import filedialog, messagebox

# 拖拽支持（v2.7）：tkinterdnd2 为可选依赖——打包环境缺它时静默退回无拖拽，
# 不影响其余任何功能。CTk + TkinterDnD 的混编写法是社区通行方案（本地 spike 验证过）。
try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    _DND_OK = True
except Exception:
    TkinterDnD = None  # type: ignore[assignment]
    DND_FILES = None   # type: ignore[assignment]
    _DND_OK = False

if _DND_OK:
    class _RootBase(ctk.CTk, TkinterDnD.DnDWrapper):  # type: ignore[misc]
        """customtkinter 根窗口 + tkdnd 拖放能力的混编基类。"""
else:
    _RootBase = ctk.CTk  # type: ignore[assignment]  # 缺依赖时退回纯 CTk


# =====================================================
# 路径解析（兼容 PyInstaller --onefile 打包）
# =====================================================

def _is_frozen():
    return getattr(sys, 'frozen', False)


def _resource_dir():
    if _is_frozen():
        return sys._MEIPASS   # type: ignore[attr-defined]  # PyInstaller 运行时注入
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _app_dir():
    if _is_frozen():
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


DEFAULT_CONFIG_PATH = os.path.join(_resource_dir(), "config", "default_config.json")
USER_CONFIG_PATH    = os.path.join(_app_dir(), "user_config.json")

COL_PLACEHOLDER = "（选择表格后自动识别）"
PDF_COL_PLACEHOLDER = "（选清单后自动识别）"
PDF_RECV_NONE = "（不需要）"

# UI 泵用的「本轮没有此类消息」哨兵（不能用 None：扫描失败时 payload 可能为空）
_MISSING = object()

APP_VERSION = "2.7.1"
# 匿名反馈问卷地址（问卷 URL 确定后替换此处即可，一行改动 + 打 tag 发版）
FEEDBACK_URL = "https://f.wps.cn/g/pBOAWUQc/"
# 在线 FAQ（「❓ 常见问题」按钮）：链接 Gitee 镜像而非 GitHub——国内办公网络访问
# GitHub 不稳定，用户点帮助按钮却打不开比没有按钮更糟。内容与 GitHub 同步发版。
FAQ_URL = "https://gitee.com/Marsandsea/Excelrouter/blob/main/docs/FAQ.md"

# 状态文字颜色（浅色模式, 深色模式）
C_OK    = ("#15803d", "#4ade80")
C_WARN  = ("#b45309", "#fbbf24")
C_ERR   = ("#b91c1c", "#f87171")
C_MUTED = ("gray40", "gray60")

# ── 品牌调色板（v2.7.1，浅色/深色双主题集中取值）──────────────
# 主色「移动蓝」：使用者是移动系办公人群，品牌联想自然；比 ctk 默认 blue
# (#3B8ED0) 更饱和明亮，主按钮/徽章/进度条统一从这套取色，形成记忆点。
# 完成态用明确的浅绿横幅承载——「显眼」最重要的时刻是完成瞬间。
PRIMARY        = ("#0E7FD1", "#38BDF8")   # 主按钮、步骤徽章、进度条、分段选择高亮
PRIMARY_HOVER  = ("#0B6BB0", "#7DD3FC")
OK_BANNER_BG   = ("#DCFCE7", "#14532D")   # 完成摘要横幅底色（文字仍用 C_OK）
CARD_BORDER    = ("gray82", "gray28")     # 步骤卡片 1px 淡边框
# ACCENT 保留为 PRIMARY 的别名：徽章等既有引用不逐处改动，取色来源已经换血
ACCENT  = PRIMARY

def _init_fonts():
    """把 CTk 默认字体族设为微软雅黑 UI（Windows 中文渲染明显好于 tk 默认 Roboto 回退）。

    CTkFont 未显式传 family 时取 ThemeManager.theme["CTkFont"]["family"]，这是
    唯一真正的全局默认入口（实测改 FontManager._default_font 无效）。主题 dict 是
    set_default_color_theme 时载入的运行时副本，改它对已建/将建组件一致生效；
    无此字体的系统由 tk 字体回退机制自然落到雅黑/宋体，不额外探测。
    """
    try:
        import tkinter.font as tkfont
        if "Microsoft YaHei UI" in list(tkfont.families()):
            ctk.ThemeManager.theme["CTkFont"]["family"] = "Microsoft YaHei UI"
    except Exception:
        pass

def _ghost_button(parent, **kw):
    """描边次要按钮：比主按钮弱一级的操作（浏览 / 扫描字段 / 打开文件夹等）。"""
    return ctk.CTkButton(parent, fg_color="transparent", border_width=1,
                         border_color=("gray60", "gray40"),
                         text_color=("gray15", "gray85"),
                         hover_color=("gray90", "gray25"), **kw)


def _flat_button(parent, **kw):
    """纯文字按钮：折叠开关、页脚工具位等最低视觉权重的操作。"""
    return ctk.CTkButton(parent, fg_color="transparent",
                         text_color=("gray20", "gray80"),
                         hover_color=("gray85", "gray25"), **kw)

# 兜底默认配置（通用、不绑定任何业务）
FALLBACK_CONFIG = {
    "input_path": "", "output_path": "",
    "header_mode": "auto", "header_row": 1,
    "grid_keys": [], "id_keys": [],
    "split_column": "", "selected_values": [],
    "person_column": "", "to_person": False, "person_file_filter": [],
    "value_alias_map": {},
    "skip_values": ["合计", "小计", "总计", "平均", ""],
    "merge_across_files": False,
    "make_zip": True,
    "exact_match": True,
    "preserve_format": True,
    "keep_formulas": False,
    "auto_open_output": True,
    # ── PDF 加密分发（v2.6）──
    "ui_mode": "excel",
    "pdf_input_paths": [],
    "pdf_mapping_path": "",
    "pdf_grid_column": "",
    "pdf_password_column": "",
    "pdf_receiver_column": "",
    "pdf_watermark": True,
    "pdf_watermark_text": "{grid} {date}",
    "pdf_watermark_opacity": 0.15,
    "pdf_watermark_angle": 45,
    # ── 界面个性化（v2.7）──
    "window_geometry": "",     # 上次关闭时的窗口大小位置；空 = 按屏幕自适应
    "ui_scale": 100,           # 界面缩放百分比（大字号=115），重启后生效
    # ── 界面个性化（v2.7.1）──
    "appearance_mode": "system",   # system / light / dark，页脚按钮循环切换
}

ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")


def load_config():
    for path in (USER_CONFIG_PATH, DEFAULT_CONFIG_PATH):
        if os.path.exists(path):
            try:
                with open(path, encoding='utf-8') as f:
                    cfg = json.load(f)
                merged = dict(FALLBACK_CONFIG)
                merged.update(cfg)
                return merged
            except Exception:
                continue
    return dict(FALLBACK_CONFIG)


def save_config(cfg):
    with open(USER_CONFIG_PATH, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


class App(_RootBase):
    def __init__(self):
        super().__init__()
        # 字体统一必须在任何组件创建前、root 建立后执行：
        # tkfont.families() 需要 root；而组件的 CTkFont 在构造时读主题值，晚了改无效。
        _init_fonts()
        if _DND_OK:
            # 加载 tkdnd 二进制（守卫条件已保证 TkinterDnD 非 None）
            self.TkdndVersion = TkinterDnD._require(self)  # type: ignore[union-attr]
        # 窗口标题只放简名：完整品牌名+副标已经在正文首行显示一次，
        # 标题栏再写一遍会让人觉得“这句话重复了”。
        self.title(f"ExcelRouter v{APP_VERSION}")
        self._set_window_icon()
        self._stop_flag = False
        self._running = False
        self.cfg = load_config()
        self._init_geometry()
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self._first_run = not os.path.exists(USER_CONFIG_PATH) and not self.cfg.get("input_path")
        # 界面缩放（大字号）：启动时应用；运行时切换只存配置、下次启动生效
        scale = self.cfg.get("ui_scale", 100)
        if isinstance(scale, int) and scale != 100:
            ctk.set_widget_scaling(scale / 100)
        # 深浅色（v2.7.1）：支持手动固定，不再只能跟随系统
        mode = self.cfg.get("appearance_mode", "system")
        if mode in ("light", "dark"):
            ctk.set_appearance_mode(mode)
        self._mode = self.cfg.get("ui_mode", "excel")   # "excel" / "pdf"
        self._pdf_paths = [p for p in self.cfg.get("pdf_input_paths", []) if os.path.exists(p)]
        self._scanned_map_path = None   # 最近扫描过的映射清单路径（防过期结果）
        self._columns = []
        self._adv_open = False
        self._log_open = False
        self._scanned_path = None     # 最近一次扫描的输入路径（防重复扫描/过期结果）
        self._out_auto = ""           # 最近一次自动推荐的输出目录（用户没改过才允许覆盖）
        self._last_output = None      # 最近一次成功运行的输出目录（「打开输出文件夹」用）
        self._tpl_path = None         # 最近一次扫描识别出的模板文件全路径（取值枚举用）
        self._values_all = []         # 当前拆分字段枚举到的全部取值
        self._value_vars = {}         # 取值 -> BooleanVar（勾选结果写回 selected_values）
        self._values_open = False     # 取值勾选面板展开状态
        self._values_panel = None     # 展开中的勾选面板
        self._enum_token = 0          # 取值枚举序号：防止过期结果回填
        self._enum_after = None       # 防抖 after id
        self._warn_count = 0          # 本轮运行日志里 ⚠ 行数（完成后给「处理详情」打角标）
        self._last_summary = None     # core 收尾的 [SUMMARY] JSON（完成摘要用）
        self._pdf_grid_rows = 0       # 映射清单网格行数（分发预览用）
        self._pdf_n_cols = 0          # 映射清单列数（分发预览用）
        self._ui_q = queue.Queue()            # 子线程 → 主线程的消息队列
        self._prog_indeterminate = False      # 进度条当前是否处于不定态动画
        self._build_ui()
        self._update_input_ui()
        self._apply_mode()
        p = self.cfg.get("input_path", "")
        saved_out = self._output_var.get().strip()
        if not saved_out:
            self._suggest_output()
        elif saved_out and os.path.normcase(saved_out) == os.path.normcase(self._auto_output_for(p) or ""):
            self._out_auto = saved_out    # 上次存的就是自动值：换输入时允许跟着更新
        if p and os.path.exists(p):
            self._scan_input()        # 上次的输入还在：启动即自动扫描，打开就能直接开始
        mp = self.cfg.get("pdf_mapping_path", "")
        if mp and os.path.exists(mp):
            self._scan_mapping()      # 上次的映射清单还在：同样启动即恢复列下拉
        if self._first_run and not p:
            # 首启一句话引导（D3）：三步卡片本身就是向导，不再做独立新手页
            tip = "点「📄 选一个 Excel 文件」"
            if _DND_OK:
                tip += "，或直接把文件 / 文件夹拖进窗口"
            self._set_in_status(f"第一次用？{tip}即可开始", C_MUTED)
        self.after(100, self._pump_ui)

    def _init_geometry(self):
        """窗口初始尺寸：优先恢复上次关闭时的大小位置；首次按屏幕高度自适应。

        目标用户大量使用 1366×768 的老笔记本，固定 820×800 会让窗口下沿
        （含主按钮的操作区）超出屏幕——这是小屏适配的实际 bug，高度必须
        按可用屏高压下来（滚动区会负责吸收内容）。
        """
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        h = min(800, max(600, sh - 90))
        default = f"820x{h}"
        saved = str(self.cfg.get("window_geometry", ""))
        m = re.fullmatch(r"(\d+)x(\d+)([+-]\d+)([+-]\d+)", saved)
        if m and int(m.group(1)) <= sw and int(m.group(2)) <= sh:
            self.geometry(saved)
        else:
            self.geometry(default)
        self.minsize(720, min(640, h))

    def _on_close(self):
        """关闭时静默记住窗口大小位置（下次启动恢复现场），失败不影响退出。"""
        try:
            cfg = self._collect_config()
            cfg["window_geometry"] = self.geometry()
            cfg["ui_scale"] = self.cfg.get("ui_scale", 100)
            save_config(cfg)
        except Exception:
            pass
        self.destroy()

    def _set_window_icon(self):
        """设置标题栏 / 任务栏 / 最小化时的窗口图标。

        `--icon app.ico`（build.bat / CI）只改 exe 文件本身的图标，运行起来的
        Tk 窗口图标要靠 iconbitmap() 单独设置，否则标题栏/任务栏还是 Tk 默认
        羽毛图标——这就是“exe 图标换了，但程序里和最小化时的图标没跟着换”的
        原因。app.ico 需要被 --add-data 打包进去才能在冻结后找到（见 build.bat）。
        """
        for base in (_resource_dir(), _app_dir()):
            ico = os.path.join(base, "app.ico")
            if os.path.exists(ico):
                try:
                    self.iconbitmap(ico)
                    return
                except Exception:
                    continue

    # ── UI 构建 ──────────────────────────────────────────
    def _build_ui(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        self._build_header()                                   # row 0 品牌区（含模式切换）
        self._body = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self._body.grid(row=1, column=0, padx=12, pady=(4, 0), sticky="nsew")
        self._body.grid_columnconfigure(0, weight=1)
        # 两个模式各占一个 Frame（同一格），切换时 grid()/grid_remove()
        self._excel_frame = ctk.CTkFrame(self._body, fg_color="transparent")
        self._excel_frame.grid(row=0, column=0, sticky="ew")
        self._excel_frame.grid_columnconfigure(0, weight=1)
        self._pdf_frame = ctk.CTkFrame(self._body, fg_color="transparent")
        self._pdf_frame.grid(row=0, column=0, sticky="ew")
        self._pdf_frame.grid_columnconfigure(0, weight=1)
        self._build_step_input(self._excel_frame)              # ① 选表格
        self._build_step_field(self._excel_frame)              # ② 选字段（含批量选项）
        self._build_adv_area(self._excel_frame)                # ▸ 高级设置
        self._build_pdf_step_input(self._pdf_frame)            # PDF ① 选 PDF
        self._build_pdf_step_map(self._pdf_frame)              # PDF ② 选映射清单
        self._build_action(self)                               # row 2 ③ 开始（固定底部）
        self._build_bottom(self)                               # row 3-5 工具条 / 日志 / 页脚

    def _build_header(self):
        # 品牌文案已定稿（副标/标语/特性行），只调整排版：作者信息移到页脚
        head = ctk.CTkFrame(self, fg_color="transparent")
        head.grid(row=0, column=0, padx=20, pady=(14, 4), sticky="ew")
        brand = ctk.CTkFrame(head, fg_color="transparent")
        brand.pack(anchor="w")
        # 品牌色块（v2.7.1）：标题左侧的圆角主色标记，最轻量的品牌记忆点
        ctk.CTkFrame(brand, width=6, height=30, corner_radius=3,
                     fg_color=PRIMARY).pack(side="left", padx=(0, 10), pady=(2, 0))
        ctk.CTkLabel(brand, text="ExcelRouter",
                     font=ctk.CTkFont(size=24, weight="bold")).pack(side="left")
        ctk.CTkLabel(brand, text="Excel 智能拆分工具",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=("gray25", "gray75")).pack(side="left", padx=(10, 0), pady=(7, 0))
        ctk.CTkLabel(head, text="整个文件夹一键拆完：按部门、区域、工号等字段自动拆分，打包分发",
                     font=ctk.CTkFont(size=12),
                     text_color=("gray30", "gray70")).pack(anchor="w", pady=(4, 0))
        ctk.CTkLabel(head, text="保留原格式 · 跨文件自动合并 · 单个文件也能拆 · 可同时拆到每个人",
                     font=ctk.CTkFont(size=11), text_color="gray").pack(anchor="w", pady=(1, 0))
        # 功能模式切换：Excel 拆分（默认） / PDF 按网格加密分发
        self._mode_seg = ctk.CTkSegmentedButton(head, values=["Excel 拆分", "PDF 加密分发"],
                                                command=self._on_mode_change,
                                                selected_color=PRIMARY,
                                                selected_hover_color=PRIMARY_HOVER,
                                                height=30, font=ctk.CTkFont(size=12))
        self._mode_seg.set("PDF 加密分发" if self._mode == "pdf" else "Excel 拆分")
        self._mode_seg.pack(anchor="w", pady=(8, 0))

    def _step_card(self, parent, row, num, title, padx=8, pady=(0, 10)):
        """带编号徽章的步骤卡片：编号即真实操作顺序，是界面的导航主线。

        返回 (card, 标题 Label)：标题引用给需要随模式改文案的卡片用（如③操作卡）。
        """
        card = ctk.CTkFrame(parent, corner_radius=12, border_width=1,
                            border_color=CARD_BORDER)
        card.grid(row=row, column=0, padx=padx, pady=pady, sticky="ew")
        card.grid_columnconfigure(0, weight=1)
        bar = ctk.CTkFrame(card, fg_color="transparent")
        bar.grid(row=0, column=0, padx=12, pady=(10, 2), sticky="w")
        ctk.CTkLabel(bar, text=str(num), width=26, height=26, corner_radius=13,
                     fg_color=PRIMARY, text_color="white",
                     font=ctk.CTkFont(size=13, weight="bold")).pack(side="left")
        title_lbl = ctk.CTkLabel(bar, text=title,
                                 font=ctk.CTkFont(size=14, weight="bold"))
        title_lbl.pack(side="left", padx=(8, 0))
        return card, title_lbl

    def _build_step_input(self, parent):
        card, _ = self._step_card(parent, 0, 1, "选择要拆的表格")

        btns = ctk.CTkFrame(card, fg_color="transparent")
        btns.grid(row=1, column=0, padx=12, pady=(6, 2), sticky="w")
        ctk.CTkButton(btns, text="📄 选一个 Excel 文件", width=170,
                      command=self._pick_file).pack(side="left")
        ctk.CTkButton(btns, text="📁 选整个文件夹（批量拆）", width=190,
                      command=self._pick_folder).pack(side="left", padx=(10, 0))

        self._input_var = ctk.StringVar(value=self.cfg.get("input_path", ""))
        entry = ctk.CTkEntry(card, textvariable=self._input_var)
        entry.grid(row=2, column=0, padx=12, pady=(6, 2), sticky="ew")
        entry.bind("<Return>",   lambda e: self._on_path_edited())
        entry.bind("<FocusOut>", lambda e: self._on_path_edited())
        # 拖拽：整个卡片和路径框都接收拖入的文件/文件夹
        self._enable_drop(card, self._on_drop_input)
        self._enable_drop(entry, self._on_drop_input)

        self._in_status = ctk.CTkLabel(card, text=self._idle_input_text(),
                                       font=ctk.CTkFont(size=11), text_color=C_MUTED,
                                       anchor="w", justify="left")
        self._in_status.grid(row=3, column=0, padx=12, pady=(0, 10), sticky="w")

    def _build_step_field(self, parent):
        card, _ = self._step_card(parent, 1, 2, "按哪个字段拆分")
        self._field_card = card   # 取值勾选面板的挂载点

        row = ctk.CTkFrame(card, fg_color="transparent")
        row.grid(row=1, column=0, padx=12, pady=(6, 2), sticky="w")
        ctk.CTkLabel(row, text="拆分字段").pack(side="left")
        self._split_var = ctk.StringVar(value=self.cfg.get("split_column") or COL_PLACEHOLDER)
        self._split_menu = ctk.CTkOptionMenu(row, variable=self._split_var,
                                             values=[self._split_var.get()], width=220,
                                             command=self._on_split_col_change)
        self._split_menu.pack(side="left", padx=(8, 8))
        _ghost_button(row, text="🔄 重新识别", width=104,
                      command=self._scan_input).pack(side="left")

        ctk.CTkLabel(card, text="拆分字段的每个取值各生成一个文件：比如按「部门」拆 → 销售部.xlsx、财务部.xlsx…",
                     font=ctk.CTkFont(size=11), text_color=C_MUTED).grid(
            row=2, column=0, padx=12, pady=(2, 0), sticky="w")
        self._status = ctk.CTkLabel(card, text="", font=ctk.CTkFont(size=11),
                                    text_color=C_MUTED, anchor="w")
        self._status.grid(row=3, column=0, padx=12, pady=(0, 6), sticky="w")

        # 取值预览（v2.7）：选定拆分字段后自动枚举该列取值——点开始之前就能看到
        # 「将拆成哪几组」，并可展开勾选只拆其中一部分（替代高级设置里手填文本的主路径）
        vrow = ctk.CTkFrame(card, fg_color="transparent")
        vrow.grid(row=4, column=0, padx=12, pady=(0, 4), sticky="w")
        self._values_status = ctk.CTkLabel(vrow, text="选定拆分字段后，这里会列出将拆出的分组",
                                           font=ctk.CTkFont(size=11), text_color=C_MUTED)
        self._values_status.pack(side="left")
        self._values_btn = _flat_button(vrow, text="▸ 查看/勾选分组", width=110,
                                        command=self._toggle_values)
        self._values_btn.pack(side="left", padx=(6, 0))
        self._values_btn.pack_forget()      # 枚举到取值后才出现
        # 展开中的勾选面板 grid 到 row=5（_toggle_values 控制）

        self._batch_frame = self._build_batch(card)
        self._batch_frame.grid(row=6, column=0, padx=12, pady=(2, 12), sticky="ew")
        self._batch_frame.grid_remove()   # 仅输入为文件夹时显示（_update_input_ui 控制）

    def _build_batch(self, parent):
        f = ctk.CTkFrame(parent)
        f.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(f, text="批量选项（拆整个文件夹时）",
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=("gray25", "gray75")).grid(row=0, column=0, padx=12, pady=(8, 0), sticky="w")

        self._zip_var = ctk.BooleanVar(value=self.cfg.get("make_zip", True))
        ctk.CTkCheckBox(f, text="每个分组打包成 ZIP（方便整体转发）",
                        variable=self._zip_var).grid(row=1, column=0, padx=12, pady=(8, 4), sticky="w")

        self._person_var = ctk.BooleanVar(value=self.cfg.get("to_person", False))
        ctk.CTkCheckBox(f, text="同时拆到人（每人再单独生成一个文件）",
                        variable=self._person_var,
                        command=self._update_person_state).grid(row=2, column=0, padx=12, pady=4, sticky="w")

        prow = ctk.CTkFrame(f, fg_color="transparent")
        prow.grid(row=3, column=0, padx=(38, 12), pady=(0, 2), sticky="w")
        ctk.CTkLabel(prow, text="按哪个字段区分人").pack(side="left")
        self._pcol_var = ctk.StringVar(value=self.cfg.get("person_column") or COL_PLACEHOLDER)
        self._pcol_menu = ctk.CTkOptionMenu(prow, variable=self._pcol_var,
                                            values=[self._pcol_var.get()], width=150)
        self._pcol_menu.pack(side="left", padx=(6, 10))
        ctk.CTkLabel(prow, text="只处理文件名含").pack(side="left")
        self._pfilter_var = ctk.StringVar(value=", ".join(self.cfg.get("person_file_filter", [])))
        self._pfilter_entry = ctk.CTkEntry(prow, textvariable=self._pfilter_var, width=150,
                                           placeholder_text="如 工资,费用")
        self._pfilter_entry.pack(side="left", padx=6)
        ctk.CTkLabel(prow, text="的表（留空＝全部）").pack(side="left")

        ctk.CTkLabel(f, text="「到人」只对上面指定的表生效，其余文件只进汇总。",
                     font=ctk.CTkFont(size=11), text_color=C_MUTED).grid(
            row=4, column=0, padx=(38, 12), pady=(0, 8), sticky="w")
        self._update_person_state()
        return f

    def _update_person_state(self):
        """「同时拆到人」未勾选时置灰它的子选项，减少一次性摆在面前的决策数。"""
        state = "normal" if self._person_var.get() else "disabled"
        self._pcol_menu.configure(state=state)
        self._pfilter_entry.configure(state=state)

    # ── 拖拽（v2.7，tkinterdnd2 可选依赖）────────────────────
    def _enable_drop(self, widget, handler):
        """给控件注册文件拖放；tkinterdnd2 缺失时静默跳过（不影响其余功能）。"""
        if not _DND_OK:
            return
        try:
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind("<<Drop>>", handler)
        except Exception:
            pass

    @staticmethod
    def _parse_drop_paths(data):
        """解析 tkdnd 的拖放数据：Windows 下格式为 `{带空格的路径} 普通路径` 混合。"""
        return [m.group(1) or m.group(2)
                for m in re.finditer(r"\{([^}]*)\}|(\S+)", data or "")]

    def _idle_input_text(self):
        """①卡片未选输入时的提示语；支持拖拽时优先教拖拽（零学习成本的操作）。"""
        if _DND_OK:
            return "还没有选择文件（可点上方按钮，或直接把文件 / 文件夹拖进来）"
        return "还没有选择文件（也可以把路径粘贴到上面的输入框）"

    def _on_drop_input(self, event):
        """Excel 模式拖放：接受单个 Excel 文件或文件夹（多个文件时取第一个所在项）。"""
        paths = self._parse_drop_paths(getattr(event, "data", ""))
        target = next((p for p in paths
                       if os.path.isdir(p) or p.lower().endswith((".xlsx", ".xls"))), None)
        if target is None:
            self._set_in_status("⚠ 请拖入 Excel 文件（.xlsx / .xls）或装表格的文件夹", C_WARN)
            return
        self._input_var.set(os.path.normpath(target))
        self._after_pick()

    def _on_drop_pdfs(self, event):
        """PDF 模式拖放：接受一批 PDF。"""
        paths = [p for p in self._parse_drop_paths(getattr(event, "data", ""))
                 if p.lower().endswith(".pdf")]
        if not paths:
            self._pdf_in_status.configure(text="⚠ 请拖入 PDF 文件", text_color=C_WARN)
            return
        self._pdf_paths = [os.path.normpath(p) for p in paths]
        self._update_pdf_in_status()
        self._suggest_pdf_output()
        self._update_pdf_preview()

    def _on_drop_mapping(self, event):
        """映射清单拖放：接受单个 xlsx。"""
        paths = self._parse_drop_paths(getattr(event, "data", ""))
        target = next((p for p in paths if p.lower().endswith(".xlsx")), None)
        if target is None:
            self._pdf_map_status.configure(text="⚠ 映射清单请拖入 Excel 文件（.xlsx）",
                                           text_color=C_WARN)
            return
        self._pdf_map_var.set(os.path.normpath(target))
        self._scan_mapping()

    # ── 取值预览（v2.7，A1）──────────────────────────────
    def _on_split_col_change(self, _choice):
        """拆分字段变更：防抖 300ms 后重新枚举该列取值（连续切换不重复读文件）。"""
        self._enum_token += 1
        if self._enum_after is not None:
            self.after_cancel(self._enum_after)
        self._enum_after = self.after(300, self._enum_values)

    def _enum_values(self):
        """子线程枚举当前拆分字段的全部取值（复用扫描时识别的模板文件）。"""
        self._enum_after = None
        col = self._split_var.get()
        self._values_all, self._value_vars = [], {}
        self._refresh_values_ui()
        if (col in ("", COL_PLACEHOLDER) or not self._tpl_path
                or not os.path.exists(self._tpl_path)):
            return
        token = self._enum_token
        tpl, cfg = self._tpl_path, self._collect_config()
        self._values_status.configure(text="正在枚举该字段的取值…", text_color=C_MUTED)

        def work():
            try:
                from core.splitter import list_values
                vals = list_values(tpl, cfg, col)
            except Exception as e:
                vals = None
                # 不静默吞掉：失败原因进日志，界面只显示笼统提示（与 _scan_input 同理）
                self._ui_q.put(("log", f"⚠ 枚举「{col}」的取值失败：{e}"))
            self._ui_q.put(("values", (col, token, vals)))

        threading.Thread(target=work, daemon=True).start()

    def _on_values(self, payload):
        """枚举结果回填：过期（已换字段/又发起新枚举）结果直接丢弃。"""
        col, token, vals = payload
        if token != self._enum_token or col != self._split_var.get():
            return
        if vals is None:
            self._values_status.configure(
                text="⚠ 没能枚举取值（不影响拆分，可直接开始）", text_color=C_WARN)
            return
        self._values_all = vals
        # 默认全选；高级设置里手填过「只拆这些取值」且能对上号时，按它预选
        wanted = set(self._split_list(self._selvals_var.get()))
        use_filter = bool(wanted) and any(w in vals for w in wanted)
        self._value_vars = {}
        for v in vals:
            var = ctk.BooleanVar(value=(v in wanted) if use_filter else True)
            var.trace_add("write", lambda *_: self._sync_selvals())
            self._value_vars[v] = var
        self._refresh_values_ui()
        if self._values_panel is not None:      # 面板开着就重建内容
            self._close_values_panel()
            self._toggle_values()

    def _refresh_values_ui(self):
        """按 _values_all 更新取值摘要行：几组、示例、异常多的告警。"""
        n = len(self._values_all)
        if not n:
            self._values_btn.pack_forget()
            self._close_values_panel()
            if self._split_var.get() not in ("", COL_PLACEHOLDER) and self._tpl_path:
                self._values_status.configure(
                    text="选定拆分字段后，这里会列出将拆出的分组", text_color=C_MUTED)
            return
        shown = "、".join(self._values_all[:6]) + ("…" if n > 6 else "")
        if n > 50:
            self._values_status.configure(
                text=f"⚠ 将拆出 {n} 组：{shown}——取值异常多，若这是工号 / 姓名类字段，建议换个分组字段",
                text_color=C_WARN)
        elif n == 1:
            self._values_status.configure(
                text=f"只会拆出 1 组「{self._values_all[0]}」——确认字段没选错？",
                text_color=C_WARN)
        else:
            self._values_status.configure(
                text=f"将拆成 {n} 组：{shown}", text_color=C_OK)
        self._values_btn.pack(side="left", padx=(6, 0))

    def _toggle_values(self):
        """展开 / 收起取值勾选面板。"""
        if self._values_panel is not None:
            self._close_values_panel()
            return
        self._values_open = True
        self._values_btn.configure(text="▾ 收起分组")
        panel = ctk.CTkFrame(self._field_card)
        panel.grid(row=5, column=0, padx=12, pady=(0, 8), sticky="ew")
        panel.grid_columnconfigure(0, weight=1)
        bar = ctk.CTkFrame(panel, fg_color="transparent")
        bar.grid(row=0, column=0, padx=8, pady=(6, 0), sticky="w")
        _flat_button(bar, text="全选", width=46,
                     command=lambda: self._set_all_values(True)).pack(side="left")
        _flat_button(bar, text="清空", width=46,
                     command=lambda: self._set_all_values(False)).pack(side="left")
        ctk.CTkLabel(bar, text="只拆勾选的分组；全选＝拆分全部",
                     font=ctk.CTkFont(size=11), text_color=C_MUTED).pack(side="left", padx=8)
        box = ctk.CTkScrollableFrame(panel, height=140)
        box.grid(row=1, column=0, padx=8, pady=(2, 8), sticky="ew")
        for v, var in self._value_vars.items():
            ctk.CTkCheckBox(box, text=v, variable=var).pack(anchor="w")
        self._values_panel = panel

    def _close_values_panel(self):
        if self._values_panel is not None:
            self._values_panel.destroy()
            self._values_panel = None
        self._values_open = False
        self._values_btn.configure(text="▸ 查看/勾选分组")

    def _set_all_values(self, on):
        for var in self._value_vars.values():
            var.set(on)

    def _sync_selvals(self):
        """勾选状态 → selected_values（高级设置文本框同步显示，双向一致）。

        全选或全不选都等价于「拆分全部」，保持文本框为空——与「留空=自动枚举
        全部取值」的既有语义一致，也避免无谓地存一长串取值。
        """
        if not self._value_vars:
            return
        checked = [v for v, var in self._value_vars.items() if var.get()]
        if not checked or len(checked) == len(self._value_vars):
            self._selvals_var.set("")
        else:
            self._selvals_var.set(", ".join(checked))

    def _build_adv_area(self, parent):
        self._adv_btn = _flat_button(parent, text="▸ 高级设置（一般用不到）", width=200,
                                     anchor="w", command=self._toggle_adv)
        self._adv_btn.grid(row=2, column=0, padx=8, pady=(0, 4), sticky="w")
        self._adv_frame = self._build_advanced(parent)   # 展开时 grid 到 row=3

    def _build_advanced(self, parent):
        f = ctk.CTkFrame(parent)
        f.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(f, text="以下选项保持默认即可；只有扫描不到字段、或需要归并取值时才需要调整。",
                     font=ctk.CTkFont(size=11), text_color=C_MUTED).grid(
            row=0, column=0, columnspan=2, padx=12, pady=(8, 2), sticky="w")

        ctk.CTkLabel(f, text="表头识别").grid(row=1, column=0, padx=12, pady=8, sticky="w")
        mode_map = {"auto": "自动", "row": "指定行号", "keyword": "关键词"}
        self._header_seg = ctk.CTkSegmentedButton(f, values=["自动", "指定行号", "关键词"],
                                                  command=self._on_header_mode)
        self._header_seg.set(mode_map.get(self.cfg.get("header_mode", "auto"), "自动"))
        self._header_seg.grid(row=1, column=1, padx=8, pady=8, sticky="w")

        self._hrow_frame = ctk.CTkFrame(f, fg_color="transparent")
        self._hrow_frame.grid(row=2, column=0, columnspan=2, padx=12, pady=2, sticky="w")
        ctk.CTkLabel(self._hrow_frame, text="表头行号（1 基）：").grid(row=0, column=0)
        self._hrow_var = ctk.StringVar(value=str(self.cfg.get("header_row", 1)))
        ctk.CTkEntry(self._hrow_frame, textvariable=self._hrow_var, width=80).grid(row=0, column=1, padx=6)

        self._kw_frame = ctk.CTkFrame(f, fg_color="transparent")
        self._kw_frame.grid(row=3, column=0, columnspan=2, padx=12, pady=2, sticky="ew")
        self._kw_frame.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(self._kw_frame, text="必含关键词A（逗号分隔）").grid(row=0, column=0, padx=4, pady=4, sticky="w")
        self._grid_keys_var = ctk.StringVar(value=", ".join(self.cfg.get("grid_keys", [])))
        ctk.CTkEntry(self._kw_frame, textvariable=self._grid_keys_var).grid(row=0, column=1, padx=6, pady=4, sticky="ew")
        ctk.CTkLabel(self._kw_frame, text="必含关键词B（逗号分隔）").grid(row=1, column=0, padx=4, pady=4, sticky="w")
        self._id_keys_var = ctk.StringVar(value=", ".join(self.cfg.get("id_keys", [])))
        ctk.CTkEntry(self._kw_frame, textvariable=self._id_keys_var).grid(row=1, column=1, padx=6, pady=4, sticky="ew")

        ctk.CTkLabel(f, text="跳过值（逗号分隔）").grid(row=4, column=0, padx=12, pady=8, sticky="w")
        self._skip_var = ctk.StringVar(value=", ".join(self.cfg.get("skip_values", [])))
        ctk.CTkEntry(f, textvariable=self._skip_var).grid(row=4, column=1, padx=8, pady=8, sticky="ew")

        ctk.CTkLabel(f, text="只拆这些取值（逗号分隔，留空=全部）").grid(row=5, column=0, padx=12, pady=8, sticky="w")
        self._selvals_var = ctk.StringVar(value=", ".join(self.cfg.get("selected_values", [])))
        ctk.CTkEntry(f, textvariable=self._selvals_var).grid(row=5, column=1, padx=8, pady=8, sticky="ew")

        ctk.CTkLabel(f, text="取值归并映射（JSON，选填）").grid(row=6, column=0, padx=12, pady=(8, 2), sticky="nw")
        self._alias_box = ctk.CTkTextbox(f, height=64, wrap="none")
        self._alias_box.grid(row=6, column=1, padx=8, pady=(8, 2), sticky="ew")
        alias = self.cfg.get("value_alias_map", {})
        if alias:
            self._alias_box.insert("1.0", json.dumps(alias, ensure_ascii=False, indent=2))

        opts = ctk.CTkFrame(f, fg_color="transparent")
        opts.grid(row=7, column=0, columnspan=2, padx=8, pady=8, sticky="w")
        self._exact_var    = ctk.BooleanVar(value=self.cfg.get("exact_match", True))
        self._merge_var    = ctk.BooleanVar(value=self.cfg.get("merge_across_files", True))
        self._preserve_var = ctk.BooleanVar(value=self.cfg.get("preserve_format", True))
        self._keep_formula_var = ctk.BooleanVar(value=self.cfg.get("keep_formulas", False))
        self._auto_open_var = ctk.BooleanVar(value=self.cfg.get("auto_open_output", True))
        ctk.CTkCheckBox(opts, text="精确匹配", variable=self._exact_var).grid(row=0, column=0, padx=8, pady=4)
        ctk.CTkCheckBox(opts, text="跨文件合并汇总", variable=self._merge_var).grid(row=0, column=1, padx=8, pady=4)
        ctk.CTkCheckBox(opts, text="保留格式", variable=self._preserve_var).grid(row=0, column=2, padx=8, pady=4)
        ctk.CTkCheckBox(opts, text="完成后打开输出", variable=self._auto_open_var).grid(row=0, column=3, padx=8, pady=4)
        ctk.CTkCheckBox(opts, text="保留公式（收件人可见计算过程，仅同行公式，需先勾选保留格式）",
                        variable=self._keep_formula_var).grid(
            row=1, column=0, columnspan=4, padx=8, pady=(0, 4), sticky="w")
        ctk.CTkLabel(opts, text="提示：「跨文件合并汇总」会把所有结果留在内存里最后统一写盘，宽表 + 大批量时明显更吃内存；默认的「按原表各自拆分」更稳。",
                     font=ctk.CTkFont(size=11), text_color=C_MUTED,
                     wraplength=560, justify="left").grid(
            row=2, column=0, columnspan=4, padx=8, pady=(0, 6), sticky="w")

        self._on_header_mode(self._header_seg.get())
        return f

    # ── PDF 加密分发模式的步骤卡 ─────────────────────────
    def _build_pdf_step_input(self, parent):
        card, _ = self._step_card(parent, 0, 1, "选择要分发的 PDF")

        btns = ctk.CTkFrame(card, fg_color="transparent")
        btns.grid(row=1, column=0, padx=12, pady=(6, 2), sticky="w")
        ctk.CTkButton(btns, text="📄 选 PDF 文件（可多选）", width=190,
                      command=self._pick_pdfs).pack(side="left")
        _ghost_button(btns, text="清空", width=60,
                      command=self._clear_pdfs).pack(side="left", padx=(10, 0))
        self._enable_drop(card, self._on_drop_pdfs)

        self._pdf_in_status = ctk.CTkLabel(card, text="", font=ctk.CTkFont(size=11),
                                           text_color=C_MUTED, anchor="w", justify="left")
        self._pdf_in_status.grid(row=2, column=0, padx=12, pady=(0, 10), sticky="w")
        self._update_pdf_in_status()

    def _build_pdf_step_map(self, parent):
        card, _ = self._step_card(parent, 1, 2, "选择密码映射清单（Excel）")

        prow = ctk.CTkFrame(card, fg_color="transparent")
        prow.grid(row=1, column=0, padx=12, pady=(6, 2), sticky="ew")
        prow.grid_columnconfigure(1, weight=1)
        ctk.CTkButton(prow, text="📋 选映射清单", width=120,
                      command=self._pick_mapping).grid(row=0, column=0)
        self._pdf_map_var = ctk.StringVar(value=self.cfg.get("pdf_mapping_path", ""))
        entry = ctk.CTkEntry(prow, textvariable=self._pdf_map_var)
        entry.grid(row=0, column=1, padx=8, sticky="ew")
        entry.bind("<Return>",   lambda e: self._on_map_edited())
        entry.bind("<FocusOut>", lambda e: self._on_map_edited())
        self._enable_drop(entry, self._on_drop_mapping)
        _ghost_button(prow, text="生成模板", width=90,
                      command=self._make_pdf_template).grid(row=0, column=2)

        ctk.CTkLabel(card, text="清单里每行一个网格：网格名、专属密码，可选接收人；没有清单可点「生成模板」，没有密码列开始时可选自动生成随机密码",
                     font=ctk.CTkFont(size=11), text_color=C_MUTED).grid(
            row=2, column=0, padx=12, pady=(2, 0), sticky="w")

        crow = ctk.CTkFrame(card, fg_color="transparent")
        crow.grid(row=3, column=0, padx=12, pady=(6, 2), sticky="w")
        ctk.CTkLabel(crow, text="网格列").pack(side="left")
        self._pdfgrid_var = ctk.StringVar(value=self.cfg.get("pdf_grid_column") or PDF_COL_PLACEHOLDER)
        self._pdfgrid_menu = ctk.CTkOptionMenu(crow, variable=self._pdfgrid_var,
                                               values=[self._pdfgrid_var.get()], width=140)
        self._pdfgrid_menu.pack(side="left", padx=(6, 12))
        ctk.CTkLabel(crow, text="密码列").pack(side="left")
        self._pdfpwd_var = ctk.StringVar(value=self.cfg.get("pdf_password_column") or PDF_COL_PLACEHOLDER)
        self._pdfpwd_menu = ctk.CTkOptionMenu(crow, variable=self._pdfpwd_var,
                                              values=[self._pdfpwd_var.get()], width=140)
        self._pdfpwd_menu.pack(side="left", padx=(6, 12))
        ctk.CTkLabel(crow, text="接收人列").pack(side="left")
        self._pdfrecv_var = ctk.StringVar(value=self.cfg.get("pdf_receiver_column") or PDF_RECV_NONE)
        self._pdfrecv_menu = ctk.CTkOptionMenu(crow, variable=self._pdfrecv_var,
                                               values=[self._pdfrecv_var.get()], width=140)
        self._pdfrecv_menu.pack(side="left", padx=(6, 0))

        self._pdf_map_status = ctk.CTkLabel(card, text="", font=ctk.CTkFont(size=11),
                                            text_color=C_MUTED, anchor="w")
        self._pdf_map_status.grid(row=4, column=0, padx=12, pady=(0, 6), sticky="w")

        wm = ctk.CTkFrame(card)
        wm.grid(row=5, column=0, padx=12, pady=(2, 12), sticky="ew")
        self._pdfwm_var = ctk.BooleanVar(value=self.cfg.get("pdf_watermark", True))
        ctk.CTkCheckBox(wm, text="加网格专属水印（泄露可溯源）",
                        variable=self._pdfwm_var,
                        command=self._update_wm_state).grid(row=0, column=0, padx=12, pady=(8, 4), sticky="w")
        wrow = ctk.CTkFrame(wm, fg_color="transparent")
        wrow.grid(row=1, column=0, padx=(38, 12), pady=(0, 8), sticky="w")
        ctk.CTkLabel(wrow, text="水印文字").pack(side="left")
        self._pdfwmtext_var = ctk.StringVar(value=self.cfg.get("pdf_watermark_text", "{grid} {date}"))
        self._pdfwmtext_entry = ctk.CTkEntry(wrow, textvariable=self._pdfwmtext_var, width=220)
        self._pdfwmtext_entry.pack(side="left", padx=(6, 8))
        ctk.CTkLabel(wrow, text="{grid}＝网格名  {date}＝日期",
                     font=ctk.CTkFont(size=11), text_color=C_MUTED).pack(side="left")
        self._update_wm_state()

    def _update_wm_state(self):
        self._pdfwmtext_entry.configure(state="normal" if self._pdfwm_var.get() else "disabled")

    def _update_pdf_in_status(self):
        n = len(self._pdf_paths)
        if not n:
            self._pdf_in_status.configure(text="还没有选择 PDF；每个网格都会拿到全部所选文件（各自密码+水印）",
                                          text_color=C_MUTED)
            return
        names = [os.path.basename(p) for p in self._pdf_paths]
        shown = "、".join(names[:3]) + ("…" if n > 3 else "")
        self._pdf_in_status.configure(text=f"✓ 已选 {n} 个：{shown}", text_color=C_OK)
        self._update_pdf_preview()   # PDF 个数变化会改变「N 网格 × M PDF」预览

    def _pick_pdfs(self):
        paths = filedialog.askopenfilenames(filetypes=[("PDF 文件", "*.pdf")])
        if paths:
            self._pdf_paths = [os.path.normpath(p) for p in paths]
            self._update_pdf_in_status()
            self._suggest_pdf_output()

    def _clear_pdfs(self):
        self._pdf_paths = []
        self._update_pdf_in_status()

    def _suggest_pdf_output(self):
        """PDF 模式的默认输出：第一个 PDF 所在目录下「分发结果」。规则同 _suggest_output。"""
        cur = self._output_var.get().strip()
        if cur and cur != self._out_auto:
            return
        if self._pdf_paths:
            auto = os.path.join(os.path.dirname(self._pdf_paths[0]), "分发结果")
            self._output_var.set(auto)
            self._out_auto = auto

    def _pick_mapping(self):
        path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx")])
        if path:
            self._pdf_map_var.set(os.path.normpath(path))
            self._scan_mapping()

    def _on_map_edited(self):
        p = self._pdf_map_var.get().strip()
        if p == self._scanned_map_path:
            return
        if p and os.path.exists(p):
            self._scan_mapping()
        elif p:
            self._scanned_map_path = p
            self._pdf_map_status.configure(text="⚠ 找不到这个文件，请检查路径", text_color=C_WARN)

    def _scan_mapping(self):
        """子线程读映射清单表头，经 UI 泵回填三个列下拉（机制同 _scan_input）。"""
        p = self._pdf_map_var.get().strip()
        self._scanned_map_path = p
        if not p or not os.path.exists(p):
            return
        self._pdf_map_status.configure(text="正在读取清单…", text_color=C_MUTED)

        def work():
            cols, n_rows = [], 0
            try:
                from core.pdf_dist import list_mapping_columns
                cols = list_mapping_columns(p)
                # 顺带数网格行数（分发预览「N 个网格 × M 个 PDF」用）；清单是小表，直接全读
                from openpyxl import load_workbook
                wb = load_workbook(p, read_only=True)
                try:
                    n_rows = max(0, (wb.worksheets[0].max_row or 1) - 1)
                finally:
                    wb.close()
            except Exception:
                cols, n_rows = [], 0
            self._ui_q.put(("pdf_scan", (cols, p, n_rows)))

        threading.Thread(target=work, daemon=True).start()

    def _on_pdf_scan(self, payload):
        cols, p, n_rows = payload
        if p != self._pdf_map_var.get().strip():
            return    # 结果已过期
        self._pdf_n_cols = len(cols)
        self._pdf_grid_rows = n_rows
        if not cols:
            self._pdf_map_status.configure(
                text="⚠ 没读到表头：请确认清单第 1 行是列名（如 网格 / 密码 / 接收人）", text_color=C_WARN)
            return
        self._pdfgrid_menu.configure(values=cols)
        if self._pdfgrid_var.get() not in cols:
            self._pdfgrid_var.set(self._recommend_col(cols, ("网格", "部门", "区域", "单位", "分组", "组")) or cols[0])
        self._pdfpwd_menu.configure(values=cols)
        if self._pdfpwd_var.get() not in cols:
            self._pdfpwd_var.set(self._recommend_col(cols, ("密码", "口令", "pwd", "密")) or cols[-1])
        recv_values = [PDF_RECV_NONE] + cols
        self._pdfrecv_menu.configure(values=recv_values)
        # 「（不需要）」是未选状态的占位，也参与推荐；用户真不需要可再手动选回去
        if self._pdfrecv_var.get() not in cols:
            self._pdfrecv_var.set(self._recommend_col(cols, ("接收", "姓名", "负责人", "联系")) or PDF_RECV_NONE)
        self._update_pdf_preview()

    def _update_pdf_preview(self):
        """分发预览（C3）：选完清单和 PDF 后，直接告诉用户将产出多少个加密副本。"""
        g, m, cols_n = self._pdf_grid_rows, len(self._pdf_paths), self._pdf_n_cols
        if not cols_n or not g:
            return
        extra = f"；{g} 个网格 × {m} 个 PDF = {g * m} 个加密副本" if m else ""
        self._pdf_map_status.configure(
            text=f"✓ 识别到 {cols_n} 列 / {g} 个网格{extra}，确认三个下拉选的对不对",
            text_color=C_OK)

    def _make_pdf_template(self):
        """一键生成「网格 → 密码」映射清单模板并自动载入（C1：解决没有清单的第一步卡点）。"""
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile="映射清单模板.xlsx",
            filetypes=[("Excel 文件", "*.xlsx")])
        if not path:
            return
        try:
            from core.pdf_dist import write_mapping_template
            write_mapping_template(path)
        except Exception as e:
            messagebox.showerror("生成失败", f"生成模板失败：{e}")
            return
        self._pdf_map_var.set(os.path.normpath(path))
        self._scan_mapping()

    @staticmethod
    def _recommend_col(cols, keywords):
        for kw in keywords:
            for c in cols:
                if kw.lower() in str(c).lower():
                    return c
        return None

    # ── 模式切换 ─────────────────────────────────────────
    def _on_mode_change(self, label):
        self._mode = "pdf" if label == "PDF 加密分发" else "excel"
        self._apply_mode()

    def _apply_mode(self):
        """按当前模式显示对应步骤区并更新③卡片文案；运行中禁止切换按钮已够复杂，
        简单起见运行中也允许切界面，但开始按钮在运行中本来就是禁用态。"""
        if self._mode == "pdf":
            self._excel_frame.grid_remove()
            self._pdf_frame.grid()
            self._action_title.configure(text="开始分发")
            if not self._running:
                self._start_btn.configure(text="🔐 开始分发")
                self._run_status.configure(
                    text="完成 ① ② 后点「开始分发」；每个网格生成专属密码+水印的副本，并输出分发清单",
                    text_color=C_MUTED)
        else:
            self._pdf_frame.grid_remove()
            self._excel_frame.grid()
            self._action_title.configure(text="开始拆分")
            if not self._running:
                self._start_btn.configure(text="▶ 开始拆分")
                self._run_status.configure(
                    text="完成 ① ② 后直接点「开始拆分」；保存位置不用改，结果会放进自动创建的「拆分结果」文件夹",
                    text_color=C_MUTED)

    def _build_action(self, parent):
        card, self._action_title = self._step_card(parent, 2, 3, "开始拆分", padx=20, pady=(6, 0))

        orow = ctk.CTkFrame(card, fg_color="transparent")
        orow.grid(row=1, column=0, padx=12, pady=(4, 2), sticky="ew")
        orow.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(orow, text="结果保存到").grid(row=0, column=0, sticky="w")
        self._output_var = ctk.StringVar(value=self.cfg.get("output_path", ""))
        ctk.CTkEntry(orow, textvariable=self._output_var).grid(row=0, column=1, padx=8, sticky="ew")
        _ghost_button(orow, text="浏览", width=60,
                      command=self._browse_output).grid(row=0, column=2)

        btns = ctk.CTkFrame(card, fg_color="transparent")
        btns.grid(row=2, column=0, padx=12, pady=(8, 2), sticky="w")
        # 主按钮是全页唯一核心动作（v2.7.1）：加高配粗 + 品牌主色，保证第一眼锁定
        self._start_btn = ctk.CTkButton(btns, text="▶ 开始拆分", width=220, height=44,
                                        corner_radius=10,
                                        font=ctk.CTkFont(size=15, weight="bold"),
                                        fg_color=PRIMARY, hover_color=PRIMARY_HOVER,
                                        command=self._start)
        self._start_btn.grid(row=0, column=0)
        self._stop_btn = ctk.CTkButton(btns, text="⏹ 停止", width=90,
                                       fg_color="gray40", hover_color="gray30",
                                       command=self._stop)
        self._stop_btn.grid(row=0, column=1, padx=(10, 0))
        self._stop_btn.grid_remove()          # 只在运行时出现
        self._open_btn = _ghost_button(btns, text="📂 打开输出文件夹", width=140,
                                       command=self._open_output)
        self._open_btn.grid(row=0, column=2, padx=(10, 0))
        self._open_btn.grid_remove()          # 成功后出现

        self._progress = ctk.CTkProgressBar(card, corner_radius=6,
                                           progress_color=PRIMARY)
        self._progress.grid(row=3, column=0, padx=12, pady=(8, 2), sticky="ew")
        self._progress.set(0)
        self._run_status = ctk.CTkLabel(card, text="完成 ① ② 后直接点「开始拆分」；保存位置不用改，结果会放进自动创建的「拆分结果」文件夹",
                                        font=ctk.CTkFont(size=11), text_color=C_MUTED, anchor="w")
        self._run_status.grid(row=4, column=0, padx=12, pady=(0, 2), sticky="w")
        # 完成摘要（v2.7，A4；v2.7.1 横幅化）：浅绿底圆角横幅承载「完成瞬间」的
        # 明确成功反馈——摘要内容与逻辑不变，只加视觉承载。
        self._summary_banner = ctk.CTkFrame(card, corner_radius=10,
                                            fg_color=OK_BANNER_BG)
        self._summary_banner.grid(row=5, column=0, padx=12, pady=(0, 10), sticky="ew")
        self._summary_banner.grid_remove()
        self._summary_lbl = ctk.CTkLabel(self._summary_banner, text="",
                                         font=ctk.CTkFont(size=12, weight="bold"),
                                         text_color=C_OK, anchor="w", justify="left")
        self._summary_lbl.grid(row=0, column=0, padx=14, pady=8, sticky="w")

    def _build_bottom(self, parent):
        bar = ctk.CTkFrame(parent, fg_color="transparent")
        bar.grid(row=3, column=0, padx=20, pady=(4, 0), sticky="ew")
        bar.grid_columnconfigure(1, weight=1)
        self._log_btn = _flat_button(bar, text="▸ 处理详情", width=100, anchor="w",
                                     command=self._toggle_log)
        self._log_btn.grid(row=0, column=0, sticky="w")
        self._copylog_btn = _flat_button(bar, text="📋 复制日志", width=90, anchor="w",
                                         command=self._copy_log)
        self._copylog_btn.grid(row=0, column=1, sticky="w", padx=(4, 0))
        self._copylog_btn.grid_remove()      # 只在失败时出现（配合反馈问卷降低求助门槛）
        right = ctk.CTkFrame(bar, fg_color="transparent")
        right.grid(row=0, column=2, sticky="e")
        _flat_button(right, text="❓ 常见问题", width=90,
                     command=lambda: webbrowser.open(FAQ_URL)).pack(side="left")
        _flat_button(right, text="💬 反馈建议", width=90,
                     command=self._open_feedback).pack(side="left")
        self._appearance_btn = _flat_button(right, text=self._appearance_btn_text(),
                                            width=96, command=self._cycle_appearance)
        self._appearance_btn.pack(side="left", padx=(6, 0))
        _flat_button(right, text="🔍 大字号" if self.cfg.get("ui_scale", 100) == 100 else "🔍 标准字号",
                     width=90, command=self._toggle_scale).pack(side="left", padx=(6, 0))
        _flat_button(right, text="保存配置", width=80,
                     command=self._save_cfg).pack(side="left", padx=(6, 0))

        self._log_box = ctk.CTkTextbox(parent, state="disabled", wrap="none", height=190)
        # 展开时 grid 到 row=4（_toggle_log 控制），默认收起

        # 「数据不出本地」是核心卖点，字号/颜色比作者署名这类次要信息更突出
        footer = ctk.CTkFrame(parent, fg_color="transparent")
        footer.grid(row=5, column=0, padx=22, pady=(2, 8), sticky="w")
        ctk.CTkLabel(footer, text="🔒 数据仅在本机处理，不上传任何服务器",
                     font=ctk.CTkFont(size=12, weight="bold"), text_color=C_OK).pack(anchor="w")
        ctk.CTkLabel(footer, text="作者 AbeLin · MIT 开源",
                     font=ctk.CTkFont(size=10), text_color="gray").pack(anchor="w", pady=(1, 0))

    # ── 自适应 / 折叠 ───────────────────────────────────
    def _update_input_ui(self):
        """输入类型跟随路径自动推断：是文件夹才显示批量选项。"""
        if os.path.isdir(self._input_var.get().strip()):
            self._batch_frame.grid()
        else:
            self._batch_frame.grid_remove()

    def _toggle_adv(self):
        self._adv_open = not self._adv_open
        if self._adv_open:
            self._adv_frame.grid(row=3, column=0, padx=8, pady=(0, 10), sticky="ew")
            self._adv_btn.configure(text="▾ 高级设置（一般用不到）")
        else:
            self._adv_frame.grid_remove()
            self._adv_btn.configure(text="▸ 高级设置（一般用不到）")

    def _log_btn_text(self):
        """日志折叠钮文案：本轮有 ⚠ 提醒时带角标，把藏在折叠日志里的警告提到可见。"""
        arrow = "▾" if self._log_open else "▸"
        badge = f"（⚠ {self._warn_count}）" if self._warn_count else ""
        return f"{arrow} 处理详情{badge}"

    def _toggle_log(self, show=None):
        show = (not self._log_open) if show is None else show
        if show == self._log_open:
            return
        self._log_open = show
        if show:
            self._log_box.grid(row=4, column=0, padx=20, pady=(2, 0), sticky="ew")
        else:
            self._log_box.grid_remove()
        self._log_btn.configure(text=self._log_btn_text())

    def _on_header_mode(self, label):
        if label == "指定行号":
            self._hrow_frame.grid()
            self._kw_frame.grid_remove()
        elif label == "关键词":
            self._hrow_frame.grid_remove()
            self._kw_frame.grid()
        else:
            self._hrow_frame.grid_remove()
            self._kw_frame.grid_remove()

    # ── 选择输入 / 输出 ─────────────────────────────────
    def _pick_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if path:
            self._input_var.set(os.path.normpath(path))
            self._after_pick()

    def _pick_folder(self):
        path = filedialog.askdirectory()
        if path:
            self._input_var.set(os.path.normpath(path))
            self._after_pick()

    def _after_pick(self):
        self._update_input_ui()
        self._suggest_output()
        self._scan_input()

    def _on_path_edited(self):
        """手动改路径框（回车/失焦）后同步界面；路径没变就什么都不做。"""
        p = self._input_var.get().strip()
        if p == self._scanned_path:
            return
        self._update_input_ui()
        if p and os.path.exists(p):
            self._suggest_output()
            self._scan_input()
        elif p:
            self._scanned_path = p
            self._set_in_status("⚠ 找不到这个路径，请检查有没有写错", C_WARN)

    def _browse_output(self):
        path = filedialog.askdirectory()
        if path:
            self._output_var.set(os.path.normpath(path))

    @staticmethod
    def _auto_output_for(p):
        """输入路径对应的默认输出目录；无法推导时返回空串。

        文件 → 同目录下「拆分结果」；文件夹 → 文件夹里面的「拆分结果」——结果永远出现在
        用户数据旁边，找得到、不污染上层目录。放输入里面是安全的：run_split 递归扫描时
        会整体跳过 output_root 子树（重复运行也不会把旧结果吃回去）。
        """
        if os.path.isdir(p):
            return os.path.join(os.path.normpath(p), "拆分结果")
        if os.path.isfile(p):
            return os.path.join(os.path.dirname(p), "拆分结果")
        return ""

    def _suggest_output(self):
        """按输入自动推荐输出目录；用户手动定过（当前值≠上次推荐值）就不再覆盖。"""
        cur = self._output_var.get().strip()
        if cur and cur != self._out_auto:
            return
        auto = self._auto_output_for(self._input_var.get().strip())
        if auto:
            self._output_var.set(auto)
            self._out_auto = auto

    # ── 扫描输入（文件数 + 字段），子线程执行 ───────────
    def _scan_input(self):
        p = self._input_var.get().strip()
        self._scanned_path = p
        if not p or not os.path.exists(p):
            self._set_in_status(self._idle_input_text() if not p
                                else "⚠ 找不到这个路径，请检查有没有写错",
                                C_MUTED if not p else C_WARN)
            return
        is_dir = os.path.isdir(p)
        self._set_in_status("正在读取…", C_MUTED)
        self._set_scan_status("正在识别表格里的字段…", C_MUTED)
        cfg = self._collect_config()
        out_now = self._output_var.get().strip()

        def work():
            # 计数镜像 run_split 的规则：输出目录子树整体跳过（结果放在输入里时数字才对）
            out_abs = os.path.normcase(os.path.abspath(out_now)) if out_now else None
            tpl, n, total_bytes = None, 0, 0
            if is_dir:
                for root, dirs, files in os.walk(p):
                    if out_abs:
                        ra = os.path.normcase(os.path.abspath(root))
                        if ra == out_abs or ra.startswith(out_abs + os.sep):
                            dirs[:] = []
                            continue
                    for fn in sorted(files):
                        if fn.lower().endswith(('.xlsx', '.xls')) and not fn.startswith('~$'):
                            n += 1
                            try:
                                total_bytes += os.path.getsize(os.path.join(root, fn))
                            except OSError:
                                pass
                            if tpl is None:
                                tpl = os.path.join(root, fn)
            else:
                tpl, n = p, 1
                try:
                    total_bytes = os.path.getsize(p)
                except OSError:
                    pass
            cols = []
            if tpl:
                try:
                    from core.splitter import list_columns
                    cols = list_columns(tpl, cfg)
                except Exception as e:
                    # 不静默吞掉：扫描失败原因（文件被占用/损坏等）打进日志，
                    # 否则界面只会显示笼统的「未能识别字段」，用户无从判断到底是
                    # 真的表头识别不出来，还是文件被 Excel/网盘同步占用等其它原因。
                    cols = []
                    self._ui_q.put(("log", f"⚠ 扫描「{os.path.basename(tpl)}」失败：{e}"))
            self._ui_q.put(("scan", (cols, n, os.path.basename(tpl) if tpl else "",
                                     is_dir, p, tpl or "", total_bytes)))

        threading.Thread(target=work, daemon=True).start()

    def _on_scan(self, payload):
        cols, n, tpl_name, is_dir, p, tpl_full, total_bytes = payload
        if p != self._input_var.get().strip():
            return    # 结果已过期（用户又换了输入）
        self._tpl_path = tpl_full or None    # 取值枚举复用这份模板文件
        # 换了输入，旧的取值预览随之作废
        self._values_all, self._value_vars = [], {}
        self._refresh_values_ui()
        mb = total_bytes / 1048576
        size_txt = f"，共 {mb:.0f} MB" if mb >= 1 else ""
        heavy = "（量较大，预计需要几分钟）" if (n > 50 or mb > 200) else ""
        name = os.path.basename(os.path.normpath(p)) or p
        if is_dir and n == 0:
            self._set_in_status(f"⚠ 文件夹「{name}」里没有找到 Excel 文件（.xlsx / .xls）", C_WARN)
            self._set_scan_status("", C_MUTED)
            return
        if is_dir:
            self._set_in_status(f"✓ 文件夹「{name}」：找到 {n} 个 Excel{size_txt}，全部一起拆{heavy}", C_OK)
        else:
            self._set_in_status(f"✓ 已选择：{name}{size_txt}{heavy}", C_OK)
        self._on_columns(cols, tpl_name)

    def _on_columns(self, cols, tpl_name=None):
        self._columns = cols
        if cols:
            self._split_menu.configure(values=cols)
            if self._split_var.get() not in cols:
                self._split_var.set(self._recommend_split(cols))
            self._pcol_menu.configure(values=cols)
            if self._pcol_var.get() not in cols:
                self._pcol_var.set(self._recommend_person(cols))
            src = f"（来自「{tpl_name}」）" if tpl_name else ""
            self._set_scan_status(f"✓ 识别到 {len(cols)} 个字段{src}，确认下拉框选的对不对", C_OK)
            self._on_split_col_change(None)   # 字段已定 → 自动枚举取值（v2.7 取值预览）
        else:
            self._warn_scan_failed()

    def _warn_scan_failed(self):
        """扫描不到字段时的提示：「关键词」模式下关键词没匹配上是个常见且好定位的
        原因（比如换了一批表头写法不同的表格，还留着上次填的旧关键词），单独
        点名出来，比统一甩一句「未能识别字段」让用户少走弯路。"""
        if self._header_seg.get() == "关键词":
            kws = [k for k in (self._grid_keys_var.get().strip(), self._id_keys_var.get().strip()) if k]
            if kws:
                if not self._adv_open:
                    self._toggle_adv()
                self._set_scan_status(
                    f"未能识别字段：当前「表头识别」用的是「关键词」模式，但表格前几行里"
                    f"没找到关键词「{' / '.join(kws)}」——已展开下方「高级设置」，"
                    f"核对这两个关键词是否适用于这份表格，或改回「自动」识别", C_WARN)
                return
        self._set_scan_status("未能识别字段：点「🔄 重新识别」重试，或到「高级设置」手动指定表头行", C_WARN)

    @staticmethod
    def _recommend_split(cols):
        """猜一个最像「分组」的字段做默认值，猜不到就用第一个（用户仍需自己确认）。"""
        for kw in ("部门", "网格", "区域", "分公司", "门店", "班组", "科室", "组织", "单位", "类别", "分类"):
            for c in cols:
                if kw in str(c):
                    return c
        return cols[0]

    @staticmethod
    def _recommend_person(cols):
        for kw in ("姓名", "名字", "工号", "负责人", "经办"):
            for c in cols:
                if kw in str(c):
                    return c
        return cols[0]

    def _find_stale_results(self, input_dir, output_root):
        """输入文件夹里「不被本次输出目录覆盖」的历史拆分结果目录。

        结果默认放在输入里面、由 run_split 按 output_root 跳过；但如果用户后来把输出
        改到了别处，留在输入里的旧结果就会被当成数据拆一遍（静默重复）。开跑前找出来确认。
        """
        out = os.path.normcase(os.path.abspath(output_root))
        hits = []
        for root, dirs, _ in os.walk(input_dir):
            for d in list(dirs):
                full = os.path.normcase(os.path.abspath(os.path.join(root, d)))
                if full == out or full.startswith(out + os.sep):
                    dirs.remove(d)    # 本次输出覆盖的子树，run_split 会跳过，不算
                    continue
                if d == "拆分结果" or d.endswith("_拆分结果") or re.fullmatch(r"\d{8}结果", d):
                    hits.append(os.path.join(root, d))
                    dirs.remove(d)    # 命中即整树算一个，不再往里翻
        return hits

    @staticmethod
    def _find_open_locks(path, cap=4):
        """检测输入里正被 Excel/WPS 打开的文件（存在 ~$同名 锁文件）。

        返回原文件名列表（最多 cap 个）；目录扫描限前 2000 个目录，防止巨型目录卡界面。
        """
        hits = []
        if os.path.isfile(path):
            if os.path.exists(os.path.join(os.path.dirname(path),
                                           "~$" + os.path.basename(path))):
                hits.append(os.path.basename(path))
            return hits
        n_walk = 0
        for _root, _dirs, files in os.walk(path):
            n_walk += 1
            if n_walk > 2000 or len(hits) >= cap:
                break
            names = {f for f in files if f.lower().endswith(('.xlsx', '.xls'))}
            locked = {f[2:] for f in names if f.startswith('~$')}
            for f in sorted(locked & names):
                hits.append(f)
                if len(hits) >= cap:
                    break
        return hits

    def _set_in_status(self, text, color=C_MUTED):
        self._in_status.configure(text=text, text_color=color)

    def _set_scan_status(self, text, color=C_MUTED):
        self._status.configure(text=text, text_color=color)

    # ── 配置收集 ───────────────────────────────────────
    def _split_list(self, text):
        return [k.strip() for k in text.split(",") if k.strip()]

    def _parse_alias(self):
        raw = self._alias_box.get("1.0", "end").strip()
        if not raw:
            return {}, True
        try:
            obj = json.loads(raw)
            if isinstance(obj, dict):
                return obj, True
        except Exception:
            pass
        return {}, False

    def _collect_config(self):
        mode_map = {"自动": "auto", "指定行号": "row", "关键词": "keyword"}
        try:
            header_row = int(self._hrow_var.get())
        except (TypeError, ValueError):
            header_row = 1

        split_col = "" if self._split_var.get() == COL_PLACEHOLDER else self._split_var.get()
        pcol = "" if self._pcol_var.get() == COL_PLACEHOLDER else self._pcol_var.get()
        alias, _ = self._parse_alias()

        recv = self._pdfrecv_var.get()
        return {
            "input_path":  self._input_var.get().strip(),
            "output_path": self._output_var.get().strip(),
            "ui_mode":     self._mode,
            "pdf_input_paths": list(self._pdf_paths),
            "pdf_mapping_path": self._pdf_map_var.get().strip(),
            "pdf_grid_column": "" if self._pdfgrid_var.get() == PDF_COL_PLACEHOLDER else self._pdfgrid_var.get(),
            "pdf_password_column": "" if self._pdfpwd_var.get() == PDF_COL_PLACEHOLDER else self._pdfpwd_var.get(),
            "pdf_receiver_column": "" if recv in (PDF_RECV_NONE, PDF_COL_PLACEHOLDER) else recv,
            "pdf_watermark": self._pdfwm_var.get(),
            "pdf_watermark_text": self._pdfwmtext_var.get().strip() or "{grid} {date}",
            "pdf_watermark_opacity": self.cfg.get("pdf_watermark_opacity", 0.15),
            "pdf_watermark_angle": self.cfg.get("pdf_watermark_angle", 45),
            # 界面个性化键随每次保存透传，避免跑一次成功后丢失（与 pdf 水印参数同理）
            "ui_scale": self.cfg.get("ui_scale", 100),
            "window_geometry": self.cfg.get("window_geometry", ""),
            "appearance_mode": self.cfg.get("appearance_mode", "system"),
            "header_mode": mode_map.get(self._header_seg.get(), "auto"),
            "header_row":  header_row,
            "grid_keys":   self._split_list(self._grid_keys_var.get()),
            "id_keys":     self._split_list(self._id_keys_var.get()),
            "split_column": split_col,
            "selected_values": self._split_list(self._selvals_var.get()),
            "person_column": pcol,
            "to_person":   self._person_var.get(),
            "person_file_filter": self._split_list(self._pfilter_var.get()),
            "value_alias_map": alias,
            "skip_values": self._split_list(self._skip_var.get()),
            "merge_across_files": self._merge_var.get(),
            "make_zip":    self._zip_var.get(),
            "exact_match": self._exact_var.get(),
            "preserve_format": self._preserve_var.get(),
            "keep_formulas": self._keep_formula_var.get(),
            "auto_open_output": self._auto_open_var.get(),
        }

    def _save_cfg(self):
        try:
            save_config(self._collect_config())
            messagebox.showinfo("已保存", f"配置已保存到：\n{USER_CONFIG_PATH}\n下次启动自动加载。\n（每次成功拆分后也会自动记住当前配置）")
        except Exception as e:
            messagebox.showerror("保存失败", f"无法保存配置：{e}")

    # ── 运行 ───────────────────────────────────────────
    def _log(self, msg):
        self._log_box.configure(state="normal")
        self._log_box.insert("end", msg + "\n")
        self._log_box.see("end")
        self._log_box.configure(state="disabled")

    def _stop_indeterminate(self):
        """进度条从不定态动画切回真实进度模式。"""
        if self._prog_indeterminate:
            self._prog_indeterminate = False
            self._progress.stop()
            self._progress.configure(mode="determinate")

    def _pump_ui(self):
        """主线程 UI 泵：每 100ms 批量消费子线程消息（日志/进度/扫描/完成）。

        子线程绝不直接碰 Tk 控件；日志批量合并插入，进度只取最新值——
        避免高频 after(0) 回调打满主线程事件队列（大文件时界面假死的诱因之一）。
        """
        logs, progress = [], None
        done: Any = _MISSING
        scan: Any = _MISSING
        pdf_scan: Any = _MISSING
        values: Any = _MISSING
        try:
            while True:
                kind, payload = self._ui_q.get_nowait()
                if kind == "log":
                    logs.append(payload)
                elif kind == "progress":
                    progress = payload
                elif kind == "done":
                    done = payload
                elif kind == "scan":
                    scan = payload
                elif kind == "pdf_scan":
                    pdf_scan = payload
                elif kind == "values":
                    values = payload
        except queue.Empty:
            pass
        if logs:
            # [SUMMARY] 机器可读行不直接展示：提取出来供完成摘要渲染（v2.7）；
            # 提取失败也不报错，退回原有纯日志行为
            out_lines = []
            for payload in logs:
                for line in str(payload).splitlines():
                    if line.startswith("[SUMMARY] "):
                        try:
                            self._last_summary = json.loads(line[len("[SUMMARY] "):])
                        except Exception:
                            pass
                        continue
                    if "⚠" in line:
                        self._warn_count += 1
                    out_lines.append(line)
            if out_lines:
                joined = "\n".join(out_lines)
                self._log(joined)
                if self._running:
                    # 状态行镜像最新一条日志：不展开「处理详情」也能看到进展
                    last = next((ln.strip() for ln in reversed(joined.splitlines()) if ln.strip()), "")
                    if last:
                        self._run_status.configure(text=("⏳ " + last)[:70], text_color=C_MUTED)
        if progress is not None:
            self._stop_indeterminate()
            self._progress.set(progress)
        if scan is not _MISSING:
            self._on_scan(scan)
        if pdf_scan is not _MISSING:
            self._on_pdf_scan(pdf_scan)
        if values is not _MISSING:
            self._on_values(values)
        if done is not _MISSING:
            self._on_done(*done)
        self.after(100, self._pump_ui)

    def _enter_running(self, busy_text, first_log):
        """进入运行态的公共 UI 准备：两种模式共用（按钮/进度/日志/状态行）。"""
        self._stop_flag = False
        self._running = True
        self._last_output = None
        self._warn_count = 0
        self._last_summary = None
        self._summary_lbl.configure(text="")
        self._summary_banner.grid_remove()
        self._copylog_btn.grid_remove()
        self._log_btn.configure(text=self._log_btn_text())
        self._start_btn.configure(state="disabled", text=busy_text)
        self._stop_btn.grid()
        self._stop_btn.configure(state="normal")
        self._open_btn.grid_remove()
        self._log_box.configure(state="normal")
        self._log_box.delete("1.0", "end")
        self._log_box.configure(state="disabled")
        # 点击的瞬间就要看到反馈：先用不定态动画，第一条真实进度回来后自动切换（见 _pump_ui）
        self._prog_indeterminate = True
        self._progress.configure(mode="indeterminate")
        self._progress.start()
        self._run_status.configure(text=first_log, text_color=C_MUTED)
        self._log("▶ " + first_log)

    def _start_pdf(self):
        cfg = self._collect_config()
        if not cfg["pdf_input_paths"]:
            messagebox.showwarning("先选择 PDF", "请先在第 ① 步选择要分发的 PDF 文件。")
            return
        missing = [p for p in cfg["pdf_input_paths"] if not os.path.exists(p)]
        if missing:
            messagebox.showwarning("PDF 不存在", "以下文件找不到了，请重新选择：\n" + "\n".join(missing[:3]))
            return
        if not cfg["pdf_mapping_path"] or not os.path.exists(cfg["pdf_mapping_path"]):
            messagebox.showwarning("先选择映射清单", "请在第 ② 步选择「网格 → 密码」的 Excel 映射清单。")
            return
        if not cfg["pdf_grid_column"]:
            messagebox.showwarning("先选择网格列", "请在第 ② 步选好「网格列」。\n如果下拉框没有内容，重新选一次映射清单。")
            return
        if not cfg["pdf_password_column"]:
            # 没选密码列 → 提供自动生成随机密码（C2），另存新清单不改动原文件
            if not messagebox.askyesno("自动生成密码？",
                                       "没有选择密码列。要我为每个网格自动生成随机密码，\n"
                                       "另存一份「含密码」清单并继续分发吗？\n（8 位随机密码只存在清单里，丢失无法找回）"):
                return
            try:
                from core.pdf_dist import fill_random_passwords
                out_path, npwd = fill_random_passwords(cfg["pdf_mapping_path"],
                                                       cfg["pdf_grid_column"])
            except Exception as e:
                messagebox.showerror("生成失败", f"自动生成密码失败：{e}")
                return
            cfg["pdf_mapping_path"] = out_path
            cfg["pdf_password_column"] = "密码"
            self._pdf_map_var.set(out_path)
            self._pdfpwd_var.set("密码")
            self._pdf_map_status.configure(
                text=f"🔑 已为 {npwd} 个网格生成随机密码 → {os.path.basename(out_path)}（含明文，保管好）",
                text_color=C_OK)
        if not cfg["output_path"]:
            self._suggest_pdf_output()
            cfg["output_path"] = self._output_var.get().strip()
        if not cfg["output_path"]:
            messagebox.showwarning("先选择保存位置", "请在第 ③ 步选择结果保存到哪个文件夹。")
            return
        try:
            os.makedirs(cfg["output_path"], exist_ok=True)
        except OSError as e:
            messagebox.showerror("无法创建保存位置",
                                 f"没法在这里创建结果文件夹：\n{cfg['output_path']}\n\n{e}\n\n"
                                 "请点「浏览」换一个能写入的位置（比如桌面或文档）。")
            return

        self._enter_running("⏳ 正在分发…", "已开始，正在读取映射清单…")

        def run():
            try:
                from core.pdf_dist import run_pdf_dist
                output_path = run_pdf_dist(
                    cfg,
                    log_fn=lambda m: self._ui_q.put(("log", m)),
                    progress_fn=lambda v: self._ui_q.put(("progress", v)),
                    stop_flag=lambda: self._stop_flag,
                )
            except Exception as e:
                self._ui_q.put(("log", f"\n❌ 运行出错：{e}"))
                output_path = None
            self._ui_q.put(("done", (cfg, output_path)))

        threading.Thread(target=run, daemon=True).start()

    def _start(self):
        if self._mode == "pdf":
            return self._start_pdf()
        cfg = self._collect_config()
        if not cfg["input_path"] or not os.path.exists(cfg["input_path"]):
            messagebox.showwarning("先选择表格", "请先在第 ① 步选择要拆的 Excel 文件或文件夹。")
            return
        if not cfg["split_column"]:
            messagebox.showwarning("先选择拆分字段",
                                   "请在第 ② 步选择「拆分字段」——按哪一列拆分。\n如果下拉框还没有内容，点「🔄 扫描字段」。")
            return
        if not cfg["output_path"]:
            self._suggest_output()
            cfg["output_path"] = self._output_var.get().strip()
        if not cfg["output_path"]:
            messagebox.showwarning("先选择保存位置", "请在第 ③ 步选择结果保存到哪个文件夹。")
            return
        # 输出目录若等于输入文件夹本身或在它上层，核心扫描时会把全部文件当输出跳过 → 提前拦住
        if os.path.isdir(cfg["input_path"]):
            try:
                inp = os.path.normcase(os.path.abspath(cfg["input_path"]))
                out = os.path.normcase(os.path.abspath(cfg["output_path"]))
                if inp == out or inp.startswith(out + os.sep):
                    messagebox.showwarning("换个保存位置",
                                           "结果保存位置不能选输入文件夹自己或它的上层文件夹，\n否则找不到要拆的文件。建议直接用自动推荐的位置。")
                    return
            except Exception:
                pass
            # 输入里有历史拆分结果、而本次输出没盖住它 → 会被当数据重复拆，先问一声
            try:
                stale = self._find_stale_results(cfg["input_path"], cfg["output_path"])
            except Exception:
                stale = []
            if stale:
                shown = "\n".join(stale[:3]) + ("\n…" if len(stale) > 3 else "")
                if not messagebox.askyesno(
                        "发现以前的拆分结果",
                        f"输入文件夹里有以前拆出来的结果：\n{shown}\n\n"
                        "这次运行会把里面的文件也当成数据一起拆，内容可能重复。\n"
                        "建议先把它们移走/删除，或把保存位置改回默认。\n\n仍要继续吗？"):
                    return
        # 误选保护（A1）：取值预览发现分组异常多时，开跑前再确认一次
        if len(self._values_all) > 50:
            if not messagebox.askyesno(
                    "分组特别多",
                    f"将按「{cfg['split_column']}」拆出 {len(self._values_all)} 个分组（文件）。\n"
                    "取值异常多常见于选错了字段（比如选到工号 / 姓名列）。\n\n仍要继续吗？"):
                return
        # 占用预检（B2）：输入里被 Excel/WPS 打开的文件（~$ 锁），读到的是未保存的旧内容
        locks = self._find_open_locks(cfg["input_path"])
        if locks:
            shown = "、".join(locks[:3]) + (" 等" if len(locks) > 3 else "")
            if not messagebox.askyesno(
                    "文件正被打开",
                    f"检测到以下文件正被 Excel/WPS 打开：{shown}\n"
                    "拆分读到的是「打开前保存的内容」，刚改的数不会生效。\n\n建议先关闭这些文件。仍要继续吗？"):
                return
        # 保存位置提前建好：位置只读/无权限时现在就报，而不是跑到一半失败
        try:
            os.makedirs(cfg["output_path"], exist_ok=True)
        except OSError as e:
            messagebox.showerror("无法创建保存位置",
                                 f"没法在这里创建结果文件夹：\n{cfg['output_path']}\n\n{e}\n\n"
                                 "请点「浏览」换一个能写入的位置（比如桌面或文档）。")
            return
        _, alias_ok = self._parse_alias()
        if not alias_ok:
            messagebox.showwarning("高级设置有误", "「高级设置 → 取值归并映射」不是合法 JSON，请修正或清空。")
            return

        self._enter_running("⏳ 正在拆分…", "已开始，正在扫描输入…")

        def run():
            try:
                from core.splitter import run_split
                output_path = run_split(
                    cfg,
                    log_fn=lambda m: self._ui_q.put(("log", m)),
                    progress_fn=lambda v: self._ui_q.put(("progress", v)),
                    stop_flag=lambda: self._stop_flag,
                )
            except Exception as e:
                self._ui_q.put(("log", f"\n❌ 运行出错：{e}"))
                output_path = None
            self._ui_q.put(("done", (cfg, output_path)))

        threading.Thread(target=run, daemon=True).start()

    def _stop(self):
        self._stop_flag = True
        self._log("正在停止，请等待当前文件处理完成...")

    def _copy_log(self):
        """把日志全文复制进剪贴板，方便粘贴到反馈问卷或发给维护者。"""
        try:
            self.clipboard_clear()
            self.clipboard_append(self._log_box.get("1.0", "end").strip())
            self._run_status.configure(text="📋 日志已复制，可直接粘贴到反馈问卷", text_color=C_OK)
        except Exception:
            pass

    def _appearance_btn_text(self):
        return {"system": "🌓 跟随系统", "light": "☀️ 浅色", "dark": "🌙 深色"}.get(
            self.cfg.get("appearance_mode", "system"), "🌓 跟随系统")

    def _cycle_appearance(self):
        """深浅色循环切换：系统 → 浅色 → 深色 → 系统；即时生效并持久化。"""
        order = ["system", "light", "dark"]
        cur = self.cfg.get("appearance_mode", "system")
        new = order[(order.index(cur) + 1) % len(order)]
        self.cfg["appearance_mode"] = new
        ctk.set_appearance_mode(new if new != "system" else "System")
        self._appearance_btn.configure(text=self._appearance_btn_text())
        try:
            save_config(self._collect_config())
        except Exception as e:
            messagebox.showerror("保存失败", f"无法保存外观设置：{e}")

    def _toggle_scale(self):
        """大字号切换：写入配置，下次启动生效（现有控件已按原缩放创建，不即时重排）。"""
        new = 100 if self.cfg.get("ui_scale", 100) != 100 else 115
        self.cfg["ui_scale"] = new
        try:
            save_config(self._collect_config() | {"ui_scale": new})
            messagebox.showinfo("字号设置", "字号设置已保存，下次启动程序时生效。")
        except Exception as e:
            messagebox.showerror("保存失败", f"无法保存设置：{e}")

    def _open_feedback(self):
        """打开匿名反馈页；顺手把版本号复制进剪贴板，方便用户粘贴到问卷。"""
        try:
            self.clipboard_clear()
            self.clipboard_append(f"ExcelRouter v{APP_VERSION}")
        except Exception:
            pass
        webbrowser.open(FEEDBACK_URL)
        self._log("💬 已打开反馈页，版本号已复制到剪贴板，粘贴到问卷即可。")

    def _open_output(self):
        path = self._last_output or self._output_var.get().strip()
        if path and os.path.exists(path):
            try:
                subprocess.Popen(f'explorer "{os.path.normpath(path)}"')
            except Exception:
                pass

    def _render_summary(self, output_path, is_pdf):
        """完成摘要（A4）：用 core 收尾的 [SUMMARY] 渲染多行结果概览。

        摘要行是文本契约（[SUMMARY] + JSON），解析失败/缺字段时逐字段容错，
        哪项拿不到就不显示哪项，绝不影响完成提示本身。
        """
        s = self._last_summary or {}
        lines, has_warn = [], False
        if is_pdf:
            if s.get("grids_ok") is not None:
                t = f"📊 {s['grids_ok']} 个网格成功"
                if s.get("grids_fail"):
                    t += f"，{s['grids_fail']} 个失败（详见处理详情）"
                    has_warn = True
                if s.get("files"):
                    t += f"，共 {s['files']} 个加密副本"
                lines.append(t)
                lines.append("🔒 分发清单含明文密码，只留给自己用，勿随文件一起发出去")
        else:
            if s.get("groups") is not None:
                t = f"📊 拆出 {s['groups']} 个分组 · {s.get('files', '?')} 个文件 · {s.get('rows', '?')} 行"
                if s.get("zips"):
                    t += f" · {s['zips']} 个 ZIP"
                lines.append(t)
                bad = [f"{s[k]} 个{lab}" for k, lab in
                       (("skipped_sheets", "sheet 被跳过"), ("failed_files", "文件读取失败"),
                        ("failed_saves", "文件保存失败")) if s.get(k)]
                if bad:
                    lines.append("⚠ " + "，".join(bad) + "（点「处理详情」看原因）")
                    has_warn = True
        if self._warn_count:
            self._log_btn.configure(text=self._log_btn_text())
        lines.append(f"📁 {output_path}")
        self._summary_lbl.configure(text="\n".join(lines))
        # 有跳过/失败时横幅退回中性底色，避免「绿底＋橙字」的违和
        self._summary_banner.configure(
            fg_color=("gray90", "gray20") if has_warn else OK_BANNER_BG)
        self._summary_lbl.configure(text_color=C_WARN if has_warn else C_OK)
        self._summary_banner.grid()

    def _on_done(self, cfg, output_path):
        self._running = False
        is_pdf = cfg.get("ui_mode") == "pdf"
        self._start_btn.configure(state="normal",
                                  text="🔐 开始分发" if self._mode == "pdf" else "▶ 开始拆分")
        self._stop_btn.grid_remove()
        self._stop_indeterminate()
        self._progress.set(1 if output_path else 0)
        if output_path:
            self._last_output = output_path
            self._open_btn.grid()
            if self._stop_flag:
                self._run_status.configure(text="⏹ 已停止：处理完的部分已保存，可点「打开输出文件夹」查看",
                                           text_color=C_WARN)
            else:
                self._run_status.configure(text="✅ 分发完成！结果已保存" if is_pdf else "✅ 拆分完成！结果已保存",
                                           text_color=C_OK)
                self._log("💬 用得顺手或踩了坑？点「反馈建议」匿名告诉作者（1 分钟）。")
            self._render_summary(output_path, is_pdf)
            try:
                save_config(cfg)      # 静默记住本次配置：下次打开即用（失败运行不存，避免存坏参数）
            except Exception:
                pass
        else:
            if self._warn_count:
                self._log_btn.configure(text=self._log_btn_text())
            self._run_status.configure(
                text="❌ 没有完成：原因见下方「处理详情」；可点「复制日志」连同反馈一起发给作者",
                text_color=C_ERR)
            self._copylog_btn.grid()
            self._toggle_log(show=True)
        if cfg.get("auto_open_output") and output_path and os.path.exists(output_path):
            try:
                subprocess.Popen(f'explorer "{os.path.normpath(output_path)}"')
            except Exception:
                pass


def run():
    app = App()
    app.mainloop()
