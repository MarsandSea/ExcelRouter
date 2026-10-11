"""三个 er_*.py CLI 的冒烟测试：以子进程方式跑（就是 Claude 实际会怎么调），
解析 stdout 最后一行 JSON。样本构造沿用 excel-router 上游 tests/ 的写法
（表头故意不放第一行、带样式、含合计行），确保 CLI 包装层没有偷懒漏传参数。
"""
import json
import os
import subprocess
import sys

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

# 本文件位于 <skill 目录>/tests/ 下：向上两级就是 skill 目录本身（scripts/ 与它同级）。
# 测试随 skill 一起发布上架，所以路径必须按「包内位置」推导，不能依赖仓库根目录。
SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(SKILL_DIR, "scripts")

HEADERS = ["工号", "姓名", "部门", "城市", "金额"]


def _make_book(path, rows, title="月度报表"):
    """表头在第 2 行、第 1 行是大标题——专测自动表头识别不是碰巧读对了第一行。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = title
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for i, r in enumerate(rows, start=3):
        for c, v in enumerate(r, 1):
            ws.cell(row=i, column=c, value=v)
    wb.save(path)


def _run(script, *args):
    """跑一个 er_*.py，返回 (returncode, json_obj, stderr_text)。"""
    proc = subprocess.run(
        [sys.executable, os.path.join(SCRIPTS, script), *args],
        capture_output=True, text=True, encoding="utf-8",
    )
    stdout_lines = [line for line in proc.stdout.splitlines() if line.strip()]
    assert stdout_lines, f"stdout 应至少有一行 JSON，实际为空。stderr:\n{proc.stderr}"
    obj = json.loads(stdout_lines[-1])
    return proc.returncode, obj, proc.stderr


@pytest.fixture
def two_books(tmp_path):
    inp = tmp_path / "in"
    out = tmp_path / "out"
    inp.mkdir()
    out.mkdir()
    _make_book(inp / "A.xlsx", [
        ["001", "张三", "销售部", "北京", 100],
        ["002", "李四", "销售部", "北京", 200],
        ["003", "王五", "销售部", "上海", 150],
        ["004", "赵六", "技术部", "北京", 300],
        ["", "合计", "", "", 750],
    ])
    _make_book(inp / "B.xlsx", [
        ["005", "钱七", "销售部", "广州", 120],
        ["006", "孙八", "技术部", "深圳", 220],
    ])
    return str(inp), str(out)


# ---------- er_inspect.py ----------

def test_inspect_single_file_detects_header_and_columns(tmp_path):
    f = tmp_path / "a.xlsx"
    _make_book(f, [["001", "张三", "销售部", "北京", 100]])
    rc, obj, _ = _run("er_inspect.py", "--input", str(f))
    assert rc == 0
    assert obj["ok"] is True
    assert obj["is_dir"] is False
    assert obj["excel_count"] == 1
    assert obj["header_row"] == 2
    assert obj["columns"] == HEADERS


def test_inspect_lists_values_and_skips_total_row(tmp_path):
    f = tmp_path / "a.xlsx"
    _make_book(f, [
        ["001", "张三", "销售部", "北京", 100],
        ["002", "李四", "技术部", "北京", 200],
        ["", "合计", "", "", 300],
    ])
    rc, obj, _ = _run("er_inspect.py", "--input", str(f), "--column", "部门")
    assert rc == 0
    assert obj["values"] == ["技术部", "销售部"]


def test_inspect_directory_counts_files(two_books):
    inp, _out = two_books
    rc, obj, _ = _run("er_inspect.py", "--input", inp)
    assert rc == 0
    assert obj["is_dir"] is True
    assert obj["excel_count"] == 2


def test_inspect_directory_values_are_flagged_as_sample(two_books):
    """目录输入默认只读「样本表」，取值数字本身没错，错的是把抽样结论说成全量结论
    —— 上游 v2.9.0 在 GUI 侧修的就是这个。这里 A.xlsx 只有北京/上海，
    广州/深圳在 B.xlsx 里，样本口径必然看不全，所以必须自报 values_scope。"""
    inp, _out = two_books
    rc, obj, _ = _run("er_inspect.py", "--input", inp, "--column", "城市")
    assert rc == 0
    assert obj["values_scope"] == "sample"
    assert obj["values_sample_file"] == "A.xlsx"      # 排序后固定是它，换次运行不该变
    assert "抽样" in obj["values_warning"]
    assert "广州" not in obj["values"]                # 正是抽样会漏掉的那个


def test_inspect_all_files_unions_values_across_books(two_books):
    """--all-files 读遍所有表取并集，并把口径标成 all、不再带警告。"""
    inp, _out = two_books
    rc, obj, _ = _run("er_inspect.py", "--input", inp, "--column", "城市", "--all-files")
    assert rc == 0
    assert obj["values_scope"] == "all"
    assert obj["values_from_files"] == 2
    assert set(obj["values"]) == {"北京", "上海", "广州", "深圳"}
    assert "values_warning" not in obj


def test_inspect_single_file_values_scope_is_single(tmp_path):
    """单文件没有抽样问题，不该吓唬用户报 sample。"""
    book = tmp_path / "one.xlsx"
    _make_book(str(book), [["001", "张三", "销售部", "北京", 100]])
    rc, obj, _ = _run("er_inspect.py", "--input", str(book), "--column", "城市")
    assert rc == 0
    assert obj["values_scope"] == "single"
    assert "values_warning" not in obj


def test_inspect_missing_input_reports_friendly_error(tmp_path):
    rc, obj, _ = _run("er_inspect.py", "--input", str(tmp_path / "not-exist.xlsx"))
    assert rc == 1
    assert obj["ok"] is False
    assert "找不到" in obj["error"]


# ---------- er_split.py ----------

def test_split_missing_by_reports_friendly_error(two_books):
    inp, out = two_books
    rc, obj, _ = _run("er_split.py", "--input", inp, "--output", out)
    assert rc == 1
    assert obj["ok"] is False
    assert "--by" in obj["error"]


def test_split_dry_run_does_not_write_files(two_books):
    inp, out = two_books
    rc, obj, _ = _run("er_split.py", "--input", inp, "--output", out, "--by", "部门", "--dry-run")
    assert rc == 0
    assert obj["dry_run"] is True
    assert obj["config"]["split_column"] == "部门"
    assert os.listdir(out) == []


def test_split_produces_expected_output_files(two_books):
    inp, out = two_books
    rc, obj, stderr = _run("er_split.py", "--input", inp, "--output", out, "--by", "部门")
    assert rc == 0, stderr
    assert obj["ok"] is True
    output_path = obj["output_path"]
    assert os.path.isdir(output_path)
    # 默认 merge_across_files=False：按原表各自拆分，两个源文件各自产出「部门」子目录
    produced = set(os.listdir(output_path))
    assert "运行日志.txt" in produced
    # A.xlsx 和 B.xlsx 都含「销售部」，都产出该主取值目录
    assert "销售部" in produced or "销售部.zip" in produced


def test_split_merge_flag_merges_across_files(two_books):
    inp, out = two_books
    rc, obj, stderr = _run("er_split.py", "--input", inp, "--output", out, "--by", "部门", "--merge", "--no-zip")
    assert rc == 0, stderr
    output_path = obj["output_path"]
    # merge=True 时同取值跨文件合并为扁平文件：{部门}.xlsx 直接在输出根
    assert os.path.exists(os.path.join(output_path, "销售部.xlsx"))


def test_split_selected_values_filters(two_books):
    inp, out = two_books
    rc, obj, stderr = _run("er_split.py", "--input", inp, "--output", out, "--by", "部门",
                            "--values", "技术部", "--merge", "--no-zip")
    assert rc == 0, stderr
    output_path = obj["output_path"]
    assert os.path.exists(os.path.join(output_path, "技术部.xlsx"))
    assert not os.path.exists(os.path.join(output_path, "销售部.xlsx"))


# ---------- 「像软件一样省事」的几条保证 ----------

def test_split_unknown_column_fails_instead_of_empty_success(two_books):
    """旧版：字段不存在时内核跳过所有 sheet，返回 ok:true + 0 个文件，AI 会说「拆完了」。"""
    inp, out = two_books
    rc, obj, _ = _run("er_split.py", "--input", inp, "--output", out, "--by", "科室")
    assert rc == 1 and obj["ok"] is False
    assert "科室" in obj["error"]
    assert "部门" in str(obj["available_columns"])
    assert os.listdir(out) == []          # 没开跑，不留空目录


def test_split_unmatched_values_fails(two_books):
    inp, out = two_books
    rc, obj, _ = _run("er_split.py", "--input", inp, "--output", out, "--by", "部门",
                      "--values", "不存在的部")
    assert rc == 1 and obj["ok"] is False
    assert "不存在的部" in obj["error"]


def test_split_result_is_self_contained(two_books):
    """一次调用就够交付：默认输出在数据旁边，JSON 自带摘要、文件清单和人话。"""
    inp, _ = two_books
    one = os.path.join(inp, "A.xlsx")
    rc, obj, stderr = _run("er_split.py", "--input", one, "--by", "部门")
    assert rc == 0, stderr
    assert obj["output_path"].startswith(os.path.join(inp, "拆分结果"))
    assert obj["summary"]["groups"] == 2 and obj["summary"]["rows"] == 4
    assert obj["files_total"] == 3          # 2 个分组文件 + 1 个总包
    assert all(os.path.isfile(f) for f in obj["files"])
    assert "2 组" in obj["message"] and obj["output_path"] in obj["message"]
    assert obj["opened"] is False         # 没加 --open 不弹窗


def test_split_partial_column_name_reports_resolved(tmp_path):
    p = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "所属部门", "部门编码"])
    ws.append(["张三", "销售部", "S01"])
    ws.append(["李四", "技术部", "T01"])
    wb.save(p)
    rc, obj, stderr = _run("er_split.py", "--input", str(p), "--by", "部门")
    assert rc == 0, stderr
    assert obj["split_column_resolved"] == "所属部门"
    assert any("部门编码" in w for w in obj["warnings"])   # 有歧义要说出来


def test_split_blocks_old_results_inside_input(two_books):
    """输入文件夹里留着以前的拆分结果时，直接拆会把它们当数据再拆一遍。"""
    inp, _ = two_books
    old = os.path.join(inp, "旧的", "10011200结果")
    os.makedirs(old)
    _make_book(os.path.join(old, "销售部.xlsx"), [["001", "张三", "销售部", "北京", 100]])
    rc, obj, _ = _run("er_split.py", "--input", inp, "--by", "部门")
    assert rc == 1 and obj["stale_dirs"] == [old]
    rc, obj, stderr = _run("er_split.py", "--input", inp, "--by", "部门", "--include-old-results")
    assert rc == 0, stderr


def test_split_makes_one_bundle_zip_listed_first(two_books):
    """「起码打包成一个 ZIP」：单文件拆分原来只有散文件。总包排第一，不含分组 ZIP（防体积翻倍）。"""
    import zipfile
    inp, _ = two_books
    rc, obj, stderr = _run("er_split.py", "--input", os.path.join(inp, "A.xlsx"), "--by", "部门")
    assert rc == 0, stderr
    bundle = obj["bundle_zip"]
    assert obj["files"][0] == bundle and os.path.basename(bundle).startswith("全部结果_按部门拆分")
    names = zipfile.ZipFile(bundle).namelist()
    assert sorted(names) == ["技术部.xlsx", "销售部.xlsx"]
    assert "压缩包" in obj["message"]


def test_split_no_zip_flag_skips_bundle(two_books):
    inp, _ = two_books
    rc, obj, _ = _run("er_split.py", "--input", os.path.join(inp, "A.xlsx"), "--by", "部门", "--no-zip")
    assert rc == 0 and obj["bundle_zip"] is None
    assert not any(f.endswith(".zip") for f in obj["files"])


def test_person_split_keeps_title_and_header_rows(two_books):
    """单文件按人拆：表头上面有大标题时，旧实现会①把列名行当成一个叫「姓名」的人拆出去，
    ②每个人的文件只有大标题、没有列名。真实工资表几乎都有大标题，这条必须守住。"""
    inp, _ = two_books
    rc, obj, stderr = _run("er_split.py", "--input", os.path.join(inp, "A.xlsx"),
                           "--by", "部门", "--to-person", "姓名")
    assert rc == 0, stderr
    person_dir = os.path.join(obj["output_path"], "销售部", "到人")
    people = sorted(os.listdir(person_dir))
    assert people == ["张三_销售部.xlsx", "李四_销售部.xlsx", "王五_销售部.xlsx"]
    ws = openpyxl.load_workbook(os.path.join(person_dir, "张三_销售部.xlsx")).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "月度报表"                       # 大标题
    assert list(rows[1]) == HEADERS                       # 列名
    assert ws.cell(2, 1).font.b                           # 列名格式也带过来了
    assert rows[2][1] == "张三" and len(rows) == 3
    # 每组 ZIP 也补上了（内核只给目录输入打分组 ZIP）
    assert os.path.exists(os.path.join(obj["output_path"], "销售部.zip"))


def test_split_hidden_content_default_and_keep(tmp_path):
    """v3.0：隐藏的 sheet 不发、隐藏列清空且仍隐藏（到人文件也一样）；--keep-hidden 全部照样发出。"""
    src = tmp_path / "表.xlsx"
    _make_book(src, [["001", "张三", "销售部", "北京", 100], ["002", "李四", "技术部", "上海", 200]])
    wb = openpyxl.load_workbook(src)
    wb.active.column_dimensions["D"].hidden = True                 # 城市
    ws2 = wb.create_sheet("底稿")
    ws2.append(["部门", "说明"])
    ws2.append(["销售部", "内部"])
    ws2.sheet_state = "hidden"
    wb.save(src)

    rc, obj, stderr = _run("er_split.py", "--input", str(src), "--by", "部门", "--to-person", "姓名",
                           "--output", str(tmp_path / "a"))
    assert rc == 0, stderr
    out = openpyxl.load_workbook(os.path.join(obj["output_path"], "销售部.xlsx"))
    assert out.sheetnames == ["Sheet1"]
    assert out.active["D3"].value is None and out.active.column_dimensions["D"].hidden
    person = openpyxl.load_workbook(os.path.join(obj["output_path"], "销售部", "到人", "张三_销售部.xlsx")).active
    assert person["D3"].value is None and person.column_dimensions["D"].hidden

    rc, obj, stderr = _run("er_split.py", "--input", str(src), "--by", "部门", "--keep-hidden",
                           "--output", str(tmp_path / "b"))
    assert rc == 0, stderr
    out = openpyxl.load_workbook(os.path.join(obj["output_path"], "销售部.xlsx"))
    assert out.sheetnames == ["Sheet1", "底稿"] and out.active["D3"].value == "北京"


def test_inspect_suggests_split_column_in_one_call(tmp_path):
    """桌面版②卡片的等价物：不用再逐个字段查，一次给出推荐 + 每列分组数。"""
    p = tmp_path / "t.xlsx"
    _make_book(p, [
        ["001", "张三", "销售部", "北京", 100],
        ["002", "李四", "技术部", "上海", 200],
        ["003", "王五", "销售部", "北京", 150],
    ])
    rc, obj, stderr = _run("er_inspect.py", "--input", str(p))
    assert rc == 0, stderr
    sug = obj["suggestion"]
    assert sug["split_column"] == "部门" and sug["values"] == ["技术部", "销售部"]
    assert sug["person_column"] == "姓名"
    assert obj["group_counts"]["城市"] == 2
    assert "推荐按「部门」拆" in obj["message"]
    assert "金额" not in obj["message"]   # 数字列不当候选


def test_inspect_columns_only_skips_profile(tmp_path):
    p = tmp_path / "t.xlsx"
    _make_book(p, [["001", "张三", "销售部", "北京", 100]])
    rc, obj, _ = _run("er_inspect.py", "--input", str(p), "--columns-only")
    assert rc == 0 and "suggestion" not in obj and obj["columns"]


# ---------- er_pdf_dist.py ----------

def _make_pdf(path, pages=1):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(False)
    for i in range(pages):
        pdf.add_page()
        pdf.set_font("Helvetica", size=14)
        pdf.text(30, 40, f"SOURCE page{i + 1}")
    pdf.output(str(path))


def _make_mapping(path, rows, headers=("GridName", "Pwd", "Receiver")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def test_pdf_dist_list_columns(tmp_path):
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [("GridA", "001234", "Zhang")])
    rc, obj, _ = _run("er_pdf_dist.py", "--mapping", str(mp), "--list-columns")
    assert rc == 0
    assert obj["columns"] == ["GridName", "Pwd", "Receiver"]


def test_pdf_dist_encrypts_and_writes_manifest(tmp_path):
    pdf_path = tmp_path / "report.pdf"
    _make_pdf(pdf_path)
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [("GridA", "001234", "Zhang"), ("GridB", 8888, "Li")])
    out = tmp_path / "dist"
    rc, obj, stderr = _run(
        "er_pdf_dist.py", "--pdf", str(pdf_path), "--mapping", str(mp),
        "--grid-col", "GridName", "--password-col", "Pwd", "--receiver-col", "Receiver",
        "--output", str(out), "--no-watermark",
    )
    assert rc == 0, stderr
    assert obj["ok"] is True
    assert os.path.exists(obj["manifest"])
    assert os.path.exists(os.path.join(obj["output_path"], "GridA", "report.pdf"))

    from pypdf import PdfReader
    reader = PdfReader(os.path.join(obj["output_path"], "GridA", "report.pdf"))
    assert reader.is_encrypted


def test_pdf_dist_random_password_reuses_existing_blank_column(tmp_path):
    """--random-password 不传 --password-col 时，清单里已有的空白「密码」列应被就地填满，
    而不是在表尾再追加一个同名列（追加会让 read_mapping 命中原来那个空列 → 密码全漏读）。

    上游 v2.7.2 修了这个（issue #1），包装层原本有一段规避、现已删除——
    这个用例就是那段规避的替代品：vendor 一旦退回旧版本，这里会立刻红。
    """
    mp = tmp_path / "grid.xlsx"
    _make_mapping(mp, [("GridA", None, "Zhang"), ("GridB", None, "Li")],
                  headers=("网格", "密码", "接收人"))
    rc, obj, stderr = _run("er_pdf_dist.py", "--mapping", str(mp),
                           "--grid-col", "网格", "--random-password")
    assert rc == 0, stderr
    assert obj["generated"] == 2
    out = obj["mapping_with_passwords"]
    wb = openpyxl.load_workbook(out)
    header = [c.value for c in wb.active[1]]
    assert header.count("密码") == 1, f"密码列被重复追加了：{header}"
    assert obj["password_col"] == "密码"
    # 真的填进去了，而且是填在原来那一列
    pwds = [r[header.index("密码")].value for r in wb.active.iter_rows(min_row=2)]
    assert all(p for p in pwds), f"密码没填进既有列：{pwds}"


def test_pdf_dist_missing_columns_reports_friendly_error(tmp_path):
    pdf_path = tmp_path / "report.pdf"
    _make_pdf(pdf_path)
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [("GridA", "001234", "Zhang")])
    rc, obj, _ = _run("er_pdf_dist.py", "--pdf", str(pdf_path), "--mapping", str(mp))
    assert rc == 1
    assert obj["ok"] is False
    assert "--grid-col" in obj["error"]


def _vendor_has_office():
    """vendor 内核 v2.10 起才会给 Office 文件加密；发版打 tag 时才同步，之前跳过这几条。"""
    with open(os.path.join(SCRIPTS, "vendor", "core", "pdf_dist.py"), encoding="utf-8") as f:
        return "OFFICE_EXTS" in f.read()


needs_office_core = pytest.mark.skipif(not _vendor_has_office(),
                                       reason="vendor/core 还没同步到 v2.10（打 tag 时才同步）")
ROW = ("E001", "张三", "客服部", "东区", 100)


@needs_office_core
def test_pdf_dist_encrypts_word_and_excel(tmp_path):
    """--files 混给 PDF 和 xlsx：xlsx 加了打开密码、对的密码能解开，内容不变；manifest 不混进 files。"""
    pytest.importorskip("msoffcrypto")
    import io

    import msoffcrypto
    pdf_path = tmp_path / "report.pdf"
    _make_pdf(pdf_path)
    xlsx = tmp_path / "明细.xlsx"
    _make_book(xlsx, [ROW])
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [("GridA", "001234", "Zhang")])
    rc, obj, stderr = _run("er_pdf_dist.py", "--files", str(pdf_path), str(xlsx), "--mapping", str(mp),
                           "--grid-col", "GridName", "--password-col", "Pwd",
                           "--output", str(tmp_path / "dist"))
    assert rc == 0, stderr
    assert obj["encrypted"] is True and obj["watermarked"] is True
    assert obj["files_total"] == 2
    assert all(not f.endswith("分发清单.xlsx") for f in obj["files"])
    assert "只加了密码" in obj["message"]
    with open(os.path.join(obj["output_path"], "GridA", "明细.xlsx"), "rb") as f:
        of = msoffcrypto.OfficeFile(f)
        assert of.is_encrypted()
        of.load_key(password="001234")
        buf = io.BytesIO()
        of.decrypt(buf)
    assert openpyxl.load_workbook(buf).active["A1"].value == "月度报表"


@needs_office_core
def test_pdf_dist_watermark_only_needs_no_password_column(tmp_path):
    pdf_path = tmp_path / "report.pdf"
    _make_pdf(pdf_path)
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [("GridA", "Zhang")], headers=("GridName", "接收人"))
    rc, obj, stderr = _run("er_pdf_dist.py", "--files", str(pdf_path), "--mapping", str(mp),
                           "--grid-col", "GridName", "--no-encrypt", "--output", str(tmp_path / "dist"))
    assert rc == 0, stderr
    assert obj["encrypted"] is False
    assert "明文密码" not in obj["message"]
    from pypdf import PdfReader
    assert not PdfReader(os.path.join(obj["output_path"], "GridA", "report.pdf")).is_encrypted


def test_pdf_dist_rejects_office_watermark_old_format_and_both_off(tmp_path):
    """这些都在开跑前就说清楚，不用等内核跳过文件后才发现。"""
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [("GridA", "001234", "Zhang")])
    xlsx = tmp_path / "a.xlsx"
    _make_book(xlsx, [ROW])
    old = tmp_path / "b.xls"
    old.write_bytes(b"x")
    base = ["--mapping", str(mp), "--grid-col", "GridName", "--password-col", "Pwd"]
    rc, obj, _ = _run("er_pdf_dist.py", "--files", str(xlsx), *base, "--no-encrypt")
    assert rc == 1 and "不能加水印" in obj["error"]
    rc, obj, _ = _run("er_pdf_dist.py", "--files", str(old), *base)
    assert rc == 1 and "老格式" in obj["error"]
    rc, obj, _ = _run("er_pdf_dist.py", "--pdf", str(xlsx), *base, "--no-encrypt", "--no-watermark")
    assert rc == 1 and "至少要一个" in obj["error"]


def _run_without_pandas(script, *args):
    """把 pandas / numpy 屏蔽掉再跑脚本：v3.0 起看表、拆分都不该再碰它们。"""
    path = os.path.join(SCRIPTS, script)
    code = ("import runpy, sys;"
            "sys.modules['pandas'] = None; sys.modules['numpy'] = None;"
            f"sys.argv = [{path!r}] + {list(args)!r};"
            f"runpy.run_path({path!r}, run_name='__main__')")
    proc = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, encoding="utf-8")
    lines = [line for line in proc.stdout.splitlines() if line.strip()]
    assert lines, proc.stderr
    return proc.returncode, json.loads(lines[-1]), proc.stderr


def test_inspect_and_split_do_not_need_pandas(two_books, tmp_path):
    """首次使用只装 openpyxl + xlrd 几秒钟，不再等 pandas + numpy 下载一分钟。"""
    inp, _ = two_books
    rc, obj, err = _run_without_pandas("er_inspect.py", "--input", inp, "--column", "部门", "--all-files")
    assert rc == 0, err
    assert obj["values"] == ["技术部", "销售部"]
    rc, obj, err = _run_without_pandas("er_inspect.py", "--input", os.path.join(inp, "A.xlsx"))
    assert rc == 0, err
    assert obj["group_counts"]["部门"] == 2
    rc, obj, err = _run_without_pandas("er_split.py", "--input", inp, "--by", "部门",
                                       "--output", str(tmp_path / "split"))
    assert rc == 0 and obj["ok"], err


def test_split_reports_unassigned_rows_and_keeps_them_out_of_bundle(tmp_path):
    """v3.0：没写分组的行进「请核对_无归属.xlsx」——结果 JSON 要说出来；它不进总包，也不被当成一个分组按人拆。"""
    src = tmp_path / "表.xlsx"
    _make_book(src, [["001", "张三", "销售部", "北京", 100], ["002", "李四", "", "上海", 200],
                     ["003", "王五", "技术部", "广州", 300]])
    rc, obj, stderr = _run("er_split.py", "--input", str(src), "--by", "部门", "--to-person", "姓名",
                           "--output", str(tmp_path / "out"))
    assert rc == 0, stderr
    s = obj["summary"]
    assert s["unassigned"] == 1 and s["source_rows"] == s["routed"] + s["skip_rows"] + s["unassigned"]
    assert os.path.exists(os.path.join(obj["output_path"], "请核对_无归属.xlsx"))
    assert any("请核对_无归属" in w for w in obj["warnings"]) and "请核对" in obj["message"]
    assert not os.path.isdir(os.path.join(obj["output_path"], "请核对_无归属"))     # 没被当成分组按人拆
    import zipfile
    with zipfile.ZipFile(obj["bundle_zip"]) as zf:
        assert not any("请核对" in n for n in zf.namelist())


# ---------- 方案（v3.0）：--recipe / --save-recipe ----------

def test_split_save_and_apply_recipe(two_books, tmp_path):
    inp, out = two_books
    rec = str(tmp_path / "按部门.json")
    code, obj, err = _run("er_split.py", "--input", inp, "--output", out, "--by", "部门", "--no-zip",
                          "--merge", "--save-recipe", rec)
    assert code == 0, err
    assert obj["recipe_saved"] == rec
    with open(rec, encoding="utf-8") as f:
        data = json.load(f)
    assert data["mode"] == "excel" and data["settings"]["split_column"] == "部门"
    assert data["settings"]["merge_across_files"] is True and data["settings"]["make_zip"] is False
    assert inp not in json.dumps(data, ensure_ascii=False)        # 路径不进方案（只留表名提示）

    out2 = str(tmp_path / "out2")
    code, obj, err = _run("er_split.py", "--input", inp, "--output", out2, "--recipe", rec)
    assert code == 0, err
    assert obj["recipe"] == "按部门" and obj["split_column_resolved"] == "部门"
    assert obj["summary"]["groups"] == 2 and not obj["bundle_zip"]          # 方案里的 make_zip=False 生效

    out3 = str(tmp_path / "out3")                                            # 命令行明确给的参数优先
    code, obj, err = _run("er_split.py", "--input", inp, "--output", out3, "--recipe", rec, "--by", "城市")
    assert code == 0, err
    assert obj["split_column_resolved"] == "城市"


def test_recipe_mode_mismatch_and_broken_file(two_books, tmp_path):
    inp, out = two_books
    rec = tmp_path / "拆分.json"
    _run("er_split.py", "--input", inp, "--output", out, "--by", "部门", "--no-zip", "--save-recipe", str(rec))
    code, obj, _ = _run("er_check.py", "--input", inp, "--recipe", str(rec))
    assert code == 1 and "不能用来做" in obj["error"]
    bad = tmp_path / "坏.json"
    bad.write_text("{", encoding="utf-8")
    code, obj, _ = _run("er_split.py", "--input", inp, "--output", out, "--recipe", str(bad))
    assert code == 1 and "方案用不了" in obj["error"]
