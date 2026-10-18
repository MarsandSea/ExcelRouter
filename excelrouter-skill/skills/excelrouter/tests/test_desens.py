"""数据脱敏（v3.1）的技能端到端测试：er_desens.py、er_split.py --desens、er_merge.py --desens。

子进程调用，解析 stdout 最后一行 JSON。号码全是虚构的（身份证号按规则算校验码，银行卡号按 Luhn 补末位）。
vendor/core 里还没有 desens.py（打 tag 同步之前）时整份跳过。
"""
import glob
import json
import os
import subprocess
import sys
import zipfile

import openpyxl
import pytest

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(SKILL_DIR, "scripts")
HAS_DESENS = os.path.isfile(os.path.join(SCRIPTS, "vendor", "core", "desens.py"))

pytestmark = pytest.mark.skipif(not HAS_DESENS, reason="vendor/core 里还没有 desens.py（v3.1 起才有）")


def _id_check(s17):
    w = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    return "10X98765432"[sum(int(a) * b for a, b in zip(s17, w)) % 11]


def _luhn_ok(s):
    total = 0
    for i, ch in enumerate(reversed(s)):
        d = int(ch)
        if i % 2:
            d = d * 2 - 9 if d > 4 else d * 2
        total += d
    return total % 10 == 0


ID1 = "11010119800101001" + _id_check("11010119800101001")
ID2 = "44200019851212003" + _id_check("44200019851212003")
CARD = next(p for p in ("622202000000000000" + str(x) for x in range(10)) if _luhn_ok(p))
PHONE = "13512340001"
HEAD = ["网格", "客户姓名", "联系电话", "证件号码", "银行卡号", "备注"]


def _book(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "清单"
    ws["A1"] = "九月客户清单"
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    wb.save(path)


def _rows(grid="东区"):
    return [[grid, "张三", PHONE, ID1, CARD, "已回访"],
            [grid, "欧阳娜娜", "13900005678", ID2, CARD, f"证件 {ID2}"]]


def _flat(path):
    wb = openpyxl.load_workbook(path)
    return "\n".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r if v is not None)


def _run(script, *args):
    proc = subprocess.run([sys.executable, os.path.join(SCRIPTS, script), *args],
                          capture_output=True, text=True, encoding="utf-8")
    lines = [ln for ln in proc.stdout.splitlines() if ln.strip()]
    assert lines, f"stdout 没有 JSON。stderr:\n{proc.stderr}"
    return json.loads(lines[-1])


# ---------- er_desens.py ----------

def test_desens_dry_run_lists_columns_and_writes_nothing(tmp_path):
    src = tmp_path / "清单.xlsx"
    _book(src, _rows())
    r = _run("er_desens.py", "--input", str(src), "--scene", "内部下发", "--dry-run")
    assert r["ok"] and r["dry_run"]
    kinds = {c["column"]: (c["kind"], c["method"]) for c in r["columns"]}
    assert kinds["证件号码"] == ("身份证号", "mask") and kinds["银行卡号"][1] == "mask"
    assert kinds["联系电话"][1] == "keep"                       # 内部下发：手机号照发
    assert ID1 not in json.dumps(r, ensure_ascii=False)           # 只有列名和处理方式，没有值
    assert sorted(os.listdir(tmp_path)) == ["清单.xlsx"]


def test_desens_group_notice_masks_everything(tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    _book(src / "东区.xlsx", _rows("东区"))
    _book(src / "港口.xlsx", _rows("港口"))
    r = _run("er_desens.py", "--input", str(src), "--scene", "发群通报")
    assert r["ok"], r
    assert r["files_total"] == 2 and r["summary"]["ok"] and "复检" in r["message"]
    assert os.path.isfile(r["report"])
    for f in glob.glob(os.path.join(r["output_dir"], "**", "*_脱敏.xlsx"), recursive=True):
        text = _flat(f)
        assert ID1 not in text and ID2 not in text and CARD not in text and PHONE not in text
        assert "张三" not in text
    assert ID1 in _flat(src / "东区.xlsx")                          # 原表不动


def test_desens_set_overrides_and_bad_input(tmp_path):
    src = tmp_path / "清单.xlsx"
    _book(src, _rows())
    r = _run("er_desens.py", "--input", str(src), "--scene", "发群通报", "--set", "客户姓名:照发")
    assert r["ok"], r
    text = _flat(r["files"][0])
    assert "张三" in text and ID1 not in text
    bad = _run("er_desens.py", "--input", str(src), "--set", "备注:随便")
    assert not bad["ok"] and "看不懂" in bad["error"]
    bad = _run("er_desens.py", "--input", str(src), "--scene", "自定义")
    assert not bad["ok"] and "--set" in bad["error"]


# ---------- er_split.py --desens ----------

def test_split_desens_masks_and_zips_when_clean(tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    _book(src / "东区.xlsx", _rows("东区"))
    _book(src / "港口.xlsx", _rows("港口"))
    out = tmp_path / "out"
    r = _run("er_split.py", "--input", str(src), "--output", str(out), "--by", "网格", "--desens", "内部下发")
    assert r["ok"] and not r["desens_blocked"], r
    ds = r["summary"]["desens"]
    assert ds["ok"] and ds["scene"] == "内部下发" and "复检通过" in r["message"]
    res = r["output_path"]
    xlsx = [f for f in glob.glob(os.path.join(res, "**", "*.xlsx"), recursive=True)
            if os.path.basename(f) != "脱敏报告.xlsx"]
    assert xlsx
    for f in xlsx:
        text = _flat(f)
        assert ID1 not in text and ID2 not in text and CARD not in text
        assert PHONE in text                                       # 内部下发：手机号照发
    assert r["bundle_zip"] and glob.glob(os.path.join(res, "*.zip"))
    with zipfile.ZipFile(r["bundle_zip"]) as z:
        assert not any(n.endswith("脱敏报告.xlsx") for n in z.namelist())


def test_split_desens_recheck_fail_blocks_all_zips(tmp_path):
    """备注里夹带的号码不处理（专家 --config 关掉 notes）→ 复检没过：不打组 ZIP、不打总包、ok=False。"""
    src = tmp_path / "in"
    src.mkdir()
    _book(src / "东区.xlsx", _rows("东区"))
    _book(src / "港口.xlsx", _rows("港口"))
    cfg = tmp_path / "cfg.json"
    cfg.write_text(json.dumps({"split_column": "网格", "desens": {"scene": "发群通报", "notes": False}},
                              ensure_ascii=False), encoding="utf-8")
    r = _run("er_split.py", "--input", str(src), "--output", str(tmp_path / "out"), "--config", str(cfg))
    assert not r["ok"] and r["desens_blocked"] and not r["bundle_zip"], r
    res = r["output_path"]
    assert not glob.glob(os.path.join(res, "**", "*.zip"), recursive=True)
    assert os.path.isfile(os.path.join(res, "未通过复检_请勿发出.txt"))
    assert "请勿发出" in r["message"] and any("复检没过" in w for w in r["warnings"])


def test_split_desens_set_needs_scene(tmp_path):
    src = tmp_path / "清单.xlsx"
    _book(src, _rows())
    r = _run("er_split.py", "--input", str(src), "--output", str(tmp_path / "out"), "--by", "网格",
             "--desens-set", "备注:清空")
    assert not r["ok"] and "--desens" in r["error"]


# ---------- er_merge.py --desens ----------

def test_merge_desens_copy_for_reporting(tmp_path):
    src = tmp_path / "收表"
    src.mkdir()
    _book(src / "东区.xlsx", _rows("东区"))
    _book(src / "港口.xlsx", _rows("港口"))
    r = _run("er_merge.py", "--input", str(src), "--output", str(tmp_path / "汇总"), "--desens", "对外报送")
    assert r["ok"] and r["done"], r
    assert r["desens_output"] and r["files"][0] == r["desens_output"] and r["desens"]["ok"]
    assert "上报发这一份" in r["message"]
    text = _flat(r["desens_output"])
    assert ID1 not in text and CARD not in text and PHONE not in text and "张三" not in text
    assert ID1 in _flat(r["output"])                               # 汇总表本身不动
