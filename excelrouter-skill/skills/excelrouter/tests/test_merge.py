"""er_merge.py（汇总）/ er_check.py（查收）的端到端测试：子进程调用，解析 stdout 最后一行 JSON。

样本照真实收表场景造：表头上面有大标题和「填报单位」行、各表带合计行、有人调过列顺序、
微信另存出来的「(1)」「- 副本」、WPS 的 .et、系统导出的网页格式 .xls、加了口令的 .xlsx、
文件夹里放着「应交名单」。
"""
import json
import os
import shutil
import subprocess
import sys

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(SKILL_DIR, "scripts")
sys.path.insert(0, os.path.join(SCRIPTS, "vendor"))
try:
    from core import collect as _collect
except ImportError:          # vendor 还没同步到带汇总的内核（v2.10 之前）
    _collect = None

pytestmark = pytest.mark.skipif(_collect is None, reason="vendor/core 里还没有 collect.py（v2.10 起才有）")
NEW_CORE = _collect is not None and hasattr(_collect, "unreadable")

COLS = ["序号", "网格", "客户名称", "联系电话", "走访日期", "走访结果", "备注"]


def _book(path, grid, n, order=None, blank=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "走访登记"
    cols = order or COLS
    ws.cell(1, 1, "九月客户走访登记表").font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, "填报单位：" + grid)
    for c, name in enumerate(cols, 1):
        cell = ws.cell(3, c, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E7FD1")
    for i in range(n):
        row = dict(zip(COLS, [i + 1, grid, f"{grid}客户{i + 1}", f"1380000{i:04d}",
                              None if i in blank else f"2026-09-{i + 1:02d}", "已走访", ""]))
        for c, name in enumerate(cols, 1):
            ws.cell(4 + i, c, row[name])
    ws.cell(4 + n, 1, "合计")
    wb.save(path)


def _names_list(path, names):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["应交单位"])
    for n in names:
        ws.append([n])
    wb.save(path)


def _run(script, *args):
    proc = subprocess.run([sys.executable, os.path.join(SCRIPTS, script), *args],
                          capture_output=True, text=True, encoding="utf-8")
    lines = [ln for ln in proc.stdout.splitlines() if ln.strip()]
    assert lines, f"stdout 应有一行 JSON。stderr:\n{proc.stderr}"
    return proc.returncode, json.loads(lines[-1]), proc.stderr


def _result_rows(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[3]]
    rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=4, values_only=True)]
    return head, rows, wb


@pytest.fixture
def returned(tmp_path):
    """一个真实的收表文件夹。"""
    src = tmp_path / "九月收表"
    src.mkdir()
    _book(src / "城区.xlsx", "城区", 5)
    _book(src / "东城.xlsx", "东城", 3)
    _book(src / "西城.xlsx", "西城", 4, blank=(1, 2),
          order=["网格", "序号", "客户名称", "走访日期", "联系电话", "走访结果", "备注"])
    shutil.copy(src / "西城.xlsx", src / "西城 - 副本.xlsx")
    (src / "南城.et").write_bytes(b"\xd0\xcf\x11\xe0fake")
    (src / "北城.xls").write_text("<html><table><tr><td>网格</td></tr></table></html>", encoding="utf-8")
    (src / "开发区.xlsx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1x")
    _names_list(src / "应交名单.xlsx", ["城区", "东城", "西城", "南城", "北城", "开发区", "高新区"])
    return src


# ---------- 普通汇总 ----------

def test_merge_folder(returned):
    rc, obj, _ = _run("er_merge.py", "--input", str(returned))
    assert rc == 0 and obj["ok"] and obj["done"]
    assert os.path.isfile(obj["output"]) and obj["files"] == [obj["output"]]
    assert os.path.dirname(obj["output"]) == str(returned / "汇总结果")
    head, rows, wb = _result_rows(obj["output"])
    assert head[-1] == "来源" and head[:len(COLS)] == COLS
    assert len(rows) == 5 + 3 + 4                       # 合计行不合、副本只合一份
    xc = [r for r in rows if r["来源"] == "西城"]
    assert len(xc) == 4 and xc[0]["网格"] == "西城" and xc[0]["联系电话"] == "13800000000"   # 调过列顺序也对齐
    assert "汇总说明" in wb.sheetnames
    assert obj["summary"]["rows"] == obj["summary"]["rows_written"] == 12
    msg = obj["message"]
    assert "核对：合进来 12 行，汇总表里 12 行，一致" in msg
    assert "南城.et" in msg and "北城.xls" in msg and "开发区.xlsx" in msg      # 合不了的每个都点名
    assert "网页格式" in msg and "WPS" in msg
    assert any("名单表" in w for w in obj["warnings"])                        # 名单表不当成「表头对不上」
    assert "表头对不上" not in msg
    assert any("西城 - 副本" in w for w in obj["warnings"])


def test_merge_one_file_sheets(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for m in (1, 2, 3):
        ws = wb.create_sheet(f"{m}月")
        ws.append([f"{m}月考勤"])
        ws.append(["工号", "姓名", "出勤天数"])
        ws.append(["001", "张三", 21])
        ws.append(["002", "李四", 20])
        ws.append(["合计", "", 41])
    hidden = wb.create_sheet("下拉选项")
    hidden.sheet_state = "hidden"
    hidden.append(["是", "否"])
    f = tmp_path / "一季度考勤.xlsx"
    wb.save(f)
    rc, obj, _ = _run("er_merge.py", "--input", str(f))
    assert rc == 0 and obj["done"] and obj["summary"]["mode"] == "sheets"
    out = openpyxl.load_workbook(obj["output"]).worksheets[0]
    sources = [r[-1] for r in out.iter_rows(min_row=3, values_only=True)]
    assert sources == ["1月", "1月", "2月", "2月", "3月", "3月"]


def test_merge_resumes_across_runs(returned):
    """宿主 120 秒掐命令：每次存好进度返回 done:false，原样再跑接着合，最后不留半成品。"""
    seen = []
    for _ in range(20):
        rc, obj, _ = _run("er_merge.py", "--input", str(returned), "--budget", "0")
        assert rc == 0 and obj["ok"]
        seen.append(obj["done"])
        if obj["done"]:
            break
    assert seen[0] is False and seen[-1] is True
    _head, rows, _wb = _result_rows(obj["output"])
    assert len(rows) == 12
    left = os.listdir(returned / "汇总结果")
    assert not [n for n in left if "未完成" in n]


def test_merge_restarts_when_files_change(returned):
    rc, obj, _ = _run("er_merge.py", "--input", str(returned), "--budget", "0")
    assert obj["done"] is False
    _book(returned / "高新区.xlsx", "高新区", 2)             # 中途又交来一份
    for _ in range(20):
        rc, obj, _ = _run("er_merge.py", "--input", str(returned), "--budget", "0")
        if obj["done"]:
            break
    _head, rows, _wb = _result_rows(obj["output"])
    assert len(rows) == 14
    assert any("从头重新合" in w for w in obj["warnings"])


def test_merge_blocks_mixed_old_split_results(returned):
    old = returned / "拆分结果" / "10091530结果"
    old.mkdir(parents=True)
    _book(old / "城区.xlsx", "城区", 6)
    rc, obj, _ = _run("er_merge.py", "--input", str(returned))
    assert rc == 1 and obj["stale_dirs"]
    rc, obj, _ = _run("er_merge.py", "--input", str(returned), "--include-old-results")
    assert rc == 0 and obj["summary"]["rows"] == 18


def test_merge_split_results_skips_person_copies(tmp_path):
    """合回拆分结果：同一组里「汇总」和「到人」是同样的行，只合「汇总」那份。"""
    root = tmp_path / "10111530结果"
    for grid in ("城区", "东城"):
        (root / grid / "汇总").mkdir(parents=True)
        (root / grid / "到人").mkdir(parents=True)
        _book(root / grid / "汇总" / f"{grid}_明细.xlsx", grid, 3)
        _book(root / grid / "到人" / f"张三_{grid}.xlsx", grid, 3)
    rc, obj, _ = _run("er_merge.py", "--input", str(root))
    assert rc == 0 and obj["summary"]["rows"] == 6
    assert any("到人" in w for w in obj["warnings"])


def test_merge_does_not_load_pandas(returned):
    """汇总只要 openpyxl + xlrd：首次使用不必先装一分钟的 pandas。"""
    script = os.path.join(SCRIPTS, "er_merge.py")
    code = (f"import runpy, sys\nsys.argv = [{script!r}, '--input', {str(returned)!r}]\n"
            f"try:\n    runpy.run_path({script!r}, run_name='__main__')\nexcept SystemExit:\n    pass\n"
            "print('PANDAS', 'pandas' in sys.modules, file=sys.stderr)")
    proc = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, encoding="utf-8")
    assert "PANDAS False" in proc.stderr


def test_merge_missing_input(tmp_path):
    rc, obj, _ = _run("er_merge.py", "--input", str(tmp_path / "没有这个文件夹"))
    assert rc == 1 and "找不到" in obj["error"]


# ---------- 查收 ----------

def test_check_notice(returned):
    rc, obj, _ = _run("er_check.py", "--input", str(returned), "--due", "10月15日 18:00")
    assert rc == 0 and obj["ok"]
    assert obj["notice"].startswith("【九月收表】回收进度")
    assert "应交 7" in obj["notice"] and "高新区" in obj["notice"]
    assert "西城：走访日期 漏 2 行（第 5、6 行）" in obj["notice"]
    assert "逐人催办" in obj["notice"] and "10月15日 18:00" in obj["notice"]
    assert obj["counts"]["expected"] == 7
    assert all(os.path.isfile(p) for p in obj["files"])
    assert "er_merge.py" in obj["next"] and "--checked" in obj["next"]


@pytest.mark.skipif(not NEW_CORE, reason="读不了的文件算交了、原因写人话：需要 vendor 同步到带 unreadable() 的内核")
def test_check_unreadable_files_are_not_reported_missing(returned):
    rc, obj, _ = _run("er_check.py", "--input", str(returned))
    notice = obj["notice"]
    assert "未交（1）：高新区" in notice                       # 南城交了 .et，不能写成未交
    assert "南城" in notice and "WPS 表格自己的格式" in notice
    assert "BOF" not in notice and "zip" not in notice


def test_check_single_file_needs_by(tmp_path):
    f = tmp_path / "总表.xlsx"
    _book(f, "城区", 3)
    rc, obj, _ = _run("er_check.py", "--input", str(f))
    assert rc == 1 and obj.get("need") == "by"


def test_checked_merge_only_submitted(returned):
    rc, obj, _ = _run("er_merge.py", "--input", str(returned), "--checked")
    assert rc == 0 and obj["done"]
    _head, rows, _wb = _result_rows(obj["output"])
    assert {r["来源"] for r in rows} == {"城区", "东城", "西城"}       # 「来源」写交表人
    msg = obj["message"]
    assert "还没交、没有汇总的（1）：高新区" in msg
    assert "交了但文件读不了" in msg and "南城" in msg
