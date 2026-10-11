"""
Excel 智能拆分工具 - 文件按网格加密分发（模块名沿用 pdf_dist，是对外契约）
Copyright (c) 2026 AbeLin
MIT License

把同一批文件按「网格 → 密码」映射清单，为每个网格生成一份副本，并输出分发清单：
- PDF：专属打开密码 + 专属水印（可溯源），两者可各自关掉；
- Word / Excel / PPT（docx / xlsx / pptx 等新格式）：只加打开密码。Office 没有
  靠得住的水印（Excel 根本没有水印功能），要水印的请先转成 PDF——v2.10 作者定的。

依赖：pypdf（读写/加密）、fpdf2（生成水印页）、cryptography（AES-256 后端，
缺失时降级 RC4-128 并警告）、msoffcrypto-tool（Office 加密，**延迟导入**：
没装时只跳过 Office 文件，PDF 照常——下游技能不一定装了它）。
openpyxl 读映射表、写分发清单。
"""

import io
import os
import json
import secrets
import shutil
import subprocess
import time
import zipfile

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from core.utils import safe_filename, soft_clean

MANIFEST_NAME = "分发清单.xlsx"

# Office 新格式（OOXML，本质是 zip）才能加打开密码；老格式要先另存
OFFICE_EXTS = (".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm")
OLD_OFFICE_EXTS = (".doc", ".xls", ".ppt", ".et", ".wps", ".dps")
SUPPORTED_EXTS = (".pdf",) + OFFICE_EXTS

# 清单里没指定接收人列时，按这些列名自动认（只抄进分发清单、可用于水印 {receiver}）
_RECEIVER_HINTS = ("接收人", "收件人", "联系人", "负责人", "姓名")

# Windows 水印字体候选（保持 v2.7.2 原样，只在 Windows 分支用）
_FONT_CANDIDATES = ["simhei.ttf", "Deng.ttf", "msyhbd.ttf", "simfang.ttf", "simkai.ttf"]

# fpdf2 能加载的字体扩展名。/usr/share/fonts 下混着 .pcf.gz/.bdf/.pfb，
# 返回它们会让 add_font 直接抛异常，必须先滤掉。
# （注：fpdf2 2.8.x 起 .ttc/.otc 字体集合是支持的，旧注释说“不支持 .ttc”已过期。）
_FONT_EXTS = (".ttf", ".otf", ".ttc", ".otc")

# Linux（银河麒麟 V10 / 统信 UOS 等）水印字体候选，按优先级排列。
# ★ 单体简中字面（*SC*.otf）必须排在 .ttc 合集前面：NotoSansCJK-Regular.ttc 的
#   第 0 面通常是日文面，汉字能渲染但字形是日式变体（直/骨/今 肉眼可辨）。
#   fpdf2 的 collection_font_number 能选 SC 面，但索引因构建而异，
#   优先单体文件比猜索引稳。
# 排序原则：**黑体类优先**——水印是半透明斜排小字，黑体的可读性明显好过仿宋/楷体/魏碑。
_LINUX_FONT_CANDIDATES = [
    # ① 信创机器实测存在的字体（2026-09 在银河麒麟 V10 x86_64 上实地确认）：
    #    /usr/share/fonts/cesi/ 是「中文信息处理标准」字库，
    #    /usr/share/fonts/wps-office/ 是装了 WPS for Linux 就会有的方正字库。
    "CESI_HT_GB18030.TTF",                 # CESI 黑体，信创机器上的首选
    "FZHTK.TTF",                           # 方正黑体_GBK（WPS 自带）
    "CESI_SS_GB18030.TTF", "FZSSK.TTF",    # 宋体
    "CESI_FS_GB18030.TTF", "CESI_KT_GB18030.TTF",
    "FZKTK.TTF", "FZFSK.TTF",              # 楷体 / 仿宋（兜底，水印可读性一般）
    # ② 文泉驿：中文桌面发行版几乎必装，且是「中文优先」字形
    "wqy-zenhei.ttc", "wqy-microhei.ttc", "wqy-zenhei.ttf", "wqy-microhei.ttf",
    # ③ Noto / 思源的单体简中面
    "NotoSansCJKsc-Regular.otf", "NotoSansSC-Regular.otf",
    "SourceHanSansSC-Regular.otf", "SourceHanSansCN-Regular.otf",
    "NotoSerifCJKsc-Regular.otf",
    # ④ 合集文件兜底
    "NotoSansCJK-Regular.ttc", "SourceHanSansSC.ttc",
    # ⑤ 老发行版的 AR PL
    "uming.ttc", "ukai.ttc",
]

_LINUX_FONT_DIRS = [
    "/usr/share/fonts",
    "/usr/local/share/fonts",
    os.path.expanduser("~/.local/share/fonts"),
    os.path.expanduser("~/.fonts"),
]

# 用户指定水印字体的逃生口：信创机器的字体配置千奇百怪，自动探测不中时用它。
_FONT_ENV = "ER_CJK_FONT"


def _noop(*_args, **_kwargs):
    pass


def list_mapping_columns(xlsx_path):
    """读映射清单首个 sheet 的第 1 行表头，供 GUI 下拉选择网格列/密码列/接收人列。

    映射清单是用户专门维护的小表，默认表头就在第 1 行，不做启发式识别。
    """
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()   # 有的程序生成的表把尺寸写错（如 A1:A1），不重置只读模式会少读行列
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            return [soft_clean(v) for v in row if soft_clean(v)]
        return []
    finally:
        wb.close()


def _cell_str(v):
    """单元格值转字符串：数字型密码防止 openpyxl 读成 1234.0 / 丢前导零后的 int。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return soft_clean(v)


def guess_receiver_column(header, exclude=()):
    """在清单表头里认接收人列（接收人 / 联系人 / 姓名…），认不出返回空串。"""
    for hint in _RECEIVER_HINTS:
        for h in header:
            if h and h not in exclude and hint in h:
                return h
    return ""


def read_mapping(xlsx_path, grid_col, password_col, receiver_col="", need_password=True):
    """读「网格 → 密码（→ 接收人）」映射。

    返回 (rows, warnings)：rows 为 [{"grid", "password", "receiver"}]，
    warnings 为需要记进日志的问题（空网格/空密码行跳过、重复网格覆盖）。
    need_password=False（只加水印）时不要求密码列，密码一律为空串。
    """
    if not need_password:
        password_col = ""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        header = None
        idx = {}
        rows, warnings = [], []
        seen = {}
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if header is None:
                header = [soft_clean(v) for v in row]
                for name, col in ((grid_col, "grid"), (password_col, "password"),
                                  (receiver_col, "receiver")):
                    if name and name in header:
                        idx[col] = header.index(name)
                need = [(grid_col, "grid")] + ([(password_col, "password")] if need_password else [])
                missing = [n or ("密码列" if c == "password" else "网格列")
                           for n, c in need if c not in idx]
                if missing:
                    raise ValueError(f"映射清单里找不到列：{'、'.join(missing)}")
                continue
            grid = _cell_str(row[idx["grid"]] if idx["grid"] < len(row) else None)
            pwd = ""
            if "password" in idx:
                pwd = _cell_str(row[idx["password"]] if idx["password"] < len(row) else None)
            recv = ""
            if "receiver" in idx and idx["receiver"] < len(row):
                recv = _cell_str(row[idx["receiver"]])
            if not grid and not pwd:
                continue                      # 整行空白，静默跳过
            if not grid:
                warnings.append(f"第 {r} 行网格为空，已跳过")
                continue
            if need_password and not pwd:
                warnings.append(f"第 {r} 行「{grid}」密码为空，已跳过")
                continue
            if grid in seen:
                warnings.append(f"网格「{grid}」重复出现，以第 {r} 行为准")
                rows[seen[grid]] = {"grid": grid, "password": pwd, "receiver": recv}
                continue
            seen[grid] = len(rows)
            rows.append({"grid": grid, "password": pwd, "receiver": recv})
        return rows, warnings
    finally:
        wb.close()


# =====================================================
# 映射清单模板 / 随机密码（v2.7 新增，追加式 API）
# =====================================================

# 随机密码字符集：排除 0O1lI 等易混字符，避免微信里口口相传时认错
_PWD_ALPHABET = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKMNPQRSTUVWXYZ23456789"

def gen_password(n=8):
    """生成 n 位随机密码（加密学安全的 secrets，排除易混字符）。"""
    return "".join(secrets.choice(_PWD_ALPHABET) for _ in range(n))

def write_mapping_template(path):
    """生成「网格 → 密码」映射清单模板 xlsx，返回写入路径。

    首个 sheet 是待填写的空模板（表头：网格/密码/接收人 + 两行示例，示例行
    使用前请删除）；第二个 sheet 放填写说明，避免说明文字混进数据行被
    read_mapping 误当成网格。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "映射清单"
    ws.append(["网格", "密码", "接收人"])
    ws.append(["东区营业厅（示例，请删除）", gen_password(), "张三（示例）"])
    ws.append(["西区营业厅（示例，请删除）", gen_password(), "李四（示例）"])
    for col, w in {"A": 24, "B": 16, "C": 14}.items():
        ws.column_dimensions[col].width = w
    ws.cell(row=2, column=2).number_format = "@"
    ws.cell(row=3, column=2).number_format = "@"

    note = wb.create_sheet("填写说明")
    for line in [
        "每行一个网格：网格名必填，密码为该网格的专属打开密码（只加水印时可不填），"
        "接收人选填（抄进分发清单方便照着发；水印文字里写 {receiver} 可印上接收人）。",
        "密码建议 8 位以上，不要用 123456、手机号这类容易猜到的密码。",
        "示例行（第 2、3 行）使用前请删除。",
        "没有密码思路？可以不选密码列，工具会在开始前询问是否自动生成随机密码。",
        "加密采用 AES-256，忘记密码无法找回；含明文密码的「分发清单.xlsx」只留给分发人自己用。",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 100
    wb.save(path)
    return path

def fill_random_passwords(mapping_path, grid_col, password_col="", out_path=None):
    """为映射清单里的每个网格生成随机密码，写入**新文件**（不改动原清单）。

    password_col 已存在时只填空白单元格（已手填的密码保留）；不存在时在表尾
    新增「密码」列。返回 (输出路径, 实际生成的密码个数)。
    """
    wb = load_workbook(mapping_path)
    try:
        ws = wb.worksheets[0]
        header = [soft_clean(c.value) for c in ws[1]]
        if grid_col not in header:
            raise ValueError(f"清单表头里找不到网格列「{grid_col}」")
        g_idx = header.index(grid_col)
        # v2.7.2 修 issue #1：password_col 未指定时先复用既有「密码」列，避免追加重复
        # 同名列（旧版追加后 header.index 命中原空白列，生成的密码被 read_mapping 漏读）
        if password_col and password_col in header:
            p_idx = header.index(password_col)
        elif not password_col and "密码" in header:
            p_idx = header.index("密码")      # 复用既有列，只填空白格
        else:
            p_idx = len(header)
            ws.cell(row=1, column=p_idx + 1, value="密码")
        count = 0
        for r in range(2, ws.max_row + 1):
            if not soft_clean(ws.cell(row=r, column=g_idx + 1).value):
                continue
            cell = ws.cell(row=r, column=p_idx + 1)
            if soft_clean(cell.value):
                continue                    # 已有密码的保留，不覆盖
            cell.value = gen_password()
            cell.number_format = "@"        # 强制文本，防 Excel 吞前导零
            count += 1
        if not count:
            raise ValueError("没有需要生成密码的网格（可能密码列都已填好）")
        if not out_path:
            stem, ext = os.path.splitext(mapping_path)
            out_path = f"{stem}_含密码{ext or '.xlsx'}"
        wb.save(out_path)
        return out_path, count
    finally:
        wb.close()

def _font_usable(path):
    """fpdf2 到底能不能加载这个字体文件——让它自己回答，别靠扩展名猜。

    requirements.txt 只钉了 fpdf2>=2.8.0，旧版对 .ttc 的支持不一定齐；
    而且 /usr/share/fonts 下也可能躺着损坏的字体文件。这一关是
    「优雅降级成 ???」和「分发跑到一半崩掉」的分界线。
    """
    try:
        from fpdf import FPDF
        probe = FPDF()
        probe.add_font("probe", "", path)
        return True
    except Exception:
        return False


def _find_cjk_font():
    """找一款可用于水印的中文字体，找不到返回 None。

    Windows：行为与 v2.7.2 完全一致（只扫系统 Fonts 目录里的那几个 .ttf）。
    其余平台（麒麟 / UOS / 其它 Linux）：环境变量 → 已知候选名 → fc-match → None。

    函数名、零参签名、str|None 返回值都是对外契约：下游 excelrouter-skill 的
    er_pdf_dist.py 靠猴补丁替换它来实现 --font 覆盖，别改签名。
    """
    return _find_cjk_font_win() if os.name == "nt" else _find_cjk_font_posix()


def _find_cjk_font_win():
    """Windows 上的中文字体探测（与 v2.7.2 逐字一致）。"""
    fonts_dir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    for name in _FONT_CANDIDATES:
        p = os.path.join(fonts_dir, name)
        if os.path.exists(p):
            return p
    return None


def _find_cjk_font_posix():
    """Linux/macOS 上的中文字体探测。"""
    # ① 用户显式指定（FAQ 里会教这个环境变量）
    env = os.environ.get(_FONT_ENV, "").strip()
    if env and os.path.isfile(env) and _font_usable(env):
        return env

    # ② 建一次「小写文件名 → 全路径」索引，再按候选顺序查
    index = {}
    for root_dir in _LINUX_FONT_DIRS:
        if not os.path.isdir(root_dir):
            continue
        n_walk = 0
        for root, _dirs, files in os.walk(root_dir, followlinks=False):
            n_walk += 1
            if n_walk > 800:        # 字体目录不该有这么深，防病态目录卡住
                break
            for f in files:
                if f.lower().endswith(_FONT_EXTS):
                    index.setdefault(f.lower(), os.path.join(root, f))
    for name in _LINUX_FONT_CANDIDATES:
        p = index.get(name.lower())
        if p and _font_usable(p):
            return p

    # ③ 问 fontconfig 要一款「中文 sans」——最通用的一条路
    # timeout 不是装饰：刚开机、fontconfig 缓存是冷的时候，fc-match 能卡好几秒。
    if shutil.which("fc-match"):
        try:
            r = subprocess.run(["fc-match", "-f", "%{file}", "sans-serif:lang=zh-cn"],
                               capture_output=True, text=True, timeout=3)
            p = (r.stdout or "").strip()
            if (p and p.lower().endswith(_FONT_EXTS)
                    and os.path.isfile(p) and _font_usable(p)):
                return p
        except (OSError, subprocess.SubprocessError):
            pass
    return None


_wm_cache = {}


def _make_watermark_pdf(text, w_pt, h_pt, *, angle=45, opacity=0.15, font_path=None):
    """用 fpdf2 生成一页平铺斜排水印 PDF，返回字节串。按（尺寸+文字）缓存。

    没有中文字体时退回内置 Helvetica，非 ASCII 字符替换成 '?'（上层已警告过）。
    """
    key = (round(w_pt), round(h_pt), text, angle, opacity, font_path)
    if key in _wm_cache:
        return _wm_cache[key]

    from fpdf import FPDF
    pdf = FPDF(unit="pt", format=(w_pt, h_pt))
    pdf.set_auto_page_break(False)
    pdf.add_page()
    if font_path:
        # 守卫 add_font：一个坏字体文件不该在已经写出一部分网格之后，
        # 把整轮分发打断——降级成 ??? 也比跑一半崩掉强。
        try:
            pdf.add_font("wm", "", font_path)
            pdf.set_font("wm", size=18)
            draw_text = text
        except Exception:
            font_path = None
    if not font_path:
        pdf.set_font("Helvetica", size=18)
        draw_text = text.encode("ascii", "replace").decode("ascii")
    pdf.set_text_color(128, 128, 128)
    text_w = pdf.get_string_width(draw_text)
    step_x = max(text_w + 60, 180)
    step_y = 120
    diag = (w_pt ** 2 + h_pt ** 2) ** 0.5
    with pdf.local_context(fill_opacity=opacity):
        # 绕页面中心旋转后铺一个盖满对角线的网格，保证旋转后无死角
        with pdf.rotation(angle, w_pt / 2, h_pt / 2):
            y = h_pt / 2 - diag / 2
            row = 0
            while y < h_pt / 2 + diag / 2:
                x = w_pt / 2 - diag / 2 + (step_x / 2 if row % 2 else 0)
                while x < w_pt / 2 + diag / 2:
                    pdf.text(x, y, draw_text)
                    x += step_x
                y += step_y
                row += 1
    data = bytes(pdf.output())
    _wm_cache[key] = data
    return data


def _encrypt_algorithm(log_fn):
    """优先 AES-256；cryptography 缺失时降级 RC4-128 并警告（不中断）。"""
    try:
        import cryptography  # noqa: F401
        return "AES-256"
    except ImportError:
        log_fn("⚠ 未安装 cryptography，加密降级为 RC4-128（强度较弱，建议 pip install cryptography）")
        return "RC4-128"


def _stamp_and_encrypt(reader, out_path, wm_text, password, algorithm,
                       *, opacity=0.15, angle=45, font_path=None):
    """把 reader 的所有页盖上水印、设打开密码，写到 out_path。返回页数。
    wm_text 为空不盖水印，password 为空不加密（两者至少有一个，由上层保证）。

    writer 每个网格都要新建：pypdf 的 merge_page 会就地修改页对象，
    复用会导致水印跨网格叠加。append(reader) 会把页深拷贝进 writer，
    所以改 writer.pages 不会污染共享的 reader。
    """
    writer = PdfWriter()
    writer.append(reader)
    if wm_text:
        wm_readers = {}
        for page in writer.pages:
            w, h = float(page.mediabox.width), float(page.mediabox.height)
            k = (round(w), round(h))
            if k not in wm_readers:
                data = _make_watermark_pdf(wm_text, w, h, angle=angle,
                                           opacity=opacity, font_path=font_path)
                wm_readers[k] = PdfReader(io.BytesIO(data))
            page.merge_page(wm_readers[k].pages[0])
    if password:
        writer.encrypt(user_password=password, algorithm=algorithm)
    with open(out_path, "wb") as f:
        writer.write(f)
    return len(writer.pages)


def _office_encryptor():
    """返回 msoffcrypto 的 OOXMLFile 类；没装返回 None（只跳过 Office 文件）。"""
    try:
        from msoffcrypto.format.ooxml import OOXMLFile
        return OOXMLFile
    except ImportError:
        return None


def _encrypt_office(ooxml_cls, data, out_path, password):
    """给 docx/xlsx/pptx 加打开密码（ECMA-376 Agile，AES-256），写到 out_path。

    直接对文件字节加密，不经 openpyxl 重写——原文件里的图表、图片、宏都原样保留。
    """
    with open(out_path, "wb") as f:
        ooxml_cls(io.BytesIO(data)).encrypt(password, f)


def write_manifest(entries, out_path, with_password=True):
    """写分发清单：网格 | 文件名 | 密码 | 接收人 | 页数 | 状态。

    密码列强制文本格式，防止 Excel 把 001234 显示成 1234。
    清单含明文密码——只给分发人自己用，别随文件一起发出去。
    with_password=False（只加水印）时不写密码列。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "分发清单"
    cols = ["网格", "文件名", "密码", "接收人", "页数", "状态"]
    widths = [18, 40, 16, 14, 8, 24]
    if not with_password:
        del cols[2], widths[2]
    ws.append(cols)
    for e in entries:
        vals = [e["grid"], e.get("file", ""), e.get("password", ""),
                e.get("receiver", ""), e.get("pages", ""), e.get("status", "")]
        if not with_password:
            del vals[2]
        ws.append(vals)
        if with_password:
            ws.cell(row=ws.max_row, column=3).number_format = "@"
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEF"[i]].width = w
    wb.save(out_path)


def run_pdf_dist(config, log_fn=None, progress_fn=None, stop_flag=None):
    """文件加密分发主流程。签名与 run_split 一致，GUI 队列泵直接复用。

    config 键沿用 pdf_* 前缀（对外契约）：pdf_input_paths 里可以混放 PDF 和
    docx/xlsx/pptx；pdf_encrypt（默认 True）/ pdf_watermark（默认 True）各自可关，
    但至少开一个。水印只加在 PDF 上；Office 文件只加密，只加水印时跳过并说明。
    pdf_receiver_column 留空时按列名自动认接收人列；水印文字支持 {receiver}。

    单个网格/文件失败不中断整体，记日志继续（项目错误处理约定）。
    返回输出根目录；一个都没成功时返回 None。
    """
    log_fn = log_fn or _noop
    progress_fn = progress_fn or _noop
    stop_flag = stop_flag or (lambda: False)

    in_paths = [p for p in config.get("pdf_input_paths", []) if p]
    mapping_path = config.get("pdf_mapping_path", "")
    output_root = config.get("output_path", "")
    encrypt_on = config.get("pdf_encrypt", True)
    watermark_on = config.get("pdf_watermark", True)
    if not in_paths:
        log_fn("❌ 没有选择要分发的文件")
        return None
    for p in in_paths:
        if not os.path.exists(p):
            log_fn(f"❌ 找不到文件：{p}")
            return None
    if not encrypt_on and not watermark_on:
        log_fn("❌ 「加密」和「水印」至少要选一个")
        return None
    if not mapping_path or not os.path.exists(mapping_path):
        log_fn("❌ 找不到映射清单文件")
        return None
    if not output_root:
        output_root = os.path.join(os.path.dirname(in_paths[0]), "分发结果")

    grid_col = config.get("pdf_grid_column", "")
    password_col = config.get("pdf_password_column", "")
    receiver_col = config.get("pdf_receiver_column", "")
    if not receiver_col:
        try:
            receiver_col = guess_receiver_column(list_mapping_columns(mapping_path),
                                                 exclude=(grid_col, password_col))
        except Exception:
            receiver_col = ""
        if receiver_col:
            log_fn(f"👤 接收人列：自动认出「{receiver_col}」（抄进分发清单）")
    try:
        rows, warnings = read_mapping(mapping_path, grid_col, password_col,
                                      receiver_col, need_password=encrypt_on)
    except ValueError as e:
        log_fn(f"❌ {e}")
        return None
    for w in warnings:
        log_fn(f"⚠ 映射清单：{w}")
    if not rows:
        log_fn("❌ 映射清单里没有有效的「网格 + 密码」行" if encrypt_on
               else "❌ 映射清单里没有有效的网格行")
        return None

    wm_template = config.get("pdf_watermark_text", "{grid} {date}") or "{grid} {date}"
    opacity = float(config.get("pdf_watermark_opacity", 0.15))
    angle = float(config.get("pdf_watermark_angle", 45))

    # 逐个看输入：PDF 开一个 reader 跨网格复用（writer 每网格重建，见 _stamp_and_encrypt）；
    # Office 读一次字节跨网格复用
    pdfs, offices = {}, {}
    ooxml_cls = None
    for p in in_paths:
        name = os.path.basename(p)
        ext = os.path.splitext(p)[1].lower()
        try:
            if ext == ".pdf":
                r = PdfReader(p)
                if r.is_encrypted:
                    log_fn(f"⚠ 「{name}」本身带密码，无法处理，已跳过（请先解密另存）")
                    continue
                pdfs[p] = r
            elif ext in OFFICE_EXTS:
                if not encrypt_on:
                    log_fn(f"⚠ 「{name}」：Word / Excel / PPT 只能加密、不能加水印，已跳过"
                           "（需要水印请先转成 PDF）")
                    continue
                if ooxml_cls is None:
                    ooxml_cls = _office_encryptor() or False
                    if not ooxml_cls:
                        log_fn("⚠ 未安装 msoffcrypto-tool，Word / Excel / PPT 文件跳过"
                               "（pip install msoffcrypto-tool）")
                if not ooxml_cls:
                    continue
                with open(p, "rb") as f:
                    data = f.read()
                if not zipfile.is_zipfile(io.BytesIO(data)):
                    log_fn(f"⚠ 「{name}」本身带密码或文件已损坏，已跳过（请先去掉密码另存）")
                    continue
                offices[p] = data
            elif ext in OLD_OFFICE_EXTS:
                log_fn(f"⚠ 「{name}」是老格式，加不了密码，已跳过"
                       "（请用 Word / Excel / WPS 另存为 .docx / .xlsx / .pptx）")
            else:
                log_fn(f"⚠ 「{name}」不是 PDF / Word / Excel / PPT，已跳过")
        except Exception as e:
            log_fn(f"⚠ 读取「{name}」失败：{e}，已跳过")
    sources = list(pdfs) + list(offices)
    if not sources:
        log_fn("❌ 没有可处理的文件")
        return None
    if not pdfs:
        watermark_on = False          # 只有 Office 文件：没有能加水印的，也不必找字体
    elif watermark_on and offices:
        log_fn("💧 水印只加在 PDF 上；Word / Excel / PPT 只加密（需要水印请先转成 PDF）")

    font_path = None
    if watermark_on:
        font_path = _find_cjk_font()
        if font_path:
            log_fn(f"🔤 水印字体：{font_path}")
        else:
            log_fn("⚠ 没找到可用的中文字体，水印里的中文会显示成 '?'。"
                   "Windows：确认系统字体没被精简；"
                   "Linux（麒麟/UOS）：sudo apt install -y fonts-wqy-zenhei（或 fonts-noto-cjk）；"
                   "也可以设环境变量 ER_CJK_FONT 指向一个 .ttf/.otf/.ttc 文件")
    algorithm = _encrypt_algorithm(log_fn) if encrypt_on and pdfs else "AES-256"
    what = "、".join(x for x, on in (("加密", encrypt_on), ("加水印", watermark_on)) if on)
    log_fn(f"📋 映射清单：{len(rows)} 个网格；待分发文件：{len(sources)} 个（{what}）")

    os.makedirs(output_root, exist_ok=True)
    date_str = time.strftime("%Y-%m-%d")
    entries = []
    total = len(rows) * len(sources)
    done = 0
    ok_grids, fail_grids = 0, 0
    stopped = False
    for row in rows:
        if stop_flag():
            log_fn("⏹ 已停止，剩余网格未处理")
            stopped = True
            break
        grid, receiver = row["grid"], row["receiver"]
        password = row["password"] if encrypt_on else ""
        grid_dir = os.path.join(output_root, safe_filename(grid))
        wm_text = (wm_template.replace("{grid}", grid).replace("{date}", date_str)
                   .replace("{receiver}", receiver or grid)) if watermark_on else ""
        grid_ok = True
        for src in sources:
            base = safe_filename(os.path.basename(src))
            out_path = os.path.join(grid_dir, base)
            entry = {"grid": grid, "password": password, "receiver": receiver,
                     "file": os.path.relpath(out_path, output_root)}
            try:
                os.makedirs(grid_dir, exist_ok=True)
                if src in pdfs:
                    entry["pages"] = _stamp_and_encrypt(
                        pdfs[src], out_path, wm_text, password, algorithm,
                        opacity=opacity, angle=angle, font_path=font_path)
                else:
                    _encrypt_office(ooxml_cls, offices[src], out_path, password)
                entry["status"] = "✅ 完成"
            except Exception as e:
                grid_ok = False
                log_fn(f"❌ 「{grid}」× {base} 失败：{e}")
                entry["status"] = f"失败：{e}"
            entries.append(entry)
            done += 1
            progress_fn(done / total * 0.95)
        if grid_ok:
            ok_grids += 1
            log_fn(f"{'🔐' if encrypt_on else '💧'} {grid} → {len(sources)} 个文件已{what}")
        else:
            fail_grids += 1

    manifest_path = os.path.join(output_root, MANIFEST_NAME)
    try:
        write_manifest(entries, manifest_path, with_password=encrypt_on)
        log_fn(f"📄 分发清单已生成：{MANIFEST_NAME}")
    except Exception as e:
        log_fn(f"⚠ 分发清单写入失败：{e}")
    progress_fn(1.0)

    if stopped:
        log_fn(f"⏹ 已停止：完成 {ok_grids} 个网格")
    else:
        log_fn(f"✅ 完成：{ok_grids} 个网格成功" + (f"，{fail_grids} 个失败" if fail_grids else ""))
    if encrypt_on:
        log_fn("🔒 提醒：分发清单里有明文密码，只留给自己用，不要随文件一起发出去")

    # 机器可读收尾行（v2.7）：GUI 提取后渲染完成摘要；追加式，不改返回值与契约。
    summary = {
        "mode": "pdf",
        "grids_ok": ok_grids,
        "grids_fail": fail_grids,
        "files": done,
        "encrypt": encrypt_on,
        "watermark": watermark_on,
        "stopped": stopped,
        "output": output_root if ok_grids else None,
        "manifest": manifest_path if ok_grids else None,
    }
    log_fn(f"[SUMMARY] {json.dumps(summary, ensure_ascii=False)}")
    return output_root if ok_grids else None
