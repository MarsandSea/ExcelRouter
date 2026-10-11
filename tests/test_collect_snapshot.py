"""查收通报 / 催办话术 / 回收进度表的全文快照：改 collect 的内部结构时，用户看到的文字必须一字不差。

场景：未交、空表、漏填、改动、表头对不上、读不了、重复交、旧版本、名单外、多 sheet、总表、没有名单、全部交齐。
期望值在 tests/golden/collect/*.json；有意改了文案时用下面的命令重新生成，并在提交里说明改了什么：
    py tests/test_collect_snapshot.py --update

Copyright (c) 2026 AbeLin · MIT License
"""
import datetime
import json
import os
import sys

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from core import collect  # noqa: E402

EXPECTED = os.path.join(ROOT, "tests", "golden", "collect")
HEAD = ["客户名称", "客户经理", "走访日期", "走访结果", "备注"]
NOW = datetime.datetime(2026, 10, 12, 9, 30)
MTIME = datetime.datetime(2026, 10, 11, 16, 5).timestamp()


def _book(path, rows, head=HEAD, title=None, sheet="走访", more=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if title:
        ws.append([title])
    ws.append(list(head))
    for r in rows:
        ws.append(list(r))
    for name, (h2, rs2) in (more or {}).items():
        s = wb.create_sheet(name)
        s.append(list(h2))
        for r in rs2:
            s.append(list(r))
    wb.save(path)
    os.utime(path, (MTIME, MTIME))


def _full(n, who="经理"):
    return [[f"客户{i}", who, f"2026-10-0{i % 9 + 1}", "已走访", ""] for i in range(1, n + 1)]


def _write(path, data):
    with open(path, "wb") as f:
        f.write(data)
    os.utime(path, (MTIME, MTIME))


# ---------- 场景 ----------

def _files_mixed(d):
    """文件夹 + 名单：各种状态都有一个。"""
    _book(d / "张三.xlsx", _full(4))
    rows = _full(12)
    for i in (1, 3, 4, 6, 7, 8, 9, 10, 11):     # 走访日期漏 9 行：超过 SHOW_ROWS，出现「等 N 行」
        rows[i][2] = None
    rows[3][3] = None
    _book(d / "李四_走访表.xlsx", rows)
    _book(d / "赵六.xlsx", [])
    _book(d / "孙八.xlsx", [[f"客户{i}", "经理", None, None, ""] for i in range(3)])
    _write(d / "周九.et", b"\xd0\xcf\x11\xe0fake")
    _write(d / "冯三.xls", b"<html><table><tr><td>x</td></tr></table></html>")
    _book(d / "吴十.xlsx", [["甲", "09:00", "x"]], head=["姓名", "签到时间", "地点"])
    old = _full(2)
    old[0][3] = None
    _book(d / "郑一.xlsx", old)
    os.utime(d / "郑一.xlsx", (MTIME - 3600, MTIME - 3600))
    _book(d / "郑一(2).xlsx", _full(2))
    _write(d / "钱二.xlsx", (d / "张三.xlsx").read_bytes())
    _write(d / "张三 - 副本.xlsx", (d / "张三.xlsx").read_bytes())
    _book(d / "别的科室.xlsx", _full(1))
    multi = _full(2)
    multi[1][3] = None
    _book(d / "陈七.xlsx", _full(2), more={"回访": (HEAD, multi)})
    return {"input_path": str(d), "collect_expect": "张三、李四、王五、赵六、孙八、周九、冯三、吴十、郑一、钱二、陈七",
            "collect_required": ["走访日期", "走访结果"], "collect_due": "10月15日 18:00",
            "collect_title": "九月走访"}, False


def _changes(d):
    """有下发记录：删了列、少了行、多了行、改了格子、只是排序、填写列整列没有。"""
    issued = d / "下发"
    issued.mkdir()
    names = ("张三", "李四", "王五", "赵六", "钱八", "郑十")
    fill = ["走访日期", "走访结果"]
    keys = collect.lock_keys(HEAD, fill)
    rows = {n: [[f"{n}客户{i}", n, None, None, ""] for i in range(1, 5)] for n in names}
    groups = {}
    for n, rs in rows.items():
        p = issued / f"{n}.xlsx"
        _book(p, rs)
        groups[n] = {"file": p.name, "sheets": collect.fingerprint_file(str(p), {}, keys)}
    record = {"title": "九月走访", "due": "10月15日", "fill_columns": fill, "lock_keys": keys,
              "lock_columns": [h for h in HEAD if h not in fill], "groups": groups}
    back = d / "回收"
    back.mkdir()
    (back / collect.RECORD_NAME).write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")

    def filled(rs):
        return [[r[0], r[1], "10-01", "好", r[4]] for r in rs]

    _book(back / "张三.xlsx", filled(rows["张三"])[:3])                                  # 少了 1 行
    li = filled(rows["李四"])
    li[1][0] = "改过的客户"
    li[2][1] = "别人"
    _book(back / "李四.xlsx", li)                                                         # 改了两列
    _book(back / "王五.xlsx", list(reversed(filled(rows["王五"]))))                       # 只是排序
    _book(back / "赵六.xlsx", filled(rows["赵六"]) + [["新客户", "赵六", "10-02", "好", ""]] * 2)   # 多了 2 行
    _book(back / "钱八.xlsx", [[r[1], r[2], r[3], r[4]] for r in filled(rows["钱八"])],
          head=["客户经理", "走访日期", "走访结果", "备注"])                                 # 删了「客户名称」列
    zr = filled(rows["郑十"])
    _book(back / "郑十.xlsx", [[r[0], r[1], r[2], r[4]] for r in zr],
          head=["客户名称", "客户经理", "走访日期", "备注"])                                 # 填写列「走访结果」整列没有
    return {"input_path": str(back)}, True


def _rows_mode(d):
    path = d / "总表.xlsx"
    head = ["网格", "客户名称", "走访日期", "走访结果"]
    rows = [["东区", "a", "10-01", "好"], ["东区", "b", "10-02", "好"],
            ["西区", "c", None, "好"], ["西区", "d", "10-03", None],
            ["南区", "e", None, None], ["北外", "f", "10-01", "好"], [None, "g", "10-01", "好"]]
    _book(path, rows, head=head, title="九月走访总表")
    return {"input_path": str(path), "collect_by": "网格", "collect_expect": "东区、西区、南区、中区",
            "collect_required": "走访日期、走访结果"}, False


def _no_list(d):
    for n in ("甲", "乙", "丙"):
        _book(d / f"{n}.xlsx", _full(3))
    rows = _full(3)
    rows[0][3] = None
    _book(d / "丁.xlsx", rows)
    _book(d / "戊.xlsx", [])
    return {"input_path": str(d), "collect_title": "九月走访"}, False


def _all_ok(d):
    for n in ("张三", "李四"):
        _book(d / f"{n}.xlsx", _full(2))
    return {"input_path": str(d), "collect_expect": "张三、李四", "collect_title": "九月走访"}, False


SCENARIOS = {"files_mixed": _files_mixed, "changes": _changes, "rows_mode": _rows_mode, "no_list": _no_list,
             "all_ok": _all_ok}
NEW = {"all_ok": ["李四"]}          # 模拟「比上次新交」


def _progress_cells(path):
    wb = openpyxl.load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        rows = [[v for v in r] for r in ws.iter_rows(values_only=True)]
        out[ws.title] = rows
        if ws.title == "回收进度":
            out["状态颜色"] = [ws.cell(row=r, column=3).fill.fgColor.rgb for r in range(3, ws.max_row + 1)]
    wb.close()
    return out


def snapshot(name, tmp):
    """跑一个场景，返回可以比较的数据：通报全文、逐人催办、进度表单元格、逐人的旧字段。"""
    tmp.mkdir(parents=True, exist_ok=True)
    cfg, has_fill = SCENARIOS[name](tmp)
    ctx, _res, report = collect._inspect(cfg, lambda m: None, None, None)
    if name in NEW:
        report["new"] = NEW[name]
    title, due = ctx["title"], ctx["due"]
    data = {
        "notice": collect.notice_text(report, title, due, NOW, has_fill).split("\n"),
        "reminds": {p["name"]: collect._remind(p, title, due, has_fill) for p in report["people"]},
        "progress": _progress_cells(collect.write_progress(report, str(tmp / "进度.xlsx"), title, NOW, due,
                                                           has_fill)),
        "people": [{k: p[k] for k in ("name", "status", "flags", "files", "rows", "blanks", "changes", "notes",
                                      "time")} for p in report["people"]],
        "report": {k: report[k] for k in ("extra", "dupes", "errors", "required", "unknown_required", "notes",
                                          "mode")},
        "counts": collect.counts(report),
    }
    return json.loads(json.dumps(data, ensure_ascii=False, default=str).replace(str(tmp), "<tmp>"))


def _expected_path(name):
    return os.path.join(EXPECTED, f"{name}.json")


@pytest.mark.parametrize("name", list(SCENARIOS))
def test_collect_report_snapshot(name, tmp_path):
    with open(_expected_path(name), encoding="utf-8") as f:
        expected = json.load(f)
    got = snapshot(name, tmp_path / name)
    for key in expected:
        assert got[key] == expected[key], f"{name}: {key} 和快照不一致"
    assert set(got) == set(expected)


if __name__ == "__main__" and "--update" in sys.argv:
    import pathlib
    import tempfile
    os.makedirs(EXPECTED, exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        for n in SCENARIOS:
            snap = snapshot(n, pathlib.Path(t) / n)
            with open(_expected_path(n), "w", encoding="utf-8") as f:
                json.dump(snap, f, ensure_ascii=False, indent=1)
                f.write("\n")
            print("wrote", n)


# ---------- 问题编码（v3.0）----------

def _issues(name, tmp):
    tmp.mkdir(parents=True, exist_ok=True)
    cfg, _ = SCENARIOS[name](tmp)
    _ctx, _res, report = collect._inspect(cfg, lambda m: None, None, None)
    return {p["name"]: [{k: v for k, v in i.items() if k != "why" or i["code"] == "empty"} for i in p["issues"]]
            for p in report["people"]}


def test_issue_codes_files_mixed(tmp_path):
    got = _issues("files_mixed", tmp_path)
    assert got["张三"] == [] and got["王五"] == [{"code": "missing"}]
    assert got["李四"][0] == {"code": "blank", "col": "走访日期", "count": 9, "rows": [3, 5, 6, 8, 9, 10, 11, 12, 13]}
    assert got["赵六"] == [{"code": "empty", "why": "header_only"}]
    assert got["孙八"] == [{"code": "empty", "why": "no_fill"}]
    assert got["周九"] == [{"code": "unreadable", "file": "周九.et"}]
    assert got["吴十"] == [{"code": "header_mismatch", "sheet": "走访"}]
    assert got["郑一"] == [{"code": "old_version", "file": "郑一.xlsx", "newest": "郑一(2).xlsx"}]
    assert got["钱二"] == [{"code": "same_as", "file": "钱二.xlsx", "other": "张三 - 副本.xlsx"}]
    assert got["陈七"] == [{"code": "blank", "col": "走访结果", "count": 1, "rows": [3], "sheet": "回访"}]


def test_issue_codes_changes_and_rows_mode(tmp_path):
    got = _issues("changes", tmp_path / "a")
    assert got["张三"] == [{"code": "rows_lost", "count": 1}]
    assert got["李四"] == [{"code": "rows_changed", "count": 2, "rows": [3, 4], "cols": ["客户名称", "客户经理"]}]
    assert got["赵六"] == [{"code": "rows_added", "count": 2}]
    assert got["钱八"] == [{"code": "col_deleted", "cols": ["客户名称"]}]
    assert got["郑十"] == [{"code": "no_col", "col": "走访结果", "count": 4}]
    assert got["王五"] == []
    got = _issues("rows_mode", tmp_path / "b")
    assert got["中区"] == [{"code": "not_in_total"}] and got["南区"] == [{"code": "empty", "why": "no_fill"}]


def test_every_code_has_text_and_status():
    sample = {"col": "列", "count": 2, "rows": [3, 4], "cols": ["甲"], "sheet": "表", "file": "a.xlsx",
              "why": "原因", "newest": "b.xlsx", "other": "c.xlsx"}
    for code in collect.ISSUE_STATUS:
        assert collect.issue_text(dict(sample, code=code)), code
    assert set(collect.ISSUE_TEXT) == set(collect.ISSUE_STATUS) == set(collect.ISSUE_NAG)
    assert not collect.ISSUE_NAG["same_as"] and collect.ISSUE_NAG["blank"]


def _round_book(path, blank_rows, secret=None):
    rows = _full(6)
    for r in blank_rows:
        rows[r - 2][2] = None
    if secret:
        rows[0][4] = secret
    _book(path, rows)


def test_result_json_and_round_compare(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    secret = "不该出现的单元格内容XYZ"
    _round_book(src / "东城.xlsx", [2, 3, 4, 5, 6], secret)
    _round_book(src / "西城.xlsx", [4])
    cfg = {"input_path": str(src), "collect_expect": "东城、西城、南城", "collect_required": "走访日期"}
    logs = []
    out = collect.check_collect(cfg, logs.append)
    summary = json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])
    assert summary["result"] == os.path.join(out, collect.RESULT_NAME)
    with open(summary["result"], encoding="utf-8") as f:
        text = f.read()
    data = json.loads(text)
    assert data["version"] == 1 and data["counts"]["partial"] == 2
    east = next(p for p in data["people"] if p["name"] == "东城")
    assert east["issues"] == [{"code": "blank", "col": "走访日期", "count": 5, "rows": [2, 3, 4, 5, 6]}]
    with open(os.path.join(out, collect.ROUNDS_NAME), encoding="utf-8") as f:
        rounds_text = f.read()
    assert secret not in text and secret not in rounds_text and "客户1" not in text
    assert json.loads(rounds_text)["rounds"][-1]["issues"] == {"东城": {"blank": 5}, "西城": {"blank": 1}}

    _round_book(src / "东城.xlsx", [3], secret)          # 第二轮：东城还差 1 行，西城改好了，南城新交
    _round_book(src / "西城.xlsx", [])
    _round_book(src / "南城.xlsx", [])
    collect.check_collect(cfg, lambda m: None)
    with open(os.path.join(out, collect.NOTICE_NAME), encoding="utf-8-sig") as f:
        notice = f.read()
    assert "比上次：\n  新交：南城\n  改好了：西城\n  还差：东城 漏填 5 格 → 1 格\n" in notice


def test_old_rounds_without_issues_only_show_new(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _round_book(src / "东城.xlsx", [3])
    _round_book(src / "西城.xlsx", [])
    out = src / collect.RESULT_DIR
    out.mkdir()
    (out / collect.ROUNDS_NAME).write_text(json.dumps({"rounds": [{"time": "x", "submitted": ["西城"], "ok": []}]},
                                                      ensure_ascii=False), encoding="utf-8")
    collect.check_collect({"input_path": str(src), "collect_expect": "东城、西城", "collect_required": "走访日期"},
                          lambda m: None)
    with open(out / collect.NOTICE_NAME, encoding="utf-8-sig") as f:
        notice = f.read()
    assert "比上次新交：东城。" in notice and "改好了" not in notice and "还差" not in notice
