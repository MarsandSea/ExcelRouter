"""core/desens.py：识别、各处理方式、文字里夹带的号码、复检，以及接进拆分 / 单独脱敏 / 汇总后的脱敏版。

样本里的号码全是虚构的（身份证号按规则算出校验码，银行卡号按 Luhn 补最后一位）。
"""
import glob
import json
import os
import zipfile

import openpyxl
from openpyxl.comments import Comment

from core import desens as D
from core.collect import merge_tables
from core.splitter import run_split


def _idn(s17):
    return s17 + D._id_check_digit(s17)


def _card(prefix):
    return next(prefix + str(x) for x in range(10) if D._luhn(prefix + str(x)))


ID1, ID2, ID3 = _idn("11010119800101001"), _idn("44200019851212003"), _idn("44200020010101004")
CARD1, CARD2 = _card("622202000000000000"), _card("621700000000000012")
PHONES = (13512340001, "139 0000 5678", "+8615000009999")


def _summary(logs):
    return json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])


def _book(path, rows=None, hidden_col=None, title="九月客户维系清单"):
    """第 1 行大标题、第 2 行表头。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "清单"
    ws["A1"] = title
    head = ["网格", "客户姓名", "联系电话", "证件号码", "银行卡号", "装机地址", "备注"]
    ws.append(head)
    rows = rows or [
        ["东区", "张三", PHONES[0], ID1, CARD1, "广东省中山市东区街道中山路 1 号", f"客户说新号 13600001111，证件 {ID3}"],
        ["东区", "欧阳娜娜", PHONES[1], ID2, CARD2, "广东省中山市石岐区康华路 12 号", "已回访"],
        ["港口", "李四", PHONES[2], ID3, CARD1, "广东省中山市港口镇民主路 3 号", "无"],
    ]
    for r in rows:
        ws.append(r)
    if hidden_col:
        ws.column_dimensions[hidden_col].hidden = True
    wb.save(path)


def _values(path):
    wb = openpyxl.load_workbook(path)
    return [[c for c in r] for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]


def _flat(path):
    return "\n".join(str(v) for r in _values(path) for v in r if v is not None)


# ---------- 校验 ----------

def test_validators():
    assert D.is_id_card(ID1) and not D.is_id_card(ID1[:17] + ("0" if ID1[17] != "0" else "1"))
    assert not D.is_id_card("110105199013070021")            # 13 月
    assert D._is_id15("110105900307002") and not D._is_id15("110105901307002")
    assert D.is_mobile("13800001234") and D.is_mobile("+8613800001234") and D.is_mobile("8613800001234")
    assert not D.is_mobile("12800001234") and not D.is_mobile("1380000123")
    assert D.is_bank_card(CARD1) and not D.is_bank_card(CARD1[:-1] + str((int(CARD1[-1]) + 1) % 10))
    assert not D.is_bank_card(ID1)                            # 身份证号不算卡号
    assert D.is_iccid("89860012345678901234") and not D.is_bank_card("8986001234567890123")
    assert D.is_landline("076088881234") and not D.is_landline("13800001234")
    assert D.is_name("张三") and D.is_name("阿卜杜·热合曼") and not D.is_name("中国移动通信集团")


def test_detect_column_by_name_and_values():
    assert D.detect_column("证件号码", [ID1, ID2]) == "身份证号"
    assert D.detect_column("随便什么", [ID1, ID2, ID3]) == "身份证号"           # 列名不认识，值能认
    assert D.detect_column("备注", [13800001234] * 4) == "手机号"
    assert D.detect_column("联系电话", ["076088881234", "0760-88885678"]) == "固定电话"
    assert D.detect_column("联系人", ["张三", "李四"]) == "姓名"
    assert D.detect_column("联系人", ["13800001234", "13900001234"]) == "手机号"
    assert D.detect_column("客户名称", ["中山市某某贸易有限公司"] * 3) is None    # 公司名不是人名
    assert D.detect_column("手机号码归属地", ["广东", "广东"]) is None            # 说的是号码的属性
    assert D.detect_column("证件类型", ["身份证"]) is None
    assert D.detect_column("宽带账号", []) == "账号编码"
    assert D.detect_column("金额", [100, 200]) is None


# ---------- 处理方式 ----------

def test_mask_rules():
    assert D.apply_value(13800001234, "手机号", "mask") == "138****1234"
    assert D.apply_value("+86 138-0000-1234", "手机号", "mask") == "138****1234"
    assert D.apply_value(ID1, "身份证号", "mask") == ID1[:6] + "*" * 8 + ID1[-4:]
    assert D.apply_value(CARD1, "银行卡号", "mask") == CARD1[:6] + "*" * (len(CARD1) - 10) + CARD1[-4:]
    assert D.apply_value("张三", "姓名", "mask") == "张*"
    assert D.apply_value("欧阳娜娜", "姓名", "mask") == "欧阳**"
    assert D.apply_value("广东省中山市东区街道中山路 1 号", "地址", "mask") == "广东省中山市东区街道****"
    assert D.apply_value("广东省中山市石岐区康华路 12 号", "地址", "generalize") == "广东省中山市石岐区"
    assert D.apply_value("0760-88881234", "固定电话", "mask") == "0760-****1234"
    assert D.apply_value("zhangsan@example.com", "邮箱", "mask") == "z***@example.com"
    assert D.apply_value(113.392411, "经纬度", "generalize") == 113.39
    assert D.apply_value(5234, "自定义", "generalize") == "5000-6000"
    assert D.apply_value("无", "身份证号", "mask") == "无"                     # 占位不动
    assert D.apply_value(ID1, "身份证号", "clear") is None
    assert D.apply_value(13800001234, "手机号", "generalize") == "138****1234"  # 号码泛化不了，退回打码


def test_pseudo_codes_consistent_per_seed():
    a = D.apply_value(13800001234, "手机号", "pseudo", b"k1")
    assert a == D.apply_value("138 0000 1234", "手机号", "pseudo", b"k1")      # 写法不同，同一个号同一个代号
    assert a != D.apply_value(13800001234, "手机号", "pseudo", b"k2")           # 换种子就不同
    assert a.startswith("TEL-") and len(a) == 14
    assert D.apply_value("张三", "姓名", "pseudo", b"k1") != D.apply_value("张三", "地址", "pseudo", b"k1")


def test_mask_text_and_find_kinds():
    s, n = D.mask_text(f"新号+8613600001111，证件{ID3}，卡 {CARD1}，电话 139-0000-5678")
    assert n == 4 and "13600001111" not in s and ID3 not in s and CARD1 not in s and "0000-5678" not in s
    assert D.find_kinds(s) == set()
    assert D.find_kinds(f"ID{ID1}") == {"身份证号"}
    assert D.find_kinds(13800001234) == {"手机号"} and D.find_kinds("订单 20261018000012345678") == set()
    s, n = D.mask_text(f"证件{ID3} 电话13600001111", {"身份证号"})               # 只遮指定的类型
    assert n == 1 and "13600001111" in s


def test_from_config():
    assert D.Job.from_config(None) is None and D.Job.from_config("") is None
    assert D.Job.from_config({"scene": "不存在"}) is None
    assert D.Job.from_config("发群通报").scene == "发群通报"
    j = D.Job.from_config({"columns": {"备注": "clear"}})
    assert j.scene == "自定义" and j.columns == {"备注": "clear"}
    assert D.Job.from_config({"scene": "交给AI分析", "seed": "00ff"}).fixed_seed


# ---------- 接进拆分 ----------

def _split(tmp_path, desens, **kw):
    src = tmp_path / "清单.xlsx"
    _book(src, **kw.pop("book", {}))
    out = tmp_path / "out"
    logs = []
    cfg = dict(input_path=str(src), output_path=str(out), split_column=kw.pop("by", "网格"), desens=desens,
               make_zip=True, **kw)
    res = run_split(cfg, log_fn=logs.append)
    return res, logs


def test_split_internal_scene_masks_id_and_card_keeps_phone(tmp_path):
    res, logs = _split(tmp_path, "内部下发")
    s = _summary(logs)
    assert s["groups"] == 2 and s["desens"]["ok"] and s["desens"]["scene"] == "内部下发"
    text = _flat(os.path.join(res, "东区.xlsx"))
    assert ID1 not in text and ID2 not in text and CARD1 not in text and CARD2 not in text
    assert ID3 not in text                                      # 备注里夹带的证件号也遮了
    assert "13600001111" in text and "13512340001" in text      # 号码照发（外呼要用）
    assert "张三" in text and "九月客户维系清单" in text          # 姓名照发，标题行不动
    rep = openpyxl.load_workbook(os.path.join(res, D.REPORT_NAME))
    assert rep.sheetnames == ["概要", "处理明细", "复检", "指纹"]
    det = list(rep["处理明细"].iter_rows(min_row=2, values_only=True))
    assert {(d[2], d[3], d[4]) for d in det} >= {("证件号码", "身份证号", "打码"), ("联系电话", "手机号", "照发")}
    report_text = "\n".join(str(v) for ws in rep.worksheets for r in ws.iter_rows(values_only=True) for v in r if v)
    assert all(x not in report_text for x in (ID1, ID2, CARD1, "13512340001", "欧阳娜娜"))   # 报告里没有真实值


def test_split_groups_by_real_values_then_masks(tmp_path):
    """按要打码的列拆：分组用真实值，文件里是打码后的值。"""
    res, logs = _split(tmp_path, "发群通报", by="客户姓名")
    s = _summary(logs)
    assert s["groups"] == 3 and s["desens"]["ok"]
    text = _flat(os.path.join(res, "欧阳娜娜.xlsx"))
    assert "欧阳**" in text and "欧阳娜娜" not in text
    assert "139****5678" in text and "13600001111" not in _flat(os.path.join(res, "张三.xlsx"))
    assert "广东省中山市石岐区****" in text


def test_split_hidden_column_still_cleared(tmp_path):
    res, logs = _split(tmp_path, "内部下发", book={"hidden_col": "E"})     # 银行卡号列隐藏
    assert _summary(logs)["desens"]["ok"]
    assert CARD1 not in _flat(os.path.join(res, "东区.xlsx"))


def test_split_without_desens_unchanged(tmp_path):
    res, logs = _split(tmp_path, None)
    assert "desens" not in _summary(logs)
    assert ID1 in _flat(os.path.join(res, "东区.xlsx"))
    assert not os.path.exists(os.path.join(res, D.REPORT_NAME))


def test_split_residual_blocks_zip(tmp_path):
    """不处理文字里夹带的号码时，备注里的完整号码会被复检抓到：不打 ZIP、写醒目说明、ok=False。"""
    src = tmp_path / "in"
    src.mkdir()
    for i, g in enumerate(("东区", "港口")):
        _book(src / f"{g}.xlsx", rows=[[g, "张三", PHONES[0], ID1, CARD1, "某地址某某路", f"联系 1390000{i}234"]])
    out = tmp_path / "out"
    logs = []
    res = run_split(dict(input_path=str(src), output_path=str(out), split_column="网格", make_zip=True,
                         desens={"scene": "发群通报", "notes": False}), log_fn=logs.append)
    s = _summary(logs)
    assert not s["desens"]["ok"] and s["desens"]["residual"] == 2 and s["zips"] == 0
    assert not glob.glob(os.path.join(res, "*.zip"))
    note = open(os.path.join(res, D.FAIL_NOTE), encoding="utf-8-sig").read()
    assert "备注" in note and "13900000234" not in note           # 只列位置，不列内容


def test_split_folder_zip_when_clean(tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    _book(src / "a.xlsx")
    logs = []
    res = run_split(dict(input_path=str(src), output_path=str(tmp_path / "out"), split_column="网格",
                         make_zip=True, merge_across_files=False, desens="发群通报"), log_fn=logs.append)
    s = _summary(logs)
    assert s["desens"]["ok"] and s["zips"] == 2
    with zipfile.ZipFile(os.path.join(res, "东区.zip")) as z:
        assert D.REPORT_NAME not in "".join(z.namelist())          # 报告在结果根目录，不进分组包


def test_keep_formulas_ignored_when_desens(tmp_path):
    _res, logs = _split(tmp_path, "内部下发", keep_formulas=True, preserve_format=True)
    assert any("保留公式」本次不生效" in m for m in logs)


# ---------- 单独脱敏 ----------

def test_run_desens_standalone(tmp_path):
    src = tmp_path / "表"
    src.mkdir()
    _book(src / "清单.xlsx")
    wb = openpyxl.load_workbook(src / "清单.xlsx")
    wb.active["B3"].comment = Comment(f"本人证件 {ID1}", "x")
    hid = wb.create_sheet("底稿")
    hid.append(["证件号码"])
    hid.append([ID1])
    hid.sheet_state = "hidden"
    wb.save(src / "清单.xlsx")
    logs = []
    out = D.run_desens({"input_path": str(src), "desens": "交给AI分析"}, log_fn=logs.append)
    assert os.path.dirname(out) == os.path.join(str(src), D.RESULT_DIR)
    s = _summary(logs)
    assert s["mode"] == "desens" and s["files"] == 1 and s["ok"]
    dst = os.path.join(out, "清单_脱敏.xlsx")
    wb2 = openpyxl.load_workbook(dst)
    assert wb2.sheetnames == ["清单"]                              # 隐藏 sheet 不带
    assert all(c.comment is None for r in wb2.active.iter_rows() for c in r)
    text = _flat(dst)
    assert ID1 not in text and "张三" not in text and "NAM-" in text and "IDN-" in text
    assert "广东省中山市东区" in text and "街道中山路" not in text   # 地址泛化到区县
    assert os.path.exists(os.path.join(out, D.REPORT_NAME))
    # 再跑一次：上次的「脱敏结果」不会被当成输入
    logs2 = []
    D.run_desens({"input_path": str(src), "desens": "交给AI分析"}, log_fn=logs2.append)
    assert _summary(logs2)["files"] == 1


def test_run_desens_blocked_file_not_written(tmp_path):
    src = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["序号", "说明"])
    wb.active.append([1, "联系 13800001234"])
    wb.save(src)
    logs = []
    out = D.run_desens({"input_path": str(src), "desens": {"scene": "发群通报", "notes": False}},
                       log_fn=logs.append)
    s = _summary(logs)
    assert s["files"] == 0 and s["blocked"] == 1 and not s["ok"]
    assert not os.path.exists(os.path.join(out, "x_脱敏.xlsx"))
    assert os.path.exists(os.path.join(out, D.FAIL_NOTE))


# ---------- 汇总后另存脱敏版 ----------

def test_merge_writes_desens_copy(tmp_path):
    src = tmp_path / "收回"
    src.mkdir()
    _book(src / "东区.xlsx")
    _book(src / "港口.xlsx")
    logs = []
    out = merge_tables({"input_path": str(src), "desens": "对外报送"}, log_fn=logs.append)
    s = _summary(logs)
    assert s["desens"]["ok"] and s["desens_output"].endswith("_脱敏版.xlsx")
    assert ID1 in _flat(out)                                       # 汇总表本身不动
    text = _flat(s["desens_output"])
    assert ID1 not in text and "13512340001" not in text and "张三" not in text
