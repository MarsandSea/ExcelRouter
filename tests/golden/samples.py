"""对比测试的样本与场景。样本全部按固定内容用 openpyxl 生成（.xls 用 fixtures 里提交的测试件），每次生成结果一样。

这里只放「新旧两条读表路径应当逐格一致」的场景。有意改变行为的情况（「NA」类文字、以「=」开头的文字、
跨文件列顺序不同、默认不发隐藏内容、无归属行单独成文件）由各自的专门测试断言，不放进来。

Copyright (c) 2026 AbeLin · MIT License
"""
import datetime
import os
import shutil

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
D = datetime.datetime

HEAD = ["工号", "姓名", "部门", "区域", "金额", "日期", "备注"]
TITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
HOT_FILL = PatternFill("solid", fgColor="FFF2CC")
WIDTHS = {"A": 9, "B": 10, "C": 12, "D": 8, "E": 12, "F": 13, "G": 16}


def _detail(ws, title, rows, tail_styled=0):
    """大标题（合并）+ 第 2 行表头 + 数据行；rows 里的 None 是一整行空行。"""
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(HEAD)
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(row=2, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(list(row) if row else [])
        if not row:
            continue
        r = ws.max_row
        ws.cell(row=r, column=5).number_format = "#,##0.00"
        ws.cell(row=r, column=6).number_format = "yyyy-mm-dd"
        amount = row[4]
        if isinstance(amount, (int, float)) and amount >= 1000:
            ws.cell(row=r, column=5).fill = HOT_FILL
        if row[6]:
            ws.cell(row=r, column=7).font = Font(italic=True, color="C00000")
    for _ in range(tail_styled):            # 尾部只有格式、没有内容的行
        r = ws.max_row + 1
        for c in range(1, len(HEAD) + 1):
            ws.cell(row=r, column=c).fill = HOT_FILL
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w


def build_dept(root):
    """两个同模板的表：A 有三个 sheet（两个明细 + 一个没有「部门」列的），B 带错误格和尾部带格式的空行。"""
    folder = os.path.join(root, "部门表")
    os.makedirs(folder, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    _detail(ws, "2026年10月 销售明细", [
        [1001, "张三", "销售部", "北京", 1200.5, D(2026, 10, 1), "首单"],
        [1002, "李四", "技术部", "上海", 800, D(2026, 10, 2), None],
        [1003, "王五", "销售部", "广州", 450, D(2026, 10, 3), "复购"],
        None,
        [1004, "赵六", "市场部", "北京", 300, D(2026, 10, 4), "活动"],
        [1005, "钱七", None, "深圳", 99, D(2026, 10, 5), "没写部门"],
        [1006, "孙八", "销售一部", "北京", 650, D(2026, 10, 6), None],
        [1007, "周九", "技术部", "上海", 1500, D(2026, 10, 7), "大单"],
        [None, "合计", None, None, 4999.5, None, None],
    ])
    ws2 = wb.create_sheet("按区域")
    ws2.append(["区域", "金额"])
    ws2.append(["北京", 2150.5])
    ws2.append(["上海", 2300])
    ws3 = wb.create_sheet("二月")
    _detail(ws3, "2026年2月 销售明细", [
        [1101, "郑一", "技术部", "北京", 120, D(2026, 2, 1), None],
        [1102, "王二", "销售部", "上海", 330, D(2026, 2, 2), "补录"],
    ])
    wb.save(os.path.join(folder, "A.xlsx"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    _detail(ws, "2026年10月 销售明细（二）", [
        [2001, "冯一", "销售部", "广州", 880, D(2026, 10, 9), None],
        [2002, "陈二", "技术部", "深圳", 1020, D(2026, 10, 10), "加急"],
        [2003, "褚三", "#N/A", "北京", 50, D(2026, 10, 11), None],
        [2004, "卫四", "市场部", "上海", "#DIV/0!", D(2026, 10, 12), None],
        [2005, "蒋五", "销售一部", "杭州", 410, D(2026, 10, 13), None],
        [2006, "张三", "销售部", "北京", 77, D(2026, 10, 14), "同名"],
    ], tail_styled=3)
    wb.save(os.path.join(folder, "B.xlsx"))
    return folder


def build_shift(root):
    """表头在第 3 行；日期、班组（整数）可做拆分列；金额是同行公式，最后一行是 SUM 合计。"""
    path = os.path.join(root, "排班.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "排班"
    ws.append(["2026年10月排班表"])
    ws.merge_cells("A1:F1")
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["制表：测试"])
    ws.append(["日期", "班组", "姓名", "工时", "单价", "金额"])
    for c in range(1, 7):
        ws.cell(row=3, column=c).font = Font(bold=True)
        ws.cell(row=3, column=c).fill = TITLE_FILL
    data = [
        (D(2026, 10, 1), 1, "张三", 8, 25.5),
        (D(2026, 10, 1), 2, "李四", 6, 30),
        (D(2026, 10, 2), 1, "王五", 8, 25.5),
        (D(2026, 10, 2), 3, "赵六", 4, 40),
        (D(2026, 10, 3), 2, "钱七", 7.5, 30),
        (D(2026, 10, 3), 3, "孙八", 8, 40),
    ]
    for i, (d, team, name, hours, price) in enumerate(data):
        r = 4 + i
        ws.append([d, team, name, hours, price, f"=D{r}*E{r}"])
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=6).number_format = "0.00"
    last = 3 + len(data)
    ws.append(["合计", None, None, f"=SUM(D4:D{last})", None, f"=SUM(F4:F{last})"])
    ws.column_dimensions["A"].width = 12
    wb.save(path)
    return path


def build_xls(root):
    path = os.path.join(root, "月报.xls")
    shutil.copyfile(os.path.join(FIXTURES, "monthly.xls"), path)
    return path


def build_collision(root):
    """两个取值净化成文件名后同名：「销售/部」「销售:部」都变成「销售_部」。"""
    path = os.path.join(root, "撞名.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["编号", "部门", "金额"])
    for row in (["1", "销售/部", 10], ["2", "销售:部", 20], ["3", "销售/部", 30]):
        ws.append(row)
    wb.save(path)
    return path


def build_hidden(root):
    """一个可见 sheet（备注列隐藏、一行隐藏）+ 一个隐藏的 sheet。对比时用 keep_hidden=True 跑，应与旧版一致。"""
    path = os.path.join(root, "隐藏.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    _detail(ws, "带隐藏内容的表", [
        ["H001", "张三", "销售部", "北京", 100, D(2026, 10, 1), "内部备注一"],
        ["H002", "李四", "技术部", "上海", 200, D(2026, 10, 2), "内部备注二"],
        ["H003", "王五", "销售部", "广州", 300, D(2026, 10, 3), None],
    ])
    ws.column_dimensions["G"].hidden = True
    ws.row_dimensions[4].hidden = True
    ws2 = wb.create_sheet("底稿")
    _detail(ws2, "底稿", [
        ["H101", "赵六", "技术部", "北京", 400, D(2026, 10, 4), None],
    ])
    ws2.sheet_state = "hidden"
    wb.save(path)
    return path


BUILDERS = {"dept": build_dept, "shift": build_shift, "xls": build_xls,
            "collision": build_collision, "hidden": build_hidden}

BASE = {
    "header_mode": "auto", "header_row": 1, "grid_keys": [], "id_keys": [],
    "split_column": "部门", "selected_values": [], "person_column": "", "to_person": False,
    "person_file_filter": [], "value_alias_map": {}, "skip_values": ["合计", "小计", "总计", "平均", ""],
    "merge_across_files": False, "make_zip": True, "exact_match": True, "preserve_format": True,
    "keep_formulas": False, "auto_open_output": False,
}

# (场景名, 样本, 是否只取样本里的 A.xlsx, 配置覆盖)
SCENARIOS = [
    ("dept_dir", "dept", False, {}),
    ("dept_dir_merge", "dept", False, {"merge_across_files": True}),
    ("dept_dir_merge_fast", "dept", False, {"merge_across_files": True, "preserve_format": False}),
    ("dept_person_merge", "dept", False, {"merge_across_files": True, "to_person": True, "person_column": "姓名"}),
    ("dept_person_fast", "dept", False, {"to_person": True, "person_column": "姓名", "preserve_format": False}),
    ("dept_person_filter", "dept", False, {"to_person": True, "person_column": "姓名", "person_file_filter": ["A"]}),
    ("dept_file", "dept", True, {}),
    ("dept_selected", "dept", False, {"merge_across_files": True, "selected_values": ["销售部", "技术部"]}),
    ("dept_contains", "dept", False, {"merge_across_files": True, "selected_values": ["销售", "售部"],
                                      "exact_match": False}),
    ("dept_alias", "dept", False, {"merge_across_files": True,
                                   "value_alias_map": {"销售": ["销售部", "销售一部"]}}),
    ("dept_header_row", "dept", False, {"header_mode": "row", "header_row": 2}),
    ("shift_date", "shift", False, {"split_column": "日期"}),
    ("shift_team_fast", "shift", False, {"split_column": "班组", "preserve_format": False}),
    ("shift_formulas", "shift", False, {"split_column": "班组", "keep_formulas": True}),
    ("xls_file", "xls", False, {}),
    ("collision", "collision", False, {}),
    ("hidden_keep", "hidden", False, {"keep_hidden": True}),
]
NAMES = [s[0] for s in SCENARIOS]


def make_config(name, root):
    """在 root 下生成场景 name 的输入，返回 run_split 用的配置（输出放 root/out）。"""
    _, sample, only_a, over = next(s for s in SCENARIOS if s[0] == name)
    src_root = os.path.join(root, "in")
    os.makedirs(src_root, exist_ok=True)
    inp = BUILDERS[sample](src_root)
    if only_a:
        inp = os.path.join(inp, "A.xlsx")
    out = os.path.join(root, "out")
    os.makedirs(out, exist_ok=True)
    cfg = dict(BASE, input_path=inp, output_path=out)
    cfg.update(over)
    return cfg
