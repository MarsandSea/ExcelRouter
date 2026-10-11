"""填报规则（v3.0）：「让对方填写的列」括号里写要求 → 下发时加数据验证、查收时查格式、汇总时规整写法。

Copyright (c) 2026 AbeLin · MIT License
"""
import datetime
import hashlib
import json
import os

import openpyxl

from core import collect
from core.collect import ST_INVALID, ST_OK, _coerce, parse_fill_spec

D = datetime.date(2026, 10, 11)


def test_parse_fill_spec_forms():
    logs = []
    got = dict(parse_fill_spec("走访日期（日期）、走访结果(已走访/未走访/拒访)，户数（整数 选填） 备注（选填）"
                               "、电话（手机号码）、首访（必填）、其它（YYYY-MM-DD）", logs.append))
    assert got["走访日期"] == {"type": "date", "optional": False}
    assert got["走访结果"] == {"type": "list", "options": ["已走访", "未走访", "拒访"], "optional": False}
    assert got["户数"] == {"type": "whole", "optional": True}
    assert got["备注"] == {"type": "text", "optional": True}
    assert got["电话"]["type"] == "mobile" and got["首访"] == {"type": "text", "optional": False}
    assert got["其它"]["type"] == "text" and any("YYYY-MM-DD" in m for m in logs)
    both = parse_fill_spec("结果（日期 是/否）", logs.append)
    assert both[0][1]["type"] == "list" and any("按选项" in m for m in logs)


def test_parse_without_brackets_same_as_split_names():
    for text in ("走访日期 走访结果", "走访日期、走访结果，备注", ["走访日期", "走访结果"], ""):
        assert [n for n, _r in parse_fill_spec(text)] == collect.split_names(text)
        assert all(r == {"type": "text", "optional": False} for _n, r in parse_fill_spec(text))


def test_coerce_numbers_and_options():
    num, whole = {"type": "number"}, {"type": "whole"}
    assert _coerce("1,200", num) == (True, 1200) and _coerce("3户", num) == (True, 3)
    assert _coerce("2.5", num) == (True, 2.5) and _coerce(3.0, whole) == (True, 3)
    assert not _coerce("约5", num)[0] and not _coerce("2.5", whole)[0] and not _coerce("≥3", num)[0]
    opt = {"type": "list", "options": ["已走访", "未走访"]}
    assert _coerce(" 已走访 ", opt) == (True, "已走访")
    assert _coerce("ＡＢ", {"type": "list", "options": ["ab"]}) == (True, "ab")       # 全角、大小写
    assert not _coerce("已访", opt)[0]                                                # 不猜同义词


def test_coerce_phone():
    mob, tel = {"type": "mobile"}, {"type": "phone"}
    assert _coerce("+86 138-0013-8000", mob) == (True, "13800138000")
    assert _coerce(13800138000.0, mob) == (True, "13800138000")
    assert not _coerce("1380013800", mob)[0] and not _coerce("23800138000", mob)[0]
    assert not _coerce("13800138000/13900139000", mob)[0]
    assert not _coerce("010-12345678转123", mob)[0]
    assert _coerce("010-12345678转123", tel) == (True, "010-12345678转123")
    assert _coerce("(010)12345678", tel) == (True, "010-12345678")
    assert _coerce("0755 12345678", tel) == (True, "0755-12345678")
    assert _coerce("01012345678", tel) == (True, "010-12345678")
    assert _coerce("13800138000", tel) == (True, "13800138000")


def test_coerce_dates():
    d = {"type": "date"}
    assert _coerce("2026.10.5", d, D) == (True, datetime.datetime(2026, 10, 5))
    assert _coerce("2026年10月5日", d, D)[1] == datetime.datetime(2026, 10, 5)
    assert _coerce("20261005", d, D)[1] == datetime.datetime(2026, 10, 5)
    assert _coerce("10月5日", d, D)[1] == datetime.datetime(2026, 10, 5)
    assert _coerce("1月3日", d, datetime.date(2026, 12, 20))[1] == datetime.datetime(2027, 1, 3)   # 跨年
    assert _coerce(46300, d, D)[1] == datetime.datetime(1899, 12, 30) + datetime.timedelta(days=46300)
    assert not _coerce(10.1, d, D)[0] and not _coerce(5, d, D)[0]          # 有歧义的数字、太小的序号
    assert not _coerce("国庆后", d, D)[0] and not _coerce("2026-13-01", d, D)[0]
    assert _coerce(datetime.date(2026, 1, 2), d, D) == (True, datetime.datetime(2026, 1, 2))
    assert not _coerce(datetime.datetime(1900, 1, 10, 2, 24), d, D)[0]       # 日期格式的格子里填了 10.1


# ---------- 往返：拆分下发 → 模拟填写（直接写，等于粘贴绕过验证）→ 查收 → 汇总 ----------

SPEC = "走访日期（日期）、走访结果（已走访/未走访/拒访）、户数（整数）、联系电话（手机号）、备注（选填）"


def _issue(tmp_path, spec=SPEC):
    from core.splitter import run_split
    src = tmp_path / "九月走访.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "走访"
    ws.append(["网格", "客户名称"])
    for g in ("东区", "西区", "南区"):
        for i in range(1, 4):
            ws.append([g, f"{g}客户{i}"])
    wb.save(src)
    out = run_split({"input_path": str(src), "output_path": str(tmp_path / "out"), "split_column": "网格",
                     "fill_columns": spec, "due": "10月15日", "make_zip": False}, lambda m: None)
    return out


def _fill(out, group, dst, values):
    wb = openpyxl.load_workbook(os.path.join(out, f"{group}.xlsx"))
    ws = wb.worksheets[0]
    for r, row in enumerate(values, start=2):
        for c, v in enumerate(row, start=3):
            ws.cell(row=r, column=c).value = v
    wb.save(dst)


def test_issue_writes_validations_and_record(tmp_path):
    out = _issue(tmp_path)
    wb = openpyxl.load_workbook(os.path.join(out, "东区.xlsx"))
    ws = wb.worksheets[0]
    assert [c.value for c in ws[1]] == ["网格", "客户名称", "走访日期", "走访结果", "户数", "联系电话", "备注"]
    dvs = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    assert dvs["C2:C4"].type == "date" and dvs["D2:D4"].formula1 == '"已走访,未走访,拒访"'
    assert dvs["E2:E4"].type == "whole" and dvs["F2:F4"].type == "custom"
    assert all(dv.showErrorMessage for dv in dvs.values()) and "G2:G4" not in dvs
    assert ws["C2"].number_format == "yyyy-mm-dd" and ws["F3"].number_format == "@"
    guide = wb[collect.GUIDE_SHEET]
    rows = [[c.value for c in r] for r in guide.iter_rows()]
    assert ["走访结果", "必填 · 从下拉里选", "已走访 / 未走访 / 拒访"] in rows and ["备注", "选填", None] in rows
    assert guide.protection.sheet
    with open(os.path.join(out, collect.RECORD_NAME), encoding="utf-8") as f:
        rec = json.load(f)
    assert rec["fill_columns"] == ["走访日期", "走访结果", "户数", "联系电话", "备注"]
    assert rec["required"] == ["走访日期", "走访结果", "户数", "联系电话"]
    assert rec["rules"]["走访结果"] == {"type": "list", "options": ["已走访", "未走访", "拒访"], "optional": False}
    assert "备注" not in rec["rules"]


def test_no_brackets_record_unchanged(tmp_path):
    out = _issue(tmp_path, "走访日期、走访结果")
    with open(os.path.join(out, collect.RECORD_NAME), encoding="utf-8") as f:
        rec = json.load(f)
    assert "rules" not in rec and "required" not in rec
    ws = openpyxl.load_workbook(os.path.join(out, "东区.xlsx")).worksheets[0]
    assert not ws.data_validations.dataValidation
    assert ws["C2"].number_format == "General"


def test_long_options_skip_dropdown_but_still_checked(tmp_path):
    opts = "/".join(f"选项{i:03d}很长很长很长" for i in range(30))
    out = _issue(tmp_path, f"走访结果（{opts}）")
    ws = openpyxl.load_workbook(os.path.join(out, "东区.xlsx")).worksheets[0]
    assert not ws.data_validations.dataValidation
    with open(os.path.join(out, collect.RECORD_NAME), encoding="utf-8") as f:
        assert len(json.load(f)["rules"]["走访结果"]["options"]) == 30
    back = tmp_path / "回收"
    back.mkdir()
    _fill(out, "东区", back / "东区.xlsx", [["选项001很长很长很长"], ["随便写的"], ["选项029很长很长很长"]])
    _ctx, _res, report = collect._inspect({"input_path": str(back),
                                           "collect_record": os.path.join(out, collect.RECORD_NAME)},
                                          lambda m: None, None, None)
    east = next(p for p in report["people"] if p["name"] == "东区")
    assert east["status"] == ST_INVALID
    assert east["invalid"] == [{"col": "走访结果", "code": "not_in_options", "count": 1, "rows": [3]}]
    assert "（请从下拉里选）" in collect.notice_text(report, "九月走访")     # 选项多于 6 个时不逐个列出


def test_check_finds_invalid_without_echoing_values(tmp_path):
    out = _issue(tmp_path)
    back = tmp_path / "回收"
    back.mkdir()
    secret = "国庆后某天"
    _fill(out, "东区", back / "东区.xlsx", [["2026.10.5", "已走访", "3户", "+86 138 0013 8000", None],
                                          ["10月6日", "未走访", 2, "13900139000", "x"],
                                          [46300, "拒访", "1", "13700137000", None]])
    _fill(out, "西区", back / "西区.xlsx", [[secret, "已访", "约5", "1380013800", None],
                                          ["2026-10-07", "已走访", 2.5, "13800138000", None],
                                          ["2026-10-08", "已走访", 1, "13800138000", None]])
    cfg = {"input_path": str(back), "collect_record": os.path.join(out, collect.RECORD_NAME)}
    logs = []
    res_dir = collect.check_collect(cfg, logs.append)
    summary = json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])
    assert summary["invalid"] == 1 and summary["partial"] == 0          # 备注选填，空着不算漏填
    _ctx, _res, report = collect._inspect(cfg, lambda m: None, None, None)
    people = {p["name"]: p for p in report["people"]}
    assert people["东区"]["status"] == ST_OK
    west = people["西区"]
    assert west["status"] == ST_INVALID
    assert {(i["col"], i["code"]): i["rows"] for i in west["invalid"]} == {
        ("走访日期", "not_date"): [2], ("走访结果", "not_in_options"): [2], ("户数", "not_whole"): [2, 3],
        ("联系电话", "not_mobile"): [2]}
    with open(os.path.join(res_dir, collect.NOTICE_NAME), encoding="utf-8-sig") as f:
        notice = f.read()
    assert "填写有误（1）：" in notice
    assert "走访结果 第 2 行不在选项里（应为 已走访 / 未走访 / 拒访）" in notice and "户数 第 2、3 行不是整数" in notice
    texts = [notice]
    for name in (collect.RESULT_NAME, collect.ROUNDS_NAME):
        with open(os.path.join(res_dir, name), encoding="utf-8") as f:
            texts.append(f.read())
    wb = openpyxl.load_workbook(summary["progress"])
    texts += [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    blob = "\n".join(texts)
    for v in (secret, "已访", "约5", "1380013800"):
        assert v not in blob.replace("13800138000", ""), v
    assert [c.value for c in wb["问题明细"][1]] == ["应交", "交的文件", "列", "问题", "格数", "行号"]


def test_merge_normalizes_and_leaves_source_untouched(tmp_path):
    out = _issue(tmp_path)
    back = tmp_path / "回收"
    back.mkdir()
    _fill(out, "东区", back / "东区.xlsx", [["2026.10.5", " 已走访", "3户", "+86 138 0013 8000", None],
                                          ["10月6日", "未走访", 2, "13900139000", None],
                                          [10.1, "拒访", "1,200", "13700137000", None]])
    before = hashlib.sha256((back / "东区.xlsx").read_bytes()).hexdigest()
    cfg = {"input_path": str(back), "collect_record": os.path.join(out, collect.RECORD_NAME),
           "collect_title": "九月走访"}
    logs = []
    path = collect.run_collect(cfg, logs.append)
    assert hashlib.sha256((back / "东区.xlsx").read_bytes()).hexdigest() == before
    wb = openpyxl.load_workbook(path)
    ws = wb["走访"]
    col = {c.value: c.column for c in ws[1]}
    vals = [[ws.cell(row=r, column=col[n]).value for n in ("走访日期", "走访结果", "户数", "联系电话")]
            for r in range(2, 5)]
    assert vals[0] == [datetime.datetime(2026, 10, 5), "已走访", 3, "13800138000"]
    assert vals[1][0] == datetime.datetime(2026, 10, 6) and vals[2][2] == 1200
    # 日期格式的格子里填了 10.1：读出来是 1900 年的日期，分不清是哪天，不改、标浅橙
    assert vals[2][0].year == 1900 and ws.cell(row=4, column=col["走访日期"]).fill.fgColor.rgb.endswith("FFEB9C")
    rep = [[c for c in r if c is not None] for r in wb[collect.REPORT_SHEET].iter_rows(values_only=True)]
    flat = ["|".join(map(str, r)) for r in rep]
    assert any(r.startswith("自动规整") for r in flat)
    assert any(r.startswith("走访日期|2 格|例：2026.10.5 → 2026-10-05") for r in flat)
    assert any("138****8000" in r for r in flat) and not any("13800138000" in r for r in flat)
    assert any(r.startswith("没敢改的|1 格") for r in flat)
    summary = json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])
    assert "规整了" in summary["report_text"]


def test_merge_step_without_rules_unchanged(tmp_path):
    out = _issue(tmp_path, "走访日期、走访结果")
    back = tmp_path / "回收"
    back.mkdir()
    _fill(out, "东区", back / "东区.xlsx", [["2026.10.5", "x"], ["2026.10.6", "y"], ["2026.10.7", "z"]])
    path = collect.run_collect({"input_path": str(back), "collect_record": os.path.join(out, collect.RECORD_NAME)},
                               lambda m: None)
    ws = openpyxl.load_workbook(path)["走访"]
    assert ws["C2"].value == "2026.10.5"
