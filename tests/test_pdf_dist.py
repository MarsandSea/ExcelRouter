"""core/pdf_dist.py 的集成测试：加密、水印、分发清单、容错、停止。

水印文字用例统一用 ASCII 网格名——CI（GitHub Windows runner）字体环境
不可控，避免与中文字体探测耦合；中文水印在本机冒烟里验证。
"""
import openpyxl
import pytest
from fpdf import FPDF
from pypdf import PdfReader

from core import pdf_dist
from core.pdf_dist import (
    list_mapping_columns, read_mapping, run_pdf_dist, MANIFEST_NAME, _find_cjk_font,
)


# ---------- 样本构造 ----------

def _make_pdf(path, pages=2, encrypt_pwd=None):
    pdf = FPDF()
    pdf.set_auto_page_break(False)
    for i in range(pages):
        pdf.add_page()
        pdf.set_font("Helvetica", size=14)
        pdf.text(30, 40, f"SOURCE-MARK page{i + 1}")
    if encrypt_pwd:
        from pypdf import PdfWriter
        import io
        reader = PdfReader(io.BytesIO(bytes(pdf.output())))
        writer = PdfWriter()
        writer.append(reader)
        writer.encrypt(user_password=encrypt_pwd, algorithm="AES-256")
        with open(path, "wb") as f:
            writer.write(f)
    else:
        pdf.output(str(path))


def _make_mapping(path, rows, headers=("GridName", "Pwd", "Receiver")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    wb.save(path)


@pytest.fixture
def workspace(tmp_path):
    src = tmp_path / "report.pdf"
    _make_pdf(src)
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [
        ("GridA", "001234", "Zhang"),   # 文本密码，前导零
        ("GridB", 8888, "Li"),          # 数字密码（openpyxl 读成 int/float）
        ("GridC", None, "Wang"),        # 空密码，应跳过
    ])
    out = tmp_path / "out"
    cfg = {
        "pdf_input_paths": [str(src)],
        "pdf_mapping_path": str(mp),
        "pdf_grid_column": "GridName",
        "pdf_password_column": "Pwd",
        "pdf_receiver_column": "Receiver",
        "pdf_watermark": True,
        "pdf_watermark_text": "{grid} {date}",
        "output_path": str(out),
    }
    return tmp_path, cfg


def _read_manifest(out_dir):
    wb = openpyxl.load_workbook(out_dir / MANIFEST_NAME)
    ws = wb.active
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    wb.close()
    return rows


# ---------- 用例 ----------

def test_list_mapping_columns(workspace):
    tmp, cfg = workspace
    assert list_mapping_columns(cfg["pdf_mapping_path"]) == ["GridName", "Pwd", "Receiver"]


def test_read_mapping_password_as_text(workspace):
    tmp, cfg = workspace
    rows, warnings = read_mapping(cfg["pdf_mapping_path"], "GridName", "Pwd", "Receiver")
    assert [r["grid"] for r in rows] == ["GridA", "GridB"]
    assert rows[0]["password"] == "001234"      # 前导零不丢
    assert rows[1]["password"] == "8888"        # 不会变成 8888.0
    assert any("GridC" in w for w in warnings)  # 空密码警告

    with pytest.raises(ValueError):
        read_mapping(cfg["pdf_mapping_path"], "GridName", "不存在的列")


def test_encrypt_right_and_wrong_password(workspace):
    tmp, cfg = workspace
    out = run_pdf_dist(cfg, log_fn=lambda m: None)
    assert out

    r = PdfReader(tmp / "out" / "GridA" / "report.pdf")
    assert r.is_encrypted
    assert r.decrypt("001234") != 0
    assert "SOURCE-MARK page1" in r.pages[0].extract_text()

    r2 = PdfReader(tmp / "out" / "GridB" / "report.pdf")
    assert r2.decrypt("wrong") == 0
    assert r2.decrypt("8888") != 0

    # 空密码的 GridC 不产出
    assert not (tmp / "out" / "GridC").exists()


def test_watermark_on_and_off(workspace):
    tmp, cfg = workspace
    run_pdf_dist(cfg, log_fn=lambda m: None)
    r = PdfReader(tmp / "out" / "GridA" / "report.pdf")
    r.decrypt("001234")
    for page in r.pages:                       # 每页都有该网格的水印
        assert "GridA" in page.extract_text()

    cfg2 = dict(cfg, pdf_watermark=False, output_path=str(tmp / "out2"))
    run_pdf_dist(cfg2, log_fn=lambda m: None)
    r = PdfReader(tmp / "out2" / "GridA" / "report.pdf")
    r.decrypt("001234")
    assert "GridA" not in r.pages[0].extract_text()


def test_manifest_content(workspace):
    tmp, cfg = workspace
    run_pdf_dist(cfg, log_fn=lambda m: None)
    rows = _read_manifest(tmp / "out")
    assert len(rows) == 2                       # GridC 跳过，不进清单
    by_grid = {r[0]: r for r in rows}
    assert by_grid["GridA"][2] == "001234"      # 密码为字符串、前导零保留
    assert by_grid["GridA"][3] == "Zhang"
    assert by_grid["GridA"][4] == 2             # 页数
    assert "完成" in by_grid["GridB"][5]


def test_bad_filename_and_encrypted_source(tmp_path):
    """非法文件名字符的网格照常产出（净化）；自带密码的源 PDF 跳过但不中断。"""
    ok_src = tmp_path / "ok.pdf"
    _make_pdf(ok_src)
    locked = tmp_path / "locked.pdf"
    _make_pdf(locked, encrypt_pwd="x")
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [('Grid/A:B?', "111", "")])
    logs = []
    cfg = {
        "pdf_input_paths": [str(ok_src), str(locked)],
        "pdf_mapping_path": str(mp),
        "pdf_grid_column": "GridName",
        "pdf_password_column": "Pwd",
        "pdf_watermark": False,
        "output_path": str(tmp_path / "out"),
    }
    out = run_pdf_dist(cfg, log_fn=logs.append)
    assert out
    assert any("locked.pdf" in m for m in logs)          # 带密码源被警告跳过
    safe_dir = tmp_path / "out" / "Grid_A_B_"
    assert (safe_dir / "ok.pdf").exists()                # 净化后的目录名


def test_stop_flag(workspace):
    tmp, cfg = workspace
    logs = []
    out = run_pdf_dist(cfg, log_fn=logs.append, stop_flag=lambda: True)
    assert out is None                                    # 一个网格都没做
    assert any("停止" in m for m in logs)

# ---------- 模板 / 随机密码 / 收尾摘要行（v2.7）----------

def test_write_mapping_template(tmp_path):
    """模板：首行表头三列，第二个 sheet 放说明（不影响 read_mapping 只读第一个 sheet）。"""
    from core.pdf_dist import write_mapping_template
    p = tmp_path / "tpl.xlsx"
    write_mapping_template(str(p))
    assert list_mapping_columns(str(p)) == ["网格", "密码", "接收人"]
    wb = openpyxl.load_workbook(p)
    assert "填写说明" in wb.sheetnames
    wb.close()

def test_fill_random_passwords_blanks_only(tmp_path):
    """只填空白的密码单元格；已手填的保留；原文件不被改动。"""
    from core.pdf_dist import fill_random_passwords
    mp = tmp_path / "m.xlsx"
    _make_mapping(mp, [("G1", None, ""), ("G2", "keepme", ""), ("G3", None, "")])
    out, n = fill_random_passwords(str(mp), "GridName", "Pwd")
    assert n == 2                       # 只补 G1 / G3
    assert out.endswith("_含密码.xlsx")
    rows, _ = read_mapping(out, "GridName", "Pwd")
    pw = {r["grid"]: r["password"] for r in rows}
    assert pw["G2"] == "keepme"         # 已填的不覆盖
    assert len(pw["G1"]) == 8 and len(pw["G3"]) == 8
    assert pw["G1"] != pw["G3"]
    rows0, _ = read_mapping(str(mp), "GridName", "Pwd")
    assert [r["grid"] for r in rows0] == ["G2"]   # 原文件保持原样

def test_fill_random_passwords_new_column(tmp_path):
    """清单没有密码列时新增「密码」列；密码字符集排除易混字符。"""
    from core.pdf_dist import fill_random_passwords, gen_password
    mp = tmp_path / "m2.xlsx"
    _make_mapping(mp, [("东区", "", ""), ("西区", "", "")], headers=("网格",))
    out, n = fill_random_passwords(str(mp), "网格")
    assert n == 2
    rows, _ = read_mapping(out, "网格", "密码")
    assert len(rows) == 2
    for _ in range(50):
        pwd = gen_password()
        assert len(pwd) == 8
        assert not set(pwd) & set("0O1lI")   # 无易混字符

def test_summary_line_pdf(workspace):
    """run_pdf_dist 收尾应打恰好一行 [SUMMARY] JSON。"""
    import json
    tmp, cfg = workspace
    logs = []
    out = run_pdf_dist(cfg, log_fn=logs.append)
    summary = [m for m in logs if m.startswith("[SUMMARY] ")]
    assert len(summary) == 1
    s = json.loads(summary[0][len("[SUMMARY] "):])
    assert s["mode"] == "pdf"
    assert s["grids_ok"] == 2           # GridA / GridB（GridC 空密码被跳过）
    assert s["grids_fail"] == 0
    assert s["files"] == 2
    assert s["output"] == out
    assert s["manifest"]

def test_fill_random_passwords_reuse_existing_blank_col(tmp_path):
    """issue #1 回归：清单已有空白「密码」列且未指定 password_col 时，复用该列，
    不得追加重复同名列（旧版追加后 read_mapping 命中原空白列，密码全部漏读）。"""
    from core.pdf_dist import fill_random_passwords
    mp = tmp_path / "m.xlsx"
    _make_mapping(mp, [("G1", None, ""), ("G2", None, "")], headers=("网格", "密码"))
    out, n = fill_random_passwords(str(mp), "网格")
    assert n == 2
    wb = openpyxl.load_workbook(out)
    header = [c.value for c in wb.worksheets[0][1]]
    wb.close()
    assert header.count("密码") == 1            # 无重复列
    rows, _ = read_mapping(out, "网格", "密码")
    pw = {r["grid"]: r["password"] for r in rows}
    assert len(pw) == 2 and all(p and len(p) == 8 for p in pw.values())   # 密码可读回


# ---------- 中文字体探测（v2.8.0 麒麟适配）----------
# 这一节补的是本文件开头注明的空白：中文水印此前在任何平台都没有测试覆盖。

def test_find_cjk_font_windows_unchanged(tmp_path, monkeypatch):
    """Windows 分支必须与 v2.7.2 逐字一致——候选顺序也不能变。"""
    fonts = tmp_path / "Fonts"
    fonts.mkdir()
    (fonts / "simkai.ttf").write_bytes(b"x")
    (fonts / "simhei.ttf").write_bytes(b"x")
    monkeypatch.setenv("WINDIR", str(tmp_path))
    # simhei 在候选列表里排在 simkai 前面；且 Windows 分支不做可加载性校验。
    # 直接调 _find_cjk_font_win()，这样在 Linux CI 上也能跑（不用改全局 os.name）。
    assert pdf_dist._find_cjk_font_win() == str(fonts / "simhei.ttf")


def test_find_cjk_font_linux_prefers_cjk_candidates(tmp_path, monkeypatch):
    """Linux 分支：命中候选名的才算数，随便一个 .ttf 不能冒充中文字体。"""
    (tmp_path / "wqy-zenhei.ttc").write_bytes(b"x")
    (tmp_path / "DejaVuSans.ttf").write_bytes(b"x")
    monkeypatch.setattr(pdf_dist, "_LINUX_FONT_DIRS", [str(tmp_path)])
    monkeypatch.setattr(pdf_dist, "_font_usable", lambda p: True)
    monkeypatch.setattr(pdf_dist.shutil, "which", lambda c: None)   # 屏蔽 fc-match
    monkeypatch.delenv("ER_CJK_FONT", raising=False)
    assert pdf_dist._find_cjk_font_posix() == str(tmp_path / "wqy-zenhei.ttc")

    # 只剩 DejaVu 时应当放弃（返回 None），而不是拿一个无中文字形的字体去糊弄
    (tmp_path / "wqy-zenhei.ttc").unlink()
    assert pdf_dist._find_cjk_font_posix() is None


def test_find_cjk_font_env_override(tmp_path, monkeypatch):
    """ER_CJK_FONT 是信创机器上的逃生口，优先级最高；指向不存在的文件则被忽略。"""
    real = tmp_path / "my.ttf"
    real.write_bytes(b"x")
    (tmp_path / "wqy-zenhei.ttc").write_bytes(b"x")
    monkeypatch.setattr(pdf_dist, "_LINUX_FONT_DIRS", [str(tmp_path)])
    monkeypatch.setattr(pdf_dist, "_font_usable", lambda p: True)
    monkeypatch.setattr(pdf_dist.shutil, "which", lambda c: None)

    monkeypatch.setenv("ER_CJK_FONT", str(real))
    assert pdf_dist._find_cjk_font_posix() == str(real)

    monkeypatch.setenv("ER_CJK_FONT", str(tmp_path / "并不存在.ttf"))
    assert pdf_dist._find_cjk_font_posix() == str(tmp_path / "wqy-zenhei.ttc")


def test_find_cjk_font_rejects_unloadable(tmp_path, monkeypatch):
    """★ 最关键的一条：名字对得上但文件是坏的，必须拒掉。

    这是「优雅降级成 ???」和「分发跑到一半崩掉」的分界线——本用例用真实的
    _font_usable（让 fpdf2 自己判断），不做桩。
    """
    (tmp_path / "wqy-zenhei.ttc").write_bytes(b"NOT A FONT")
    monkeypatch.setattr(pdf_dist, "_LINUX_FONT_DIRS", [str(tmp_path)])
    monkeypatch.setattr(pdf_dist.shutil, "which", lambda c: None)
    monkeypatch.delenv("ER_CJK_FONT", raising=False)
    assert pdf_dist._find_cjk_font_posix() is None


def test_watermark_survives_broken_font(workspace, monkeypatch):
    """坏字体不能中断整轮分发：降级成 Helvetica 继续跑完，所有网格都要产出。"""
    tmp_path, cfg = workspace
    bad = tmp_path / "broken.ttf"
    bad.write_bytes(b"NOT A FONT")
    monkeypatch.setattr(pdf_dist, "_find_cjk_font", lambda: str(bad))
    pdf_dist._wm_cache.clear()
    logs = []
    run_pdf_dist(cfg, log_fn=logs.append)
    out = tmp_path / "out"
    assert (out / "GridA" / "report.pdf").exists()
    assert (out / "GridB" / "report.pdf").exists()


@pytest.mark.skipif(_find_cjk_font() is None, reason="本机没有可用的中文字体")
def test_watermark_cjk_end_to_end(tmp_path):
    """中文网格名的水印必须真的嵌入中文字体，而不是退回 Helvetica 的 '???'。"""
    src = tmp_path / "report.pdf"
    _make_pdf(src)
    mp = tmp_path / "map.xlsx"
    _make_mapping(mp, [("研发一组", "001234", "张三")])
    out = tmp_path / "out"
    pdf_dist._wm_cache.clear()
    logs = []
    run_pdf_dist({
        "pdf_input_paths": [str(src)],
        "pdf_mapping_path": str(mp),
        "pdf_grid_column": "GridName",
        "pdf_password_column": "Pwd",
        "pdf_receiver_column": "Receiver",
        "pdf_watermark": True,
        "pdf_watermark_text": "{grid} {date}",
        "output_path": str(out),
    }, log_fn=logs.append)

    assert any("水印字体" in m for m in logs), logs
    dst = out / "研发一组" / "report.pdf"
    assert dst.exists()
    reader = PdfReader(str(dst))
    reader.decrypt("001234")
    raw = reader.pages[0].get_contents().get_data()
    # 退回 Helvetica 时中文会被 encode("ascii","replace") 打成 '?'
    assert b"?" * 3 not in raw
    # 中文字体是以 Type0 复合字体（子集化后）嵌进去的；退回 Helvetica 时
    # 页面里只会有 /Type1 的内置字体，不会出现 /Type0。
    fonts = reader.pages[0]["/Resources"]["/Font"]
    subtypes = [f.get_object().get("/Subtype") for f in fonts.values()]
    assert "/Type0" in subtypes, f"水印没有嵌入中文字体，说明退回了 Helvetica：{subtypes}"
