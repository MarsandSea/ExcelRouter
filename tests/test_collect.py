"""core/collect.py 的测试：查收各状态、认领、漏填行号、多轮、总表按列分人、和下发时比对、汇总。

Copyright (c) 2026 AbeLin · MIT License
"""
import json
import os
import subprocess
import sys
import time

import openpyxl

from core import collect
from core.collect import (
    ST_CHANGED, ST_EMPTY, ST_MISSING, ST_OK, ST_PARTIAL, check_collect, run_collect,
)

HEAD = ["客户名称", "客户经理", "走访日期", "走访结果", "备注"]
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _book(path, rows, head=HEAD, guide=None, title=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "走访"
    if title:
        ws.append([title])
    ws.append(list(head))
    for r in rows:
        ws.append(list(r))
    if guide:
        g = wb.create_sheet(collect.GUIDE_SHEET)
        g.append(["填报对象", guide])
        g.append(["截止时间", "10月15日"])
    wb.save(path)


def _full(n, who="经理"):
    return [[f"客户{i}", who, f"2026-10-0{i % 9 + 1}", "已走访", ""] for i in range(1, n + 1)]


def _run(cfg):
    logs = []
    out = check_collect(cfg, logs.append)
    summary = json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])
    return out, summary, logs


def _notice(out):
    with open(os.path.join(out, collect.NOTICE_NAME), encoding="utf-8-sig") as f:
        return f.read()


def _people(cfg):
    ctx, _res, report = collect._inspect(cfg, lambda m: None, None, None)
    return {p["name"]: p for p in report["people"]}, report


def test_does_not_import_pandas():
    code = "import sys; import core.collect; print('pandas' in sys.modules)"
    out = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True, check=True)
    assert out.stdout.strip() == "False"


def test_check_reports_each_status(tmp_path):
    src = tmp_path / "九月走访"
    src.mkdir()
    _book(src / "张三.xlsx", _full(4))
    rows = _full(5)
    rows[1][2] = None          # 第 3 行走访日期没填
    rows[3][2] = ""            # 第 5 行
    rows[3][3] = None          # 第 5 行走访结果
    _book(src / "李四_走访表.xlsx", rows)
    _book(src / "赵六.xlsx", [])                                   # 只有表头
    _book(src / "孙八.xlsx", [[f"客户{i}", "经理", None, None, ""] for i in range(3)])   # 该填的一格没填
    cfg = {"input_path": str(src), "collect_expect": "张三、李四、王五、赵六、孙八",
           "collect_required": ["走访日期", "走访结果"], "collect_due": "10月15日 18:00"}
    out, summary, _logs = _run(cfg)
    assert out == os.path.join(str(src), collect.RESULT_DIR)
    assert summary["expected"] == 5 and summary["submitted"] == 4 and summary["ok"] == 1
    assert summary["missing"] == 1 and summary["empty"] == 2 and summary["partial"] == 1
    people, _ = _people(cfg)
    assert people["张三"]["status"] == ST_OK
    assert people["王五"]["status"] == ST_MISSING
    assert people["赵六"]["status"] == ST_EMPTY and people["孙八"]["status"] == ST_EMPTY
    li = people["李四"]
    assert li["status"] == ST_PARTIAL
    assert {b["col"]: b["rows"] for b in li["blanks"]} == {"走访日期": [3, 5], "走访结果": [5]}
    text = _notice(out)
    assert "应交 5 人，已交 4 人，合格 1 人" in text
    assert "未交（1）：王五" in text
    assert "交了空表（2）：赵六、孙八" in text
    assert "李四：走访日期 漏 2 行（第 3、5 行），走访结果 漏 1 行（第 5 行）" in text
    assert "请以上同事于 10月15日 18:00 前补交" in text
    assert "王五：【九月走访】还没收到你的表" in text
    wb = openpyxl.load_workbook(os.path.join(out, collect.PROGRESS_NAME))
    ws = wb["回收进度"]
    assert ws.cell(row=3, column=2).value == "王五"              # 未交排在最前
    assert ws.cell(row=3, column=3).fill.fgColor.rgb.endswith("FFC7CE")
    assert [c.value for c in wb["问题明细"]["C"]][1:] == ["走访日期", "走访结果"]


def test_longest_name_wins(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _book(src / "城区二.xlsx", _full(2))
    people, _ = _people({"input_path": str(src), "collect_expect": ["城区", "城区二"]})
    assert people["城区"]["status"] == ST_MISSING
    assert people["城区二"]["files"] == ["城区二.xlsx"]
    missing, multi, others = collect.expect_check(["城区", "城区二"], ["城区二.xlsx"])
    assert missing == ["城区"] and not multi and not others


def test_guide_sheet_claims_file(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _book(src / "走访表(已填).xlsx", _full(2), guide="李四")
    _book(src / "不知道是谁的.xlsx", _full(2))
    people, report = _people({"input_path": str(src), "collect_expect": ["李四", "张三"]})
    assert people["李四"]["files"] == ["走访表(已填).xlsx"] and people["李四"]["rows"] == 2
    assert people["张三"]["status"] == ST_MISSING
    assert report["extra"] == ["不知道是谁的.xlsx"]


def test_required_columns_inferred_without_list(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    for n in ("甲", "乙", "丙"):
        _book(src / f"{n}.xlsx", _full(3))
    rows = _full(3)
    rows[0][3] = None
    _book(src / "丁.xlsx", rows)
    people, report = _people({"input_path": str(src)})
    assert "走访结果" in report["required"] and "备注" not in report["required"]
    assert people["丁"]["status"] == ST_PARTIAL and people["甲"]["status"] == ST_OK
    out, summary, _ = _run({"input_path": str(src)})
    assert summary["expected"] is None
    assert "没有应交名单" in _notice(out)


def test_name_list_file_and_rounds(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    wb = openpyxl.Workbook()
    wb.active.append(["序号", "姓名"])
    for i, n in enumerate(["张三", "李四"], 1):
        wb.active.append([i, n])
    wb.save(src / "应交名单.xlsx")
    _book(src / "张三.xlsx", _full(2))
    out, s1, _ = _run({"input_path": str(src)})
    assert s1["expected"] == 2 and s1["missing"] == 1
    assert "比上次新交" not in _notice(out)
    _book(src / "李四.xlsx", _full(2))
    out, s2, _ = _run({"input_path": str(src)})
    assert s2["ok"] == 2
    text = _notice(out)
    assert "比上次新交：李四。" in text and "全部交齐" in text
    with open(os.path.join(out, collect.ROUNDS_NAME), encoding="utf-8") as f:
        assert len(json.load(f)["rounds"]) == 2


def test_resubmitted_copy_uses_newest(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    rows = _full(2)
    rows[0][2] = None
    _book(src / "张三.xlsx", rows)
    old = time.time() - 600
    os.utime(src / "张三.xlsx", (old, old))
    _book(src / "张三(2).xlsx", _full(2))
    people, _ = _people({"input_path": str(src), "collect_expect": ["张三"], "collect_required": ["走访日期"]})
    p = people["张三"]
    assert p["files"] == ["张三(2).xlsx"] and p["status"] == ST_OK
    assert any("旧版本" in n for n in p["notes"])


def test_rows_mode_groups_one_table(tmp_path):
    path = tmp_path / "总表.xlsx"
    head = ["网格", "客户名称", "走访日期", "走访结果"]
    rows = [["东区", "a", "10-01", "好"], ["东区", "b", "10-02", "好"],
            ["西区", "c", None, "好"], ["西区", "d", "10-03", None],
            ["南区", "e", None, None], ["北外", "f", "10-01", "好"], [None, "g", "10-01", "好"]]
    _book(path, rows, head=head, title="九月走访总表")
    cfg = {"input_path": str(path), "collect_by": "网格", "collect_expect": "东区、西区、南区、中区",
           "collect_required": "走访日期、走访结果"}
    people, report = _people(cfg)
    assert people["东区"]["status"] == ST_OK and people["东区"]["rows"] == 2
    assert {b["col"]: b["rows"] for b in people["西区"]["blanks"]} == {"走访日期": [5], "走访结果": [6]}
    assert people["南区"]["status"] == ST_EMPTY
    assert people["中区"]["status"] == ST_MISSING
    assert report["extra"] == ["北外"]
    assert any("1 行没写网格" in n for n in report["notes"])
    out, summary, _ = _run(cfg)
    assert summary["ok"] == 1 and "名单外的分组（1）：北外" in _notice(out)


def test_changes_against_issued_record(tmp_path):
    head = HEAD
    issued = tmp_path / "下发"
    issued.mkdir()
    rows = {n: [[f"{n}客户{i}", n, None, None, ""] for i in range(1, 5)] for n in ("张三", "李四", "王五")}
    keys = collect.lock_keys(head, ["走访日期", "走访结果"])
    groups = {}
    for n, rs in rows.items():
        p = issued / f"{n}.xlsx"
        _book(p, rs)
        groups[n] = {"file": p.name, "sheets": collect.fingerprint_file(str(p), {}, keys)}
    record = {"title": "九月走访", "due": "10月15日", "fill_columns": ["走访日期", "走访结果"],
              "lock_keys": keys, "lock_columns": [h for h in head if h not in ("走访日期", "走访结果")],
              "groups": groups}
    back = tmp_path / "回收"
    back.mkdir()
    (back / collect.RECORD_NAME).write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")

    def filled(rs):
        return [[r[0], r[1], "10-01", "好", r[4]] for r in rs]

    _book(back / "张三.xlsx", filled(rows["张三"])[:3])                 # 删了一行
    li = filled(rows["李四"])
    li[1][0] = "改过的客户"                                             # 第 3 行客户名称被改
    _book(back / "李四.xlsx", li)
    _book(back / "王五.xlsx", list(reversed(filled(rows["王五"]))))     # 只是排了序
    people, report = _people({"input_path": str(back)})
    assert report["required"] == ["走访日期", "走访结果"]
    assert people["张三"]["status"] == ST_CHANGED and people["张三"]["changes"] == ["少了 1 行"]
    assert people["李四"]["changes"] == ["第 3 行和下发时不一样（改了「客户名称」）"]
    assert people["王五"]["status"] == ST_OK
    out, summary, _ = _run({"input_path": str(back)})
    text = _notice(out)
    assert summary["changed"] == 2 and "【九月走访】" in text and "只填黄色的列" in text


def test_run_collect_merges_and_reports(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _book(src / "张三.xlsx", _full(2, "张三"))
    rows = _full(3, "李四")
    rows[2][2] = None
    _book(src / "李四.xlsx", rows)
    _book(src / "赵六.xlsx", [])
    logs = []
    out = run_collect({"input_path": str(src), "collect_expect": "张三、李四、王五、赵六",
                       "collect_required": "走访日期", "collect_title": "九月走访"}, logs.append)
    assert os.path.basename(out).startswith("汇总_九月走访_")
    summary = json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])
    assert summary["rows"] == 5 and summary["files"] == 2
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["走访", collect.REPORT_SHEET]
    ws = wb["走访"]
    assert ws.cell(row=1, column=6).value == collect.SOURCE_COL
    assert [ws.cell(row=r, column=6).value for r in range(2, 7)] == ["张三", "张三", "李四", "李四", "李四"]
    assert ws.cell(row=6, column=3).fill.fgColor.rgb.endswith("FFC7CE")     # 李四漏填的格子标浅红
    rep = wb[collect.REPORT_SHEET]
    table = {r[0]: r for r in rep.iter_rows(min_row=2, values_only=True)}
    assert table["王五"][1] == ST_MISSING and table["赵六"][9] == "空表，未汇总"
    assert table["李四"][4] == 3 and "走访日期 1 格" in table["李四"][6]


def test_result_dir_is_not_rescanned(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _book(src / "张三.xlsx", _full(2))
    cfg = {"input_path": str(src), "collect_expect": ["张三"]}
    run_collect(cfg, lambda m: None)
    _out, summary, _ = _run(cfg)
    assert summary["submitted"] == 1 and summary["extra"] == 0


def test_title_rows_and_total_rows(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _book(src / "张三.xlsx", _full(3) + [["合计", "", "", "", ""]], title="九月走访表")
    people, _ = _people({"input_path": str(src), "collect_expect": ["张三"], "collect_required": ["走访结果"]})
    assert people["张三"]["rows"] == 3 and people["张三"]["status"] == ST_OK


def test_stop_flag_skips_outputs(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _book(src / "张三.xlsx", _full(2))
    logs = []
    out = check_collect({"input_path": str(src)}, logs.append, stop_flag=lambda: True)
    assert not os.path.exists(os.path.join(out, collect.NOTICE_NAME))
    assert '"stopped": true' in next(m for m in logs if m.startswith("[SUMMARY]"))


def test_identical_file_from_another_person_still_counts(tmp_path):
    src = tmp_path / "收"
    src.mkdir()
    _book(src / "张三.xlsx", _full(2))
    (src / "李四.xlsx").write_bytes((src / "张三.xlsx").read_bytes())
    (src / "张三 - 副本.xlsx").write_bytes((src / "张三.xlsx").read_bytes())
    people, report = _people({"input_path": str(src), "collect_expect": ["张三", "李四"]})
    assert people["李四"]["status"] == ST_OK and people["李四"]["rows"] == 2
    assert any("内容完全一样" in n for n in people["李四"]["notes"])
    assert len(people["张三"]["files"]) == 1 and len(report["dupes"]) == 1     # 同一个人的两份一样的只算一份


def test_split_issue_then_collect_round_trip(tmp_path):
    """拆分时指定填写列 → 标黄锁定、写下发记录 → 收回来（改名、改了锁定格、没交）→ 查收认得出。"""
    from core.splitter import run_split
    src = tmp_path / "九月走访.xlsx"
    rows = [[g, f"{g}客户{i}", "经理", None] for g in ("东区", "西区", "南区") for i in range(1, 4)]
    _book(src, rows, head=["网格", "客户名称", "客户经理", "走访日期"])
    out_root = tmp_path / "out"
    logs = []
    out = run_split({"input_path": str(src), "output_path": str(out_root), "split_column": "网格",
                     "fill_columns": "走访日期、走访结果", "due": "10月15日", "make_zip": False}, logs.append)
    assert '"issued": true' in next(m for m in logs if m.startswith("[SUMMARY]"))
    for name in (collect.RECORD_NAME, collect.SEND_LIST_NAME, collect.ISSUE_NOTE_NAME):
        assert os.path.exists(os.path.join(out, name))
    wb = openpyxl.load_workbook(os.path.join(out, "东区.xlsx"))
    ws = wb.worksheets[0]
    assert ws.protection.sheet
    assert [c.value for c in ws[1]] == ["网格", "客户名称", "客户经理", "走访日期", "走访结果"]   # 缺的列加在最后
    assert ws["D2"].protection.locked is False and ws["E3"].protection.locked is False
    assert ws["B2"].protection.locked is True
    assert ws["D2"].fill.fgColor.rgb.endswith(collect.FILL_COLOR)
    assert wb[collect.GUIDE_SHEET]["B1"].value == "东区"
    with open(os.path.join(out, collect.RECORD_NAME), encoding="utf-8") as f:
        record = json.load(f)
    assert list(record["groups"]) == ["东区", "南区", "西区"]
    assert record["fill_columns"] == ["走访日期", "走访结果"] and "客户名称" in record["lock_columns"]
    assert record["groups"]["东区"]["sheets"]["走访"]["rows"] == 3
    assert "客户" not in json.dumps(record, ensure_ascii=False).replace("客户名称", "").replace("客户经理", "")

    back = tmp_path / "回收"
    back.mkdir()

    def fill_back(group, dst, edit=None):
        b = openpyxl.load_workbook(os.path.join(out, f"{group}.xlsx"))
        s = b.worksheets[0]
        for r in range(2, s.max_row + 1):
            s.cell(row=r, column=4).value = "10-01"
            s.cell(row=r, column=5).value = "已走访"
        if edit:
            edit(s)
        b.save(dst)

    fill_back("东区", back / "我的表.xlsx")                                   # 改了文件名，靠「填报说明」认
    fill_back("西区", back / "西区.xlsx", lambda s: s.cell(row=3, column=2, value="别人的客户"))
    cfg = {"input_path": str(back), "collect_record": os.path.join(out, collect.RECORD_NAME)}
    people, report = _people(cfg)
    assert report["required"] == ["走访日期", "走访结果"]
    assert people["东区"]["status"] == ST_OK and people["东区"]["files"] == ["我的表.xlsx"]
    assert people["西区"]["status"] == ST_CHANGED
    assert people["西区"]["changes"] == ["第 3 行和下发时不一样（改了「客户名称」）"]
    assert people["南区"]["status"] == ST_MISSING
    out2, summary, _ = _run(cfg)
    assert "于 10月15日 前补交" in _notice(out2)


def test_renamed_without_guide_sheet_is_still_claimed(tmp_path):
    """删了「填报说明」、文件名里也没有名字（把数据贴进新表）：靠拆分列或行指纹认出是谁的。"""
    from core.splitter import run_split
    src = tmp_path / "九月走访.xlsx"
    rows = [[g, f"{g}客户{i}", "经理", None] for g in ("东区", "西区", "南区") for i in range(1, 4)]
    _book(src, rows, head=["网格", "客户名称", "客户经理", "走访日期"])
    out = run_split({"input_path": str(src), "output_path": str(tmp_path / "out"), "split_column": "网格",
                     "fill_columns": "走访日期", "make_zip": False}, lambda m: None)
    back = tmp_path / "回收"
    back.mkdir()
    for group, dst in (("东区", "新建 XLSX 工作表.xlsx"), ("西区", "已填好.xlsx")):
        s = openpyxl.load_workbook(os.path.join(out, f"{group}.xlsx")).worksheets[0]
        data = [[c.value for c in r] for r in s.iter_rows()]
        nb = openpyxl.Workbook()                      # 贴进一张新表：没有「填报说明」、没有保护
        nb.active.append(data[0][:4])
        for r in data[1:]:
            nb.active.append(r[:3] + ["10-01"])
        nb.save(back / dst)
    _book(back / "别的表.xlsx", [["北区", "无关客户", "经理", "10-01"]], head=["网格", "客户名称", "客户经理", "走访日期"])

    with open(os.path.join(out, collect.RECORD_NAME), encoding="utf-8") as f:
        record = json.load(f)
    record.pop("split_column")                        # 只剩行指纹可用
    rec_path = tmp_path / "记录.json"
    rec_path.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
    for rp in (os.path.join(out, collect.RECORD_NAME), str(rec_path)):
        people, report = _people({"input_path": str(back), "collect_record": rp})
        assert people["东区"]["files"] == ["新建 XLSX 工作表.xlsx"] and people["东区"]["status"] == ST_OK
        assert people["西区"]["files"] == ["已填好.xlsx"]
        assert people["南区"]["status"] == ST_MISSING
        assert report["extra"] == ["别的表.xlsx"]


def test_merge_tables_plain(tmp_path):
    """普通汇总：不要名单；「明细_1 / 明细_2」各算一份（不当成同一人重复交），列顺序不同照样对齐。"""
    _book(tmp_path / "明细_1.xlsx", _full(2))
    _book(tmp_path / "明细_2.xlsx", [[r[1], r[0]] + r[2:] for r in _full(3, "李四")],
          head=["客户经理", "客户名称", "走访日期", "走访结果", "备注"])
    _book(tmp_path / "发送清单.xlsx", _full(1, "王五"))          # 名字像名单，普通汇总照样合
    _book(tmp_path / "签到.xlsx", [["某人", "09:00"]], head=["姓名", "签到时间"])
    logs = []
    out = collect.merge_tables({"input_path": str(tmp_path)}, logs.append)
    summary = json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])
    assert summary["mode"] == "merge" and summary["files"] == 3 and summary["rows"] == 6
    assert "签到.xlsx「走访」：表头对不上" in summary["report_text"]
    wb = openpyxl.load_workbook(out)
    ws = wb["走访"]
    assert [c.value for c in ws[1]] == HEAD + [collect.SOURCE_COL]
    assert ws.cell(row=2, column=6).value == "发送清单"          # 按文件名排序：发送清单、明细_1、明细_2
    assert ws.cell(row=5, column=2).value == "李四" and ws.cell(row=5, column=6).value == "明细_2"
    rep = wb[collect.REPORT_SHEET]
    assert rep["A1"].value == "来源文件"                       # 没有名单：不写「应交」「状态」


def test_merge_tables_one_file_sheets(tmp_path):
    p = tmp_path / "全年.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "1月"
    wb.active.append(["姓名", "金额"])
    wb.active.append(["甲", 1])
    s = wb.create_sheet("2月")
    s.append(["姓名", "金额"])
    s.append(["乙", 2])
    wb.save(p)
    logs = []
    out = collect.merge_tables({"input_path": str(p)}, logs.append)
    rows = list(openpyxl.load_workbook(out).worksheets[0].iter_rows(values_only=True))
    assert rows == [("姓名", "金额", "来源"), ("甲", 1, "1月"), ("乙", 2, "2月")]
    assert any("把 2 个 sheet 合成了一张" in m for m in logs)
    assert os.path.dirname(out) == str(tmp_path / collect.RESULT_DIR)


def test_unreadable_files_count_as_submitted_with_plain_reason(tmp_path):
    """交了 .et / 网页格式的 .xls / 加了口令的 .xlsx：通报里不能写成「未交」，原因要是人话。

    回归：旧版只收 .xlsx / .xls，交了 .et 的人被写进「未交」发到群里；
    其余两种在催办里显示「Expected BOF record」「File is not a zip file」原话。
    """
    src = tmp_path / "收表"
    src.mkdir()
    _book(src / "张三.xlsx", _full(3))
    (src / "李四.et").write_bytes(b"\xd0\xcf\x11\xe0fake")
    (src / "王五.xls").write_text("<html><table><tr><td>客户名称</td></tr></table></html>", encoding="utf-8")
    (src / "赵六.xlsx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1encrypted")
    _book(src / "孙八.xlsm", _full(2))
    cfg = {"input_path": str(src), "collect_expect": "张三、李四、王五、赵六、孙八"}
    out, summary, _logs = _run(cfg)
    assert summary["missing"] == 0 and summary["submitted"] == 5
    text = _notice(out)
    assert "未交" not in text
    assert "WPS 表格自己的格式" in text and "网页格式" in text and "打开口令" in text
    assert "BOF" not in text and "zip" not in text
    people, _report = _people(cfg)
    assert people["孙八"]["status"] == ST_OK          # .xlsm 照常能查


def test_merge_tables_reports_unreadable_in_words(tmp_path):
    src = tmp_path / "报表"
    src.mkdir()
    _book(src / "a.xlsx", _full(2))
    (src / "b.et").write_bytes(b"\xd0\xcf\x11\xe0fake")
    logs = []
    out = collect.merge_tables({"input_path": str(src)}, logs.append)
    assert out
    summary = json.loads(next(m for m in logs if m.startswith("[SUMMARY]"))[len("[SUMMARY] "):])
    assert summary["rows"] == 2 and summary["failed_files"] == 1
    assert "WPS 表格自己的格式" in summary["report_text"]


def test_merge_copies_each_cells_own_style(tmp_path):
    """回归：与拆分同一个样式缓存 bug（按 id(cell.font) 缓存会取到别的格子的样式）。"""
    from openpyxl.styles import Font, PatternFill
    colors = ["FF0000", "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    for i in range(300):
        ws.append([f"人{i}", i])
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=colors[i % 6])
            c.font = Font(bold=i % 3 == 0, color=colors[(i + 2) % 6])
    src = tmp_path / "src.xlsx"
    wb.save(src)
    out = collect.merge_tables({"input_path": str(src)})
    a, b = openpyxl.load_workbook(src).active, openpyxl.load_workbook(out).worksheets[0]

    def sig(c):
        return c.fill.fgColor.rgb, c.font.bold, c.font.color.rgb

    assert all(sig(a.cell(r, c)) == sig(b.cell(r, c)) for r in range(2, 302) for c in (1, 2))
