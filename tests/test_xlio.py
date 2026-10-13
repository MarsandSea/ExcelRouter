"""core/xlio.py 的测试：写值防公式、错误格与整数的读法、.xls 读进内存、打包，
以及「内部技能能原样带上」的清洁度（没有会被内部打包和 SkillHub 审核拦下的导入和字眼）。

Copyright (c) 2026 AbeLin · MIT License
"""
import ast
import datetime
import os
import subprocess
import sys
import zipfile

import openpyxl

from core import xlio

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLS = os.path.join(ROOT, "tests", "golden", "fixtures", "monthly.xls")

# 内部 MobileWork 打包（tools/pack_skill.py 的 FORBIDDEN / LEAN_FORBIDDEN）和 SkillHub 审核会拦的字眼。
# 这几个文件内部技能会原样带上，一个都不能有。
BANNED = ("getpass", "getuser", "_INTEGRITY", "usage.jsonl", "msoffcrypto", "secrets", "er_guard", "ER_MW_",
          "faulthandler", "shutil", "subprocess", "os.remove", "os.unlink", "os.startfile", "os.system",
          "ProcessPoolExecutor", "os.environ", "getenv",
          "密码", "台账", "工资", "绩效", "薪", "敏感", "国家", "秘密", "保密", "涉密")
LEAN_FILES = ("core/xlio.py", "core/collect.py", "core/utils.py", "core/compare.py")


def test_put_keeps_equals_text_as_text(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    xlio.put(ws["A1"], "=A2+1")       # 表格里本来就是文字
    xlio.put(ws["A2"], 5)
    ws["A3"] = "=A2*2"                # 真公式照常直接写
    p = tmp_path / "x.xlsx"
    wb.save(p)
    r = openpyxl.load_workbook(p).active
    assert (r["A1"].value, r["A1"].data_type) == ("=A2+1", "s")
    assert (r["A2"].value, r["A2"].data_type) == (5, "n")
    assert r["A3"].data_type == "f"


def test_put_leaves_error_codes_and_plain_text_alone():
    ws = openpyxl.Workbook().active
    xlio.put(ws["A1"], "#N/A")        # 错误码和以前直接赋值一样，还是错误格
    xlio.put(ws["A2"], "普通文字")
    xlio.put(ws["A3"], "=")           # 只有一个等号，openpyxl 本来就当文字
    assert ws["A1"].data_type == "e"
    assert ws["A2"].data_type == "s"
    assert ws["A3"].data_type == "s"


def test_key_value():
    assert xlio.key_value(3.0) == 3 and isinstance(xlio.key_value(3.0), int)
    assert xlio.key_value(2.5) == 2.5
    assert xlio.key_value("#N/A") is None
    assert xlio.key_value("#DIV/0!") is None
    assert xlio.key_value("N/A") == "N/A"
    assert xlio.key_value(True) is True
    assert xlio.key_value(None) is None


def test_cell_value_reads_error_cells_as_empty(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "#DIV/0!"
    ws["A2"] = 7
    p = tmp_path / "e.xlsx"
    wb.save(p)
    r = openpyxl.load_workbook(p, data_only=True).active
    assert r["A1"].data_type == "e"
    assert xlio.cell_value(r["A1"]) is None
    assert xlio.cell_value(r["A2"]) == 7


def test_open_book_reads_xls_into_memory():
    wb, is_xls = xlio.open_book(XLS)
    assert is_xls
    ws = wb["数据"]
    assert [c.value for c in ws[1]] == ["部门", "姓名", "数量", "单价", "日期", "备注"]
    assert ws["C2"].value == 3 and isinstance(ws["C2"].value, int)    # .xls 里的整数读回整数
    assert ws["D2"].value == 2.5
    assert ws["E2"].value == datetime.datetime(2026, 10, 1)
    wb.close()


def test_open_book_reads_cached_values(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["部门", "金额"])
    wb.active.append(["销售部", 12])
    p = tmp_path / "v.xlsx"
    wb.save(p)
    book, is_xls = xlio.open_book(str(p))
    assert not is_xls
    assert book.active["B2"].value == 12
    book.close()


def test_zip_folder(tmp_path):
    d = tmp_path / "销售部"
    (d / "汇总").mkdir(parents=True)
    (d / "汇总" / "a.xlsx").write_bytes(b"x")
    (d / "b.txt").write_text("y", encoding="utf-8")
    assert xlio.zip_folder(str(d), str(tmp_path / "销售部.zip")) == 2
    with zipfile.ZipFile(tmp_path / "销售部.zip") as z:
        assert sorted(z.namelist()) == ["b.txt", "汇总/a.xlsx"]
    # 包放在目录里面也不会把自己打进去
    assert xlio.zip_folder(str(d), str(d / "自己.zip")) == 2
    with zipfile.ZipFile(d / "自己.zip") as z:
        assert "自己.zip" not in z.namelist()


def test_is_total_row():
    assert xlio.is_total_row([None, "合计", None, 100])
    assert xlio.is_total_row(["全市合计", 5])
    assert not xlio.is_total_row(["张三", "销售部", 100])


def test_header_detection_reexported_from_splitter():
    from core import collect, splitter
    assert splitter.resolve_header_row is xlio.resolve_header_row
    assert splitter.detect_header_row_auto is xlio.detect_header_row_auto
    assert splitter.find_header_row is xlio.find_header_row
    assert collect.col_keys is xlio.col_keys and collect.hidden_info is xlio.hidden_info


def test_lean_modules_have_no_banned_words():
    for rel in LEAN_FILES:
        with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
            text = f.read()
        hits = [w for w in BANNED if w in text]
        assert not hits, f"{rel} 里有内部打包 / 审核会拦的字眼：{hits}"


def test_lean_modules_do_not_import_heavy_or_risky_modules():
    for rel in LEAN_FILES:
        with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
            tree = ast.parse(f.read())
        mods = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                mods.update(a.name.split(".")[0] for a in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                mods.add(node.module.split(".")[0])
        bad = mods & {"pandas", "numpy", "shutil", "subprocess"}
        assert not bad, f"{rel} 导入了 {bad}"


def test_xlio_does_not_import_pandas():
    code = "import sys; import core.xlio; print('pandas' in sys.modules)"
    out = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True, check=True)
    assert out.stdout.strip() == "False"
