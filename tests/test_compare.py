"""core/compare.py 的测试：比较口径、关键列推荐、配对与重复、核对结果工作簿。

Copyright (c) 2026 AbeLin · MIT License
"""
import datetime

import openpyxl

from core import compare as cp


def _xlsx(path, rows, title="Sheet1", top=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in list(top) + list(rows):
        ws.append(list(r))
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f":          # 写着「=…」的文字，不是公式
                c.data_type = "s"
    wb.save(path)
    return str(path)


def test_same_numbers_dates_text():
    assert cp.same(100, "100.00")
    assert cp.same("1,000", 1000)
    assert cp.same("¥1,000.5", 1000.5)
    assert cp.same("10%", 0.1)
    assert cp.same(100, 100.004)
    assert not cp.same(100, 100.01)
    assert cp.same(datetime.datetime(2026, 10, 1), "2026-10-01")
    assert cp.same(datetime.date(2026, 10, 1), "2026年10月1日")
    assert cp.same("ＡＢＣ １", "ABC 1")
    assert cp.same(" 东城 ", "东城")
    assert cp.same(None, "")
    assert cp.same(None, "  ")
    assert not cp.same(0, None)
    assert not cp.same("东城", "西城")


def test_long_digit_text_and_leading_zero_stay_text():
    a, b = "110101199001011234", "110101199001011235"     # 18 位：按数值比会丢精度变成一样
    assert not cp.same(a, b)
    assert not cp.same("00123", 123)
    assert cp.same("00123", "00123")


def test_key_text():
    assert cp.key_text(123) == cp.key_text("123") == cp.key_text(123.0)
    assert cp.key_text("00123") != cp.key_text(123)
    assert cp.key_text(" a 01 ") == "A01"
    assert cp.key_text("１２３") == "123"
    assert cp.key_text(None) == ""


def test_diff():
    assert cp.diff("1,000", 1200) == 200
    assert cp.diff(1.5, 1.25) == -0.25
    assert cp.diff("东城", 1) is None


HEAD = ("序号", "工号", "姓名", "部门", "金额", "日期")


def _rows(n=6):
    return [(i, f"E{i:03d}", f"人{i}", "东城" if i % 2 else "西城", 100.5 + i, datetime.date(2026, 9, i))
            for i in range(1, n + 1)]


def test_suggest_prefers_id_column_and_skips_serial(tmp_path):
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [HEAD] + _rows()))
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [HEAD] + _rows()))
    keys = cp.suggest_keys(a, b)
    assert keys[0]["cols"] == ["工号"]
    assert all(k["cols"] != ["姓名"] or k["score"] < keys[0]["score"] for k in keys)
    assert all(k["cols"] != ["序号"] for k in keys)
    assert all(k["cols"] != ["金额"] for k in keys)


def test_suggest_falls_back_to_pair(tmp_path):
    head = ("姓名", "部门", "金额")
    rows = [("张三", "东城", 1), ("张三", "西城", 2), ("李四", "东城", 3), ("李四", "西城", 4)]   # 金额碰巧不重复也不推荐
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [head] + rows))
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [head] + rows))
    assert cp.suggest_keys(a, b)[0]["cols"] == ["姓名", "部门"]


def test_compare_reordered_columns_and_header_rows(tmp_path):
    rows = _rows()
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [HEAD] + rows))
    # B：表头在第 3 行，列顺序打乱；E002 金额改了，E003 部门改了，E006 没了，E007 新增；还有一行合计
    order = (3, 1, 0, 5, 4, 2)
    brows = [list(r) for r in rows[:5]] + [[7, "E007", "人7", "东城", 107.5, datetime.date(2026, 9, 7)]]
    brows[1][4] = 999
    brows[2][3] = "北城"
    brows[3][4] = f"{brows[3][4]:,.2f}"            # 同一个数写成带千分位的文字：不算变化
    b_rows = [[r[i] for i in order] for r in brows] + [["合计", None, None, None, 1000, None]]
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [[HEAD[i] for i in order]] + b_rows,
                            top=[("九月明细",), ()]))
    assert b["h"] == 3 and b["totals"] == 1
    res = cp.compare(a, b, ["工号"])
    c = cp.counts(res)
    assert c == {"rows_a": 6, "rows_b": 6, "matched": 5, "only_a": 1, "only_b": 1, "changed_rows": 2,
                 "changed_cells": 2, "unpaired": 0}
    assert res["per_col"] == {"金额": 1, "部门": 1}
    assert a["rows"][res["only_a"][0]][1][1] == "E006"


def test_duplicate_and_empty_keys_are_not_paired(tmp_path):
    head = ("工号", "金额")
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [head, ("E1", 1), ("E2", 2), ("E2", 3), (None, 4), ("E3", 5)]))
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [head, ("E1", 1), ("E2", 2), ("E3", 6)]))
    res = cp.compare(a, b, ["工号"])
    assert cp.counts(res)["unpaired"] == 4          # A 的两行 E2 + A 的空工号 + B 的 E2
    assert cp.counts(res)["only_b"] == 0            # B 的 E2 不会被当成「只在 B」
    assert res["per_col"] == {"金额": 1}


def test_columns_only_on_one_side(tmp_path):
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [("工号", "金额", "备注"), ("E1", 1, "x")]))
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [("工号", "金额", "新列"), ("E1", 1, "y")]))
    res = cp.compare(a, b, ["工号"])
    assert res["only_a_cols"] == ["备注"] and res["only_b_cols"] == ["新列"]
    assert not res["changes"]


def test_write_result(tmp_path):
    head = ("工号", "姓名", "金额")
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [head, ("E1", "张三", 100), ("E2", "李四", 200), ("E3", "=1+1", 1)]))
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [head, ("E1", "张三", 150), ("E2", "李四", 200), ("E4", "王五", 9)]))
    res = cp.compare(a, b, ["工号"])
    out = cp.write_result(res, a, b, str(tmp_path / "核对.xlsx"), "a.xlsx", "b.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == list(cp.SHEETS)
    ch = list(wb["有变化"].iter_rows(values_only=True))
    assert ch[1][:5] == ("E1", "金额", 100, 150, 50)
    row = wb["变化行"][2]
    assert row[2].value == 150 and row[2].fill.fgColor.rgb.endswith("FFC7CE")
    assert row[1].fill.fill_type is None
    assert "金额：100" in row[3].value
    only_a = list(wb["只在A"].iter_rows(values_only=True))
    assert only_a[1][:2] == ("E3", "=1+1")
    assert wb["只在A"]["B2"].data_type == "s"       # 「=…」文字没有变成公式
    assert [r[1] for r in wb["只在B"].iter_rows(min_row=2, values_only=True)] == ["王五"]


def test_read_table_picks_sheet_and_skips_hidden(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "说明"
    ws.sheet_state = "hidden"
    ws.append(["工号", "金额"])
    ws.append(["X", 1])
    ws2 = wb.create_sheet("数据")
    ws2.append(["工号", "金额"])
    ws2.append(["E1", 1])
    p = str(tmp_path / "s.xlsx")
    wb.save(p)
    assert cp.read_table(p)["sheet"] == "数据"
    assert cp.read_table(p, sheet="说明")["sheet"] == "说明"
    assert cp.data_sheets(p) == ["数据"]


def test_read_table_stops_when_too_big(tmp_path):
    p = _xlsx(tmp_path / "big.xlsx", [("工号", "金额")] + [(f"E{i}", i) for i in range(100)])
    try:
        cp.read_table(p, max_cells=50)
    except cp.TooBig:
        pass
    else:
        raise AssertionError("应当报 TooBig")
    assert len(cp.read_table(p, max_cells=1000)["rows"]) == 100


def test_numeric_id_column_still_suggested(tmp_path):
    head = ("编号", "数量")
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [head] + [(1000 + i, i * 3) for i in range(10)]))
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [head] + [(1000 + i, i * 3) for i in range(10)]))
    assert [k["cols"] for k in cp.suggest_keys(a, b)] == [["编号"]]


def test_a_few_duplicate_keys_still_suggested(tmp_path):
    rows = [(f"E{i}", i) for i in range(40)]
    a = cp.read_table(_xlsx(tmp_path / "a.xlsx", [("工号", "数量")] + rows + [rows[0]]))
    b = cp.read_table(_xlsx(tmp_path / "b.xlsx", [("工号", "数量")] + rows))
    k = cp.suggest_keys(a, b)[0]
    assert k["cols"] == ["工号"] and k["dup_a"] == 2
